#!/usr/bin/env python3
"""Rebuild 关键代码.docx for thesis appendix — excerpts from live project sources."""

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DOCX_PATH = Path(r'c:\Users\ygzn\Desktop\关键代码.docx')
BACKUP_PATH = Path(r'c:\Users\ygzn\Desktop\关键代码.docx.bak')

FONT_NAME = 'Times New Roman'
FONT_SIZE = 12


def add_line(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = 'Normal'
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(FONT_SIZE)
    r = run._element.rPr
    if r is not None:
        r.rFonts.set(qn('w:eastAsia'), FONT_NAME)


def clear_document(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        body.remove(child)


def read_lines(rel_path: str, start: int, end: int) -> list[str]:
    """Extract inclusive 1-based line range from project file."""
    path = ROOT / rel_path.replace('/', '\\') if '\\' not in rel_path else ROOT / rel_path
    path = ROOT / rel_path
    text = path.read_text(encoding='utf-8')
    lines = text.splitlines()
    return lines[start - 1 : end]


def omit() -> list[str]:
    return ['// … 中间实现省略 …']


def build_sections() -> list[str]:
    intro = [
        '基于自进化 LLM 的期货市场多智能体决策模拟系统 — 关键代码摘录',
        '项目路径：futures-agent-simulation/（agent-world-codemoo）',
        '说明：以下代码与仓库源文件一致（按行摘录）；已去除像素世界/CLI/前端杂项。',
        '',
    ]

    sections: list[tuple[str, str, list[str]]] = []

    # --- 代码1 ---
    sections.append((
        '代码1 历史行情数据下载',
        'scripts/download-commodity-data.js',
        read_lines('scripts/download-commodity-data.js', 10, 90),
    ))

    # --- 代码2 ---
    code2 = (
        read_lines('scripts/process-commodity-data.js', 11, 26)
        + omit()
        + read_lines('scripts/process-commodity-data.js', 96, 154)
        + ['']
        + ['// --- adapter/localDataLoader.js：按日期查询本地历史价 ---']
        + read_lines('adapter/localDataLoader.js', 15, 36)
        + omit()
        + read_lines('adapter/localDataLoader.js', 74, 95)
        + ['']
        + ['// --- adapter/simulationEngine.js：情景价与本地真实价合并 ---']
        + read_lines('adapter/simulationEngine.js', 665, 699)
    )
    sections.append(('代码2 历史行情数据处理与本地加载', 'scripts/process-commodity-data.js、adapter/localDataLoader.js、adapter/simulationEngine.js', code2))

    # --- 代码3 ---
    code3 = (
        read_lines('adapter/simulationEngine.js', 4, 97)
        + ['']
        + read_lines('adapter/simulationEngine.js', 99, 124)
        + ['']
        + read_lines('adapter/simulationEngine.js', 185, 186)
        + ['']
        + ['// --- adapter/tradingRounds.js：36 个历史情景（第 1 轮完整，其余见源文件）---']
        + read_lines('adapter/tradingRounds.js', 5, 21)
        + ['  // … 共 36 轮，完整定义见 adapter/tradingRounds.js']
    )
    sections.append(('代码3 多智能体设定与 36 轮历史情景', 'adapter/simulationEngine.js、adapter/tradingRounds.js', code3))

    # --- 代码4 ---
    code4 = (
        read_lines('adapter/simulationEngine.js', 265, 284)
        + omit()
        + read_lines('adapter/simulationEngine.js', 359, 430)
        + omit()
        + read_lines('adapter/simulationEngine.js', 1437, 1465)
    )
    sections.append(('代码4 Prompt 自进化（训练段每 3 轮迭代）', 'adapter/simulationEngine.js', code4))

    # --- 代码5 ---
    code5 = (
        read_lines('adapter/simulationEngine.js', 1476, 1495)
        + ['']
        + read_lines('adapter/simulationEngine.js', 1701, 1748)
        + ['']
        + ['// --- server/llmClient.js：记忆审计官 Prompt 与 DeepSeek 请求参数 ---']
        + read_lines('server/llmClient.js', 95, 115)
        + omit()
        + read_lines('server/llmClient.js', 122, 211)
    )
    sections.append(('代码5 记忆审计型智能体：同业审阅与独立决策', 'adapter/simulationEngine.js、server/llmClient.js', code5))

    # --- 代码6 ---
    sections.append((
        '代码6 智能体辩论场',
        'adapter/simulationEngine.js — buildDebateRound()',
        read_lines('adapter/simulationEngine.js', 1091, 1169),
    ))

    # --- 代码7 ---
    code7 = (
        read_lines('adapter/simulationEngine.js', 702, 714)
        + ['']
        + read_lines('adapter/simulationEngine.js', 970, 1088)
    )
    sections.append(('代码7 策略评估计分', 'adapter/simulationEngine.js', code7))

    # --- 代码8 ---
    code8 = (
        read_lines('adapter/simulationEngine.js', 1529, 1557)
        + omit()
        + read_lines('adapter/simulationEngine.js', 1719, 1792)
    )
    sections.append(('代码8 多智能体单轮仿真主流程（自进化 + 协同）', 'adapter/simulationEngine.js — updateTradingGameWithLlm()', code8))

    lines = intro[:]
    for title, file_comment, code_lines in sections:
        lines.append(title)
        lines.append(f'// 文件：{file_comment}')
        lines.extend(code_lines)
        lines.append('')
    return lines


def main() -> None:
    if not DOCX_PATH.exists():
        raise SystemExit(f'Missing: {DOCX_PATH}')

    shutil.copy2(DOCX_PATH, BACKUP_PATH)
    print(f'Backup: {BACKUP_PATH}')

    content = build_sections()
    doc = Document(str(DOCX_PATH))
    clear_document(doc)
    for line in content:
        add_line(doc, line)
    doc.save(str(DOCX_PATH))
    print(f'Updated: {DOCX_PATH} ({len(content)} paragraphs)')


if __name__ == '__main__':
    main()
