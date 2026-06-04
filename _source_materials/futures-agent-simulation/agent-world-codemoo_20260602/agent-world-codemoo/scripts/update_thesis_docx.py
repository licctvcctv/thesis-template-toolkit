#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update 论文过程稿5.27.docx per confirmed thesis revision plan."""

import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from thesis_chapter2_insert import blocks_before_27
from thesis_metrics import aggregate_stats, rankings_ordered, worst_failures

DOCX = Path(r'c:\Users\ygzn\Desktop\论文过程稿5.27.docx')
BACKUP = Path(r'c:\Users\ygzn\Desktop\论文过程稿5.27.bak.docx')

TITLE = '基于自进化LLM的期货市场多智能体决策模拟系统'
OLD_TITLE = '基于多智能体的期货交易模拟系统设计与实现'


def sanitize_xml_text(text: str) -> str:
    return ''.join(c for c in text if c in '\n\t' or ord(c) >= 32)


def insert_paragraph_after(paragraph, text='', style=None):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(sanitize_xml_text(text))
    return new_para


def insert_paragraph_before(paragraph, text='', style=None):
    new_p = OxmlElement('w:p')
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(sanitize_xml_text(text))
    return new_para


def replace_in_text(text: str) -> str:
    if not text:
        return text
    t = text
    t = t.replace(f'《{OLD_TITLE}》', f'《{TITLE}》')
    t = t.replace(OLD_TITLE, TITLE)
    t = t.replace('基于多智能体的期货交易模拟系统展开', f'基于{TITLE}展开')
    t = t.replace('22轮共88条', '36轮共144条')
    t = t.replace('22轮、88条', '36轮、144条')
    t = t.replace('统计22轮', '统计36轮')
    t = t.replace('第22轮', '第36轮')
    t = t.replace('至第22轮', '至第36轮')
    t = t.replace('22轮', '36轮')
    t = t.replace('88条', '144条')
    t = t.replace('训练段3轮', '训练段6轮')
    t = t.replace('测试段19轮', '测试段30轮')
    t = t.replace('训练段3轮、测试段19轮', '训练段6轮、测试段30轮')
    t = t.replace('13/22轮', '14/35轮')
    t = t.replace('createTradingGameDemo', 'createSimulationEngine')
    t = t.replace('tradingGameDemo', 'simulationEngine')
    t = t.replace('TradingHud', 'MarketStrategyPanel')
    t = t.replace('DeepSeek-V4-Flash', 'deepseek-v4-flash')
    t = t.replace('Qwen3.6', 'deepseek-v4-pro')
    # metrics
    stats = aggregate_stats()
    rows = rankings_ordered()
    t = t.replace('35.3%', f'{stats["avg_hit"]}%')
    t = t.replace('39.14', str(stats['avg_mdd']))
    t = t.replace('86.34', str(stats['alpha_spread']))
    t = t.replace('22.72', str(rows[0]['total']))
    t = t.replace('25.59', str(rows[0]['total']))
    t = t.replace('原始22.72', f'相对{rows[0]["total"]}')
    t = t.replace('相对25.59', f'相对{rows[0]["total"]}')
    t = t.replace('测试9.84', f'测试{rows[0]["test"]}')
    t = t.replace('训练15.76', f'训练{rows[0]["train"]}')
    t = t.replace('测试7.63', f'测试{rows[1]["test"]}')
    t = t.replace('回撤55.17', f'回撤{rows[1]["mdd"]}')
    t = t.replace('测试-17.17', f'测试{rows[3]["test"]}')
    t = t.replace('买入30次、卖出22次、观望36次',
                  f'买入{stats["actions"]["BUY"]}次、卖出{stats["actions"]["SELL"]}次、观望{stats["actions"]["HOLD"]}次')
    t = t.replace('进入采纳池', '相对排名第一')
    t = t.replace('给出淘汰或修正建议', '排名靠后需修正')
    t = t.replace('保留观察并加强', '排名第二并加强')
    t = t.replace('采纳、观察或淘汰', '第一至第四排名')
    t = t.replace('策略状态划分为采纳、观察或淘汰', '策略按相对总分排序为第一至第四')
    t = t.replace('建议标为「观察」', '排名第二')
    t = t.replace('建议标为「淘汰」', '排名第四')
    return t


def walk_replace(doc: Document):
    for p in doc.paragraphs:
        if p.text:
            p.text = replace_in_text(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text:
                        p.text = replace_in_text(p.text)


def update_table_6_3(doc: Document):
    rows = rankings_ordered()
    # table index 4 = 表6-3
    if len(doc.tables) < 5:
        return
    table = doc.tables[4]
    for i, r in enumerate(rows):
        row = table.rows[i + 1]
        row.cells[0].text = f'{r["rank"]}\n{r["role"]}\n{r["symbol"]}'
        row.cells[1].text = f'相对{r["total"]}\n原始{round(r["raw"], 2)}'
        row.cells[2].text = f'训练{r["train"]}\n测试{r["test"]}'
        row.cells[3].text = f'胜率{r["hit"]}%\n回撤{r["mdd"]}'
        row.cells[4].text = r['rank']


def update_table_6_4(doc: Document):
    if len(doc.tables) < 6:
        return
    table = doc.tables[5]
    ag = rankings_ordered()[0]
    # find trades from summary - use hardcoded key points from metrics run
    points = [
        ('起点', '2016-02-11\n开始', '累计0.00', '36轮实验起点'),
        ('训练段结束', '2018-05-09\n买入', '单轮-18.00\n累计3.46', '第7轮 promptv2 后相对累计'),
        ('阶段高点', '2022-02-24\n买入', '单轮22.00\n累计116.10', '测试段相对优势扩大'),
        ('最终节点', '2025-10-08\n买入', '相对100.7\n原始183', '36轮结束排名第一'),
    ]
    for i, pt in enumerate(points):
        if i + 1 < len(table.rows):
            row = table.rows[i + 1]
            for j, val in enumerate(pt):
                if j < len(row.cells):
                    row.cells[j].text = val


def update_table_6_5(doc: Document):
    if len(doc.tables) < 7:
        return
    table = doc.tables[6]
    fails = worst_failures()
    order = ['aggressive-breakout', 'balanced-strategist', 'conservative-hedger', 'memory-auditor']
    for i, aid in enumerate(order):
        f = fails.get(aid, {})
        if i + 1 >= len(table.rows):
            break
        row = table.rows[i + 1]
        row.cells[0].text = f.get('role', '')
        row.cells[1].text = f'第{f.get("round", "")}轮\n{f.get("delta", "")}'
        row.cells[2].text = (
            f'{f.get("role", "")}累计{f.get("neg_rounds", "")}/{f.get("total_rounds", 35)}轮失分，'
            f'最大失分来自第{f.get("round", "")}轮{f.get("action", "")}{f.get("symbol", "")}。'
        )
        row.cells[3].text = '将最大失分写入记忆检索；训练段输出 strategyPatchNotes；测试段验证泛化。'


def update_table_6_1_tc04(doc: Document):
    if len(doc.tables) < 3:
        return
    table = doc.tables[2]
    for row in table.rows:
        if 'TC04' in row.cells[0].text:
            row.cells[1].text = '统计分析'
            row.cells[2].text = '执行第36轮后检查 evaluation.analysis'
            break


def update_abstract(doc: Document):
    for p in doc.paragraphs:
        if '本文围绕' in p.text and '展开研究' in p.text:
            p.text = (
                f'随着人工智能技术与金融数据基础设施的发展，面向真实行情的可解释策略验证逐渐从单一回测转向'
                f'多智能体、多情景与可追踪 Prompt 进化的综合模拟。本文围绕「{TITLE}」展开研究，'
                f'构建接入黄金/原油/大豆真实月度价格、组织四类 LLM 智能体在 36 个历史情景中协同决策，'
                f'并在前端实时展示辩论、审计、评分与 Prompt 迭代的仿真系统。'
            )
        if p.text.startswith('系统采用Node.js') and 'WebSocket' in p.text:
            p.text = replace_in_text(p.text) + (
                ' 系统启用 DeepSeek 大模型生成结构化 JSON 决策；训练段（1–6轮）支持 strategyPatchNotes 自进化，'
                '测试段（7–36轮）冻结 Prompt 检验泛化；记忆审计官审阅 peerDecisions 后独立决策。'
            )
        if p.text.startswith('关键词：'):
            p.text = '关键词：自进化大语言模型；多智能体；期货交易；Prompt迭代；策略评估；WebSocket'


def find_body_heading(doc: Document, prefix: str, level=2):
    """Match chapter body headings, skip TOC (toc 1 / toc 2)."""
    for p in doc.paragraphs:
        name = (p.style.name or '') if p.style else ''
        if 'toc' in name.lower():
            continue
        t = p.text.strip()
        if t.startswith(prefix) and (level == 2 and 'Heading 2' in name or level == 3 and 'Heading 3' in name):
            return p
    return None


def remove_orphan_toc_insertions(doc: Document):
    """Remove theory block mistakenly placed before正文第一章（index < 180）."""
    paras = list(doc.paragraphs)
    to_del = []
    for idx, p in enumerate(paras):
        if p.text.strip() != '2.4.1 自博弈数据收集与轨迹表示' or idx >= 180:
            continue
        to_del.append(p)
        j = idx - 1
        while j >= 0:
            prev = paras[j]
            ps = prev.text.strip()
            if prev.style and 'toc' in (prev.style.name or '').lower():
                break
            if (
                ps.startswith('2.4.') or ps.startswith('2.5 ')
                or ps.startswith('程序2-') or ps.startswith('（2.')
                or ps.startswith('式（2.') or ps.startswith('（本节在原')
                or '自博弈' in ps or 'Prompt' in ps
            ):
                to_del.append(prev)
                j -= 1
            else:
                break
    seen = set()
    for p in to_del:
        el_id = id(p._element)
        if el_id in seen:
            continue
        seen.add(el_id)
        p._element.getparent().remove(p._element)


def expand_chapter2(doc: Document):
    body_hits = [
        p for p in doc.paragraphs
        if p.text.strip() == '2.4.1 自博弈数据收集与轨迹表示'
        and p.style and 'Heading 3' in p.style.name
        and doc.paragraphs.index(p) > 200
    ]
    if body_hits:
        return
    target = find_body_heading(doc, '2.7', level=2)
    if not target:
        return
    # Insert new theory block immediately BEFORE §2.7 小结
    pending = []
    for kind, text in blocks_before_27():
        style = kind if kind and kind.startswith('Heading') else None
        pending.append((style, text))
    for style, text in pending:
        insert_paragraph_before(target, text, style=style)


def expand_early_chapter2(doc: Document):
    """Thicken §2.1–2.3 per advisor feedback."""
    extras = {
        '2.1多智能体系统基础': [
            '从 MAS 形式化视角，环境可建模为序贯决策过程：环境状态 s∈S 由价格向量、波动水平、新闻事件与记忆检索结果组成；智能体 i∈N 在策略 π_i 下选择动作 a∈{BUY,SELL,HOLD} 与品种 c∈{黄金,原油,大豆}。本文不追求解析均衡解，而通过并行仿真比较 π_i 在历史样本上的收益—风险表现。',
            '四类智能体的分工具有明确金融含义：保守型提供风险下界，激进型探索趋势上界，均衡型在因子冲突时折中，记忆审计型对拥挤交易与历史失败模式进行第三方约束。多智能体并行的价值在于形成“策略谱”，而非提前假定某一类永远最优。',
        ],
        '2.2大模型智能体基础': [
            '大模型智能体可视为“语言策略函数” f_θ：将结构化市场上下文映射为 JSON 决策。与传统因子模型不同，f_θ 可同时处理非结构化新闻与结构化价格；与黑箱深度学习不同，本文强制 JSON 输出并记录 reason/evidence，使推理结果可进入评分与失败样本管线。',
            'ReAct/工具调用思想在系统中体现为工作站链路：资讯、记忆、价格、风控、辩论、执行六类信息源被组织进 Prompt，而非让模型直接访问未声明的外部 API。DeepSeek 客户端对审计官启用 thinking 模式，提高 peer 审阅深度，但对外仍只提交结构化决策字段。',
        ],
        '2.3博弈论与MARL基础': [
            '将四智能体同一轮决策视为一次 n 人博弈，策略组合 a=(a_1,…,a_4) 决定各自收益 u_i(a)。纳什均衡概念用于解释“为何不能假设三体同向即安全”：若保守、均衡、激进同时 BUY，审计官从博弈视角有动机选择 HOLD 或反向以打破拥挤均衡。',
            'MARL 抽象为马尔可夫博弈 ⟨N,S,{A_i},{P},{r_i},γ⟩。本文采用“轨迹来自 LLM 自博弈 + 评分来自真实价格”的混合范式：不更新神经网络权重，而更新 Prompt 文本与 patchNotes，可看作语言空间上的策略迭代，与 Reflexion/自进化 LLM 文献一致。',
        ],
    }
    for key, paras in extras.items():
        p = find_body_heading(doc, key, level=2)
        if not p:
            continue
        for t in reversed(paras):
            insert_paragraph_before(p, t)


def patch_para_63(doc: Document):
    """§6.1 36轮 experiment paragraph."""
    for p in doc.paragraphs:
        if '本次评估统计至' in p.text:
            p.text = (
                '本次评估基于 runs/预设运行版本 的 36 轮完整 LLM 实验（运行 ID：2026-06-01T13-38-51-040Z_b98e03，'
                '模型 deepseek-v4-flash）。训练区间为第 1–6 轮（2016-02-11→2018-05-09），'
                '测试区间为第 7–36 轮（2018-05-09→2026-01-12）。系统共形成 144 条智能体交易决策（36 轮×4 智能体），'
                '评分方法为基于模型真实决策的累计评分（含少量 LLM 不可用时的预设回退补全）。'
            )
        if '表6-3显示' in p.text:
            r0, r1, r2, r3 = rankings_ordered()
            p.text = (
                f'表6-3显示，36轮完整实验下激进型智能体相对排名第一（相对总分{r0["total"]}，'
                f'训练{r0["train"]}/测试{r0["test"]}，胜率{r0["hit"]}%，回撤{r0["mdd"]}）。'
                f'均衡型排名第二（相对{r1["total"]}，测试段{r1["test"]}为正，回撤{r1["mdd"]}）。'
                f'保守型排名第三（相对{r2["total"]}），记忆审计型排名第四（相对{r3["total"]}）。'
                f'四者差异说明必须同时观察收益、回撤与训练/测试分段，而不能只看单轮盈利。'
            )
        if '系统累计统计22轮' in p.text or '系统累计统计36轮' in p.text:
            r0 = rankings_ordered()[0]
            p.text = (
                f'从智能体评价结果看，系统累计统计36轮、144条交易决策。'
                f'激进型智能体相对排名第一（相对总分{r0["total"]}，训练{r0["train"]}，测试{r0["test"]}）。'
                f'均衡型排名第二且测试段为正；保守型与记忆审计型排名靠后，需优化审计记忆权重与观望阈值。'
            )
        if '界面中间为可视化地图' in p.text:
            p.text = (
                '系统主界面采用 AppShell 分页布局：总览、智能体、市场行情、工作站、策略评估与运行日志。'
                'WebSocket 推送 worldState；策略评估页（MarketStrategyPanel）展示四名智能体相对排名（第一至第四）、'
                '训练/测试得分、收益曲线与动作分布。'
            )
        if 'constexpress=require' in p.text.replace(' ', ''):
            p.text = (
                '（伪代码）const { createSimulationEngine } = require("../adapter/simulationEngine");\n'
                'const wss = new WebSocket.Server({ server });\n'
                'tradingDemo = createSimulationEngine({ worldState, llmClient, runRecorder });'
            )
        if 'constAGENT_PROFILES=Object.freeze' in p.text.replace(' ', ''):
            p.text = (
                '（伪代码，完整实现见附录）AGENT_PROFILES 含 role、mandate、style、riskFramework、'
                'decisionPrinciples、commodityBias、modelChain；每轮 LLM 输出 commodity+action+confidence+'
                'strategyPatchNotes（训练段）+ memoryLesson；审计官额外读取 peerDecisions。'
            )
        if 'functionscoreAgent' in p.text.replace(' ', ''):
            p.text = (
                '（伪代码）buildLlmEvaluation() 读取 llmDecisionLedger，按 presetCommodityForRound 补全品种，'
                '调用 calculateTradeScore 累计 Raw 与相对 alpha，输出 rankings（第一至第四）。'
            )
        if 'functionrenderPredictionReport' in p.text.replace(' ', ''):
            p.text = (
                '（伪代码）MarketStrategyPanel 根据 meta.trading.evaluation 渲染排名、得分柱、收益曲线与动作分布；'
                '不再使用采纳/淘汰标签，而以第一至第四相对排名展示。'
            )


def main():
    if not DOCX.exists():
        raise SystemExit(f'Missing {DOCX}')
    shutil.copy2(DOCX, BACKUP)
    print('Backup:', BACKUP)

    doc = Document(str(DOCX))
    update_abstract(doc)
    walk_replace(doc)
    remove_orphan_toc_insertions(doc)
    expand_early_chapter2(doc)
    expand_chapter2(doc)
    patch_para_63(doc)
    update_table_6_3(doc)
    update_table_6_4(doc)
    update_table_6_5(doc)
    update_table_6_1_tc04(doc)

    doc.save(str(DOCX))
    print('Updated:', DOCX)
    print('Figures: c:\\Users\\ygzn\\Desktop\\论文图表重绘\\')
    print('Screenshots: c:\\Users\\ygzn\\Desktop\\论文截图清单.md')


if __name__ == '__main__':
    main()
