#!/usr/bin/env python3
"""Build the Nanjing Vocational University of Industry Technology thesis template.

The paper builder needs a docxtpl-ready template, while the school reference
document carries the official cover logo as a normal Word picture.  This helper
keeps the existing JSON/docxtpl body workflow and replaces the cover identity
with the Nanjing template assets.
"""
from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


HERE = Path(__file__).resolve().parent
TEMPLATES = HERE.parent
BASE_TEMPLATE = TEMPLATES / "hnwlxy" / "template.docx"
SOURCE_TEMPLATE = TEMPLATES / "ml_ids_source" / "template.docx"
LOGO_PATH = HERE / "njgyzyjsdx-logo.png"
OUTPUT_TEMPLATE = HERE / "template.docx"


def clear_paragraph(paragraph):
    for run in paragraph.runs:
        run.text = ""


def set_run_font(run, size_pt: float, bold=None, font="宋体"):
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.name = font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), font)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")


def set_title_page(section):
    sect_pr = section._sectPr
    title_pg = sect_pr.find(qn("w:titlePg"))
    if title_pg is None:
        title_pg = OxmlElement("w:titlePg")
        pg_mar = sect_pr.find(qn("w:pgMar"))
        if pg_mar is not None:
            sect_pr.insert(list(sect_pr).index(pg_mar), title_pg)
        else:
            sect_pr.append(title_pg)


def remove_header_style_border(section):
    section.different_first_page_header_footer = True
    first_header = section.first_page_header
    for paragraph in first_header.paragraphs:
        paragraph.style = "Normal"
        clear_paragraph(paragraph)
        p_pr = paragraph._p.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)


def extract_school_logo():
    with zipfile.ZipFile(SOURCE_TEMPLATE) as zf:
        LOGO_PATH.write_bytes(zf.read("word/media/image1.png"))


def replace_cover_logo(doc: Document):
    cover_logo = doc.paragraphs[1]
    clear_paragraph(cover_logo)
    cover_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_logo.paragraph_format.space_before = Pt(0)
    cover_logo.paragraph_format.space_after = Pt(12)
    run = cover_logo.add_run()
    run.add_picture(str(LOGO_PATH), width=Mm(165))


def replace_cover_title(doc: Document):
    title = doc.paragraphs[2]
    clear_paragraph(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(22)
    run = title.add_run("毕业论文（设计）")
    set_run_font(run, 36, bold=False, font="方正小标宋_GBK")


def set_cell_text(cell, text, size=16, bold=True):
    for idx, paragraph in enumerate(cell.paragraphs):
        if idx == 0:
            clear_paragraph(paragraph)
        else:
            paragraph._p.getparent().remove(paragraph._p)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=bold, font="宋体")


def tune_cover_table(doc: Document):
    table = doc.tables[0]
    labels = [
        "题    目：",
        "学生姓名：",
        "学    号：",
        "学    院：",
        "专    业：",
        "班    级：",
        "指导教师：",
        "完成时间：",
    ]
    values = [
        "{{ title_zh }}",
        "{{ name }}",
        "{{ student_id }}",
        "{{ college }}",
        "{{ major }}",
        "{{ class_name }}",
        "{{ advisor }}",
        "{{ finish_date }}",
    ]
    while len(table.rows) < len(labels):
        table.add_row()
    for row, label, value in zip(table.rows, labels, values):
        set_cell_text(row.cells[0], label, size=16, bold=True)
        set_cell_text(row.cells[1], value, size=14 if label == "题    目：" else 16, bold=(label != "完成时间："))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(28)
    for cell in table.rows[-1].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(42)


def main() -> int:
    if not BASE_TEMPLATE.exists():
        raise FileNotFoundError(BASE_TEMPLATE)
    if not SOURCE_TEMPLATE.exists():
        raise FileNotFoundError(SOURCE_TEMPLATE)
    extract_school_logo()
    shutil.copy2(BASE_TEMPLATE, OUTPUT_TEMPLATE)
    doc = Document(OUTPUT_TEMPLATE)
    set_title_page(doc.sections[0])
    remove_header_style_border(doc.sections[0])
    replace_cover_logo(doc)
    replace_cover_title(doc)
    tune_cover_table(doc)
    doc.save(OUTPUT_TEMPLATE)
    print(f"wrote {OUTPUT_TEMPLATE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
