"""
北京信息科技大学毕业设计（论文）模板制作。

用法:
  cd /Users/a136/vs/45425/thesis_project
  python templates/bistu_2026/make.py

本脚本从 source.docx 生成 template.docx，仅清理旧论文身份和正文，
保留封面、任务书、声明、摘要、目录前置页和正文样式。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "source.docx"
OUTPUT = HERE / "template.docx"


def set_run_text(paragraph, text: str) -> None:
    paragraph.clear()
    if text:
        paragraph.add_run(text)


def set_cell_text(cell, text: str) -> None:
    if cell.paragraphs:
        set_run_text(cell.paragraphs[0], text)
        for paragraph in cell.paragraphs[1:]:
            set_run_text(paragraph, "")
    else:
        cell.add_paragraph(text)


def replace_text_in_runs(paragraph, mapping: dict[str, str]) -> None:
    for run in paragraph.runs:
        text = run.text
        for old, new in mapping.items():
            text = text.replace(old, new)
        run.text = text


def delete_paragraph(paragraph) -> None:
    element = paragraph._p
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def has_section_break(paragraph) -> bool:
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.find(qn("w:sectPr")) is not None


def clear_numbering(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)


def clear_toc(paragraphs) -> None:
    started = False
    kept_placeholder = False
    to_delete = []
    for paragraph in paragraphs:
        text = (paragraph.text or "").strip()
        style = paragraph.style.name if paragraph.style else ""
        if text == "目录":
            started = True
            continue
        if not started:
            continue
        if style.lower().startswith("toc") or text in {"毕业设计(论文)任务书46", "摘要I", "AbstractII"}:
            if not kept_placeholder:
                set_run_text(paragraph, "{{ toc }}")
                kept_placeholder = True
            else:
                to_delete.append(paragraph)
        if text == "引言":
            break
    for paragraph in to_delete:
        if not has_section_break(paragraph):
            delete_paragraph(paragraph)
        else:
            set_run_text(paragraph, "")


def make_template() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"missing source template: {SOURCE}")

    doc = Document(str(SOURCE))
    paragraphs = doc.paragraphs

    # Cover fields. The source uses visible underline spacing, so keep labels
    # and replace only the editable values with single-run placeholders.
    cover_map = {
        13: "题    目：          {{ title_zh }}",
        16: "学    院：        {{ college }}",
        19: "专    业：        {{ major }}",
        22: "学生姓名：       {{ name }}  {{ class_name }}/{{ student_id }}",
        25: "指导老师/督导老师：        {{ advisor }}",
        28: "起止时间：  {{ start_date }} 至 {{ end_date }}",
        38: "学院 ：{{ college }}        专业：{{ major }}      班级：{{ class_name }}",
    }
    for idx, text in cover_map.items():
        if idx < len(paragraphs):
            set_run_text(paragraphs[idx], text)

    # Originality declaration contains a stale old title in the source.
    if 42 < len(paragraphs):
        set_run_text(
            paragraphs[42],
            "本人郑重声明：所呈交的毕业设计（论文），题目为《{{ title_zh }}》，是本人在导师指导下，"
            "进行研究工作所取得的成果。尽我所知，除了文中特别加以标注的内容外，本毕业设计（论文）"
            "的研究成果不包含任何他人创作的、已公开发表或者没有公开发表的作品的内容。对本毕业设计"
            "（论文）所涉及的研究工作做出贡献的其他个人和集体，均已在文中以明确方式标明并表示了谢",
        )

    # Task-book table.
    if doc.tables:
        table = doc.tables[0]
        set_cell_text(table.cell(3, 0), "{{ name }}")
        set_cell_text(table.cell(3, 1), "{{ student_id }}")
        set_cell_text(table.cell(3, 2), "{{ advisor }}")
        set_cell_text(table.cell(3, 3), "{{ advisor_title }}")
        set_cell_text(table.cell(3, 4), "{{ advisor_unit }}")
        set_cell_text(table.cell(5, 1), "{{ title_zh }}")
        set_cell_text(table.cell(6, 1), "{{ task_main_content }}")
        set_cell_text(table.cell(7, 1), "{{ task_outcomes }}")
        set_cell_text(table.cell(8, 1), "{{ task_requirements }}")
        set_cell_text(table.cell(10, 1), "{{ task_references }}")
        set_cell_text(table.cell(11, 1), "{{ task_environment }}")
        set_cell_text(table.cell(12, 2), "{{ start_date_short }}")
        set_cell_text(table.cell(12, 7), "{{ end_date_short }}")
        set_cell_text(table.cell(14, 0), "{{ task_schedule }}")

    # Abstract and keywords.
    abstract_map = {
        64: "{{ abstract_zh_1 }}",
        65: "{{ abstract_zh_2 }}",
        66: "{{ abstract_zh_3 }}",
        67: "{{ abstract_zh_4 }}",
        69: "关键词：{{ keywords_zh }}",
        72: "{{ abstract_en_1 }}",
        73: "{{ abstract_en_2 }}",
        74: "{{ abstract_en_3 }}",
        75: "{{ abstract_en_4 }}",
        77: "Key words: {{ keywords_en }}",
    }
    for idx, text in abstract_map.items():
        if idx < len(paragraphs):
            set_run_text(paragraphs[idx], text)

    clear_toc(paragraphs)

    # Collapse old body, conclusion, acknowledgements and references into a
    # single body placeholder. Keep section-break paragraphs before the body.
    paragraphs = doc.paragraphs
    body_start = None
    for idx, paragraph in enumerate(paragraphs):
        if paragraph.style and paragraph.style.name == "Heading 1" and paragraph.text.strip() == "引言":
            body_start = idx
            break
    if body_start is None:
        raise RuntimeError("cannot find body start heading")

    set_run_text(paragraphs[body_start], "{{ body }}")
    for idx in range(len(paragraphs) - 1, body_start, -1):
        paragraph = paragraphs[idx]
        if has_section_break(paragraph):
            set_run_text(paragraph, "")
        else:
            delete_paragraph(paragraph)

    paragraphs[body_start].style = doc.styles["Normal"]
    clear_numbering(paragraphs[body_start])

    stale_title_map = {
        "老旧小区共享车位停车场收费系统设计": "{{ title_zh }}",
        "老旧小区共享车位停车场收费系统": "{{ title_zh }}",
    }
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                replace_text_in_runs(paragraph, stale_title_map)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(f"template saved: {OUTPUT}")


if __name__ == "__main__":
    make_template()
