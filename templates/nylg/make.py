#!/usr/bin/env python3
"""
南阳理工学院本科生毕业设计（论文）模板制作。

用法:
    python make.py <source.docx> [template.docx]

约定:
    - source.docx: 原始范文
    - template.docx: 可渲染模板
"""
from __future__ import annotations

import copy
import os
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from refs_maker import setup_refs_template  # noqa: E402

SOURCE_BODY_H1 = 192
SOURCE_BODY_H2 = 193
SOURCE_BODY_H3 = 194
SOURCE_BODY_SAMPLE = 317


def _p(text: str):
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    return p


def _insert_before(ref_elem, new_elem):
    ref_elem.addprevious(new_elem)
    return new_elem


def _insert_after(ref_elem, new_elem):
    ref_elem.addnext(new_elem)
    return new_elem


def _clear_para(para):
    for r in para.runs:
        r.text = ""


def _replace_first_text_run(para, text):
    found = False
    for r in para.runs:
        if not found and (r.text or "").strip():
            r.text = text
            found = True
        elif found:
            r.text = ""
    if not found:
        if para.runs:
            para.runs[0].text = text
            for r in para.runs[1:]:
                r.text = ""
        else:
            para.add_run(text)


def _replace_run(para, run_idx, text, clear_other_runs=False):
    if run_idx < len(para.runs):
        para.runs[run_idx].text = text
    if clear_other_runs:
        for i, r in enumerate(para.runs):
            if i != run_idx:
                r.text = ""


def _replace_paragraph_text(para, text):
    _replace_first_text_run(para, text)


def _replace_p_elem_text(p_elem, text):
    """Replace the first visible run in a raw <w:p> element."""
    runs = p_elem.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r")
    found = False
    for r in runs:
        texts = r.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")
        if not found and any((t.text or "").strip() for t in texts):
            for i, t in enumerate(texts):
                t.text = text if i == 0 else ""
            found = True
        elif found:
            for t in texts:
                t.text = ""
    if not found and runs:
        first = runs[0]
        texts = first.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")
        if texts:
            texts[0].text = text
            for t in texts[1:]:
                t.text = ""
        else:
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            first.append(t)


def _clear_range_text(doc, start_idx, end_idx, preserve_ids):
    """Clear text in a paragraph range while preserving control/sample nodes."""
    for i, p in enumerate(doc.paragraphs):
        if i < start_idx or i >= end_idx:
            continue
        if id(p._p) in preserve_ids:
            continue
        txt = (p.text or "").strip()
        if txt.startswith("{%p") or txt.startswith("{{") or txt in {
            "参考文献",
            "致谢",
        }:
            continue
        _clear_para(p)


def _remove_all_body_tables(doc):
    body = doc.element.body
    for elem in list(body):
        if elem.tag.endswith("}tbl"):
            body.remove(elem)


def _strip_revision_markup(docx_path):
    """Remove tracked-change wrappers and other leftover review markup."""
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    unwrap = {
        f"{{{ns}}}ins",
        f"{{{ns}}}del",
        f"{{{ns}}}moveFrom",
        f"{{{ns}}}moveTo",
    }
    drop = {
        f"{{{ns}}}commentRangeStart",
        f"{{{ns}}}commentRangeEnd",
        f"{{{ns}}}commentReference",
        f"{{{ns}}}proofErr",
    }

    tmp = Path(str(docx_path) + ".tmp")
    with zipfile.ZipFile(docx_path, "r") as zin:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                    except Exception:
                        pass
                    else:
                        etree.strip_tags(root, *unwrap)
                        etree.strip_elements(root, *drop, with_tail=False)
                        data = etree.tostring(
                            root,
                            encoding="UTF-8",
                            xml_declaration=True,
                        )
                zout.writestr(info, data)
    tmp.replace(docx_path)


def _find_para_idx(doc, text, style_name=None, start=0, end=None):
    for i, p in enumerate(doc.paragraphs[start:end], start=start):
        if (p.text or "").strip() != text:
            continue
        if style_name is not None and (p.style.name if p.style else None) != style_name:
            continue
        return i
    return None


def build_body_template(doc):
    """把正文区域改成 docxtpl 循环骨架。"""
    paras = doc.paragraphs

    h1 = paras[SOURCE_BODY_H1]._p
    h2 = paras[SOURCE_BODY_H2]._p
    h3 = paras[SOURCE_BODY_H3]._p
    body_src = paras[SOURCE_BODY_SAMPLE]._p

    # 保留一个正文样本，并复制成 sec.content / sub.content 两个模板段
    body_before = copy.deepcopy(body_src)
    body_after = copy.deepcopy(body_src)

    _replace_paragraph_text(paras[SOURCE_BODY_H1], "{{ ch.title }}")
    _replace_paragraph_text(paras[SOURCE_BODY_H2], "{{ sec.title }}")
    _replace_paragraph_text(paras[SOURCE_BODY_H3], "{{ sub.title }}")

    _replace_paragraph_text(doc.paragraphs[SOURCE_BODY_SAMPLE], "{{ para }}")
    _replace_p_elem_text(body_before, "{{ para }}")
    _replace_p_elem_text(body_after, "{{ p2 }}")

    # 先移除原始正文样本，避免在后面留下一个多余段落
    body_src.getparent().remove(body_src)

    # 插入正文循环骨架
    _insert_before(h1, _p("{%p for ch in chapters %}"))
    _insert_before(h2, _p("{%p for sec in ch.sections %}"))

    # sec.content 的正文样本
    h3.addprevious(body_before)
    body_before.addprevious(_p("{%p for para in sec.content %}"))
    body_before.addnext(_p("{%p endfor %}"))

    # subsections + sub.content 的正文样本
    h3.addprevious(_p("{%p for sub in sec.subsections %}"))
    h3.addnext(body_after)
    body_after.addprevious(_p("{%p for p2 in sub.content %}"))
    body_after.addnext(_p("{%p endfor %}"))

    # 关闭 subsections / sections / chapters
    end_subsections = _p("{%p endfor %}")
    end_sections = _p("{%p endfor %}")
    end_chapters = _p("{%p endfor %}")
    body_after.getnext().addnext(end_subsections)
    end_subsections.addnext(end_sections)
    end_sections.addnext(end_chapters)

    preserve_ids = {
        id(paras[SOURCE_BODY_H1]._p),
        id(paras[SOURCE_BODY_H2]._p),
        id(paras[SOURCE_BODY_H3]._p),
        id(body_before),
        id(body_after),
    }

    # 清理正文样本以外的正文内容，保留结构性空段落和控制标签
    h1_idx = _find_para_idx(doc, "{{ ch.title }}")
    refs_idx = _find_para_idx(doc, "参考文献", start=600)
    if h1_idx is not None and refs_idx is not None and refs_idx > h1_idx:
        _clear_range_text(doc, h1_idx + 1, refs_idx, preserve_ids)

    # 删除正文区的示例表格
    _remove_all_body_tables(doc)


def make(src_path, out_path):
    doc = Document(src_path)

    # ========== Step 1: 外封 ==========
    _replace_run(doc.paragraphs[13], 4, "{{ college }}", clear_other_runs=True)
    _replace_run(doc.paragraphs[14], 4, "{{ major }}", clear_other_runs=True)
    _replace_run(doc.paragraphs[15], 6, "{{ name }}", clear_other_runs=True)
    _replace_run(doc.paragraphs[16], 8, "{{ advisor }}", clear_other_runs=True)
    _replace_run(doc.paragraphs[26], 4, "{{ year }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[26], 6, "{{ month }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[26], 2, "", clear_other_runs=False)

    # ========== Step 2: 内封 ==========
    _replace_paragraph_text(doc.paragraphs[36], "{{ title_zh }}")
    _replace_paragraph_text(doc.paragraphs[39], "{{ title_en }}")
    _replace_run(doc.paragraphs[51], 4, "{{ page_count }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[52], 1, "{{ table_count }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[53], 3, "{{ figure_count }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[53], 4, "幅", clear_other_runs=False)

    _replace_paragraph_text(doc.paragraphs[60], "南 阳 理 工 学 院 本 科 毕 业 设 计(论文)")
    _replace_paragraph_text(doc.paragraphs[65], "{{ title_zh }}")
    _replace_paragraph_text(doc.paragraphs[67], "{{ title_en }}")
    _replace_run(doc.paragraphs[71], 2, "{{ college }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[72], 2, "{{ major }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[73], 2, "{{ name }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[74], 3, "{{ student_id }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[75], 3, "{{ advisor }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[75], 5, "{{ advisor_title }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[76], 3, "{{ reviewer }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[77], 4, "{{ year }}", clear_other_runs=False)
    _replace_run(doc.paragraphs[77], 6, "{{ month }}", clear_other_runs=False)
    _replace_paragraph_text(doc.paragraphs[85], "南阳理工学院")
    _replace_paragraph_text(doc.paragraphs[86], "Nanyang Institute of Technology")
    _replace_paragraph_text(doc.paragraphs[89], "{{ title_zh }}")
    _replace_paragraph_text(doc.paragraphs[90], "{{ major }} {{ name }}")

    # ========== Step 3: 摘要 ==========
    _replace_paragraph_text(doc.paragraphs[91], "{{ abstract_zh }}")
    _replace_paragraph_text(doc.paragraphs[92], "{{ keywords_zh }}")
    _replace_paragraph_text(doc.paragraphs[108], "{{ title_en }}")
    _replace_paragraph_text(doc.paragraphs[109], "{{ major_en }} Major {{ name_en }}")
    _replace_paragraph_text(doc.paragraphs[110], "{{ abstract_en }}")
    _replace_paragraph_text(doc.paragraphs[111], "{{ keywords_en }}")

    # ========== Step 4: 正文 ==========
    build_body_template(doc)

    # ========== Step 5: 参考文献 ==========
    setup_refs_template(doc, _find_para_idx(doc, "参考文献", start=600))

    # ========== Step 6: 致谢 ==========
    ack_idx = _find_para_idx(doc, "致谢", start=600)
    if ack_idx is not None:
        for j in range(ack_idx + 1, min(ack_idx + 6, len(doc.paragraphs))):
            p = doc.paragraphs[j]
            style_name = p.style.name if p.style else ""
            if "正文" not in style_name:
                continue
            _replace_paragraph_text(p, "{{ acknowledgement }}")
            break

    doc.save(out_path)
    _strip_revision_markup(out_path)
    print(f"完成: {out_path}")


def main():
    if len(sys.argv) < 2:
        print("用法: python make.py <source.docx> [template.docx]")
        raise SystemExit(1)

    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "template.docx"
    make(src, out)


if __name__ == "__main__":
    main()
