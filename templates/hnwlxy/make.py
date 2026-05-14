#!/usr/bin/env python3
"""
湖南文理学院本科生毕业论文（设计）模板制作器。

用法:
    cd /Users/a136/vs/45425/thesis_project
    python templates/hnwlxy/make.py templates/hnwlxy/source.docx templates/hnwlxy/template.docx

模板来源: 新要求5.9/论文模板(1).docx

数据约定:
    - title_zh, name, student_id, class_name, advisor, finish_date
    - abstract_zh_list, abstract_en_list, keywords_zh, keywords_en
    - toc_placeholder
    - chapters[].content / chapters[].sections[].content / sections[].subsections[].content
    - references, acknowledgement, appendix
"""
from __future__ import annotations

import copy
import os
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": WNS}


def _p(text: str) -> OxmlElement:
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    return p


def _page_break_p() -> OxmlElement:
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


def _insert_before(ref, elem):
    ref.addprevious(elem)
    return elem


def _insert_after(ref, elem):
    ref.addnext(elem)
    return elem


def _text(elem) -> str:
    return "".join(t.text or "" for t in elem.findall(".//w:t", NS)).strip()


def _clear_para(para) -> None:
    for run in para.runs:
        run.text = ""


def _replace_para_text(para, text: str) -> None:
    """Replace paragraph visible text while preserving the first run's formatting."""
    texts = para._p.findall(".//w:t", NS)
    if texts:
        texts[0].text = text
        for t in texts[1:]:
            t.text = ""
        return
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def _replace_cell_value(cell, text: str) -> None:
    """Replace only the value part in a cover cell, keeping underline padding when present."""
    replaced = False
    for para in cell.paragraphs:
        for run in para.runs:
            raw = run.text or ""
            if not replaced and raw.strip():
                run.text = text
                replaced = True
                continue
            if replaced and raw.strip():
                run.text = ""
        if replaced:
            break


def _replace_p_elem_text(p_elem, text: str) -> None:
    texts = p_elem.findall(".//w:t", NS)
    if not texts:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        p_elem.append(r)
        return
    texts[0].text = text
    for t in texts[1:]:
        t.text = ""


def _clone_para(para, text: str):
    elem = copy.deepcopy(para._p)
    _replace_p_elem_text(elem, text)
    return elem


def _page_break_before(para) -> None:
    para.paragraph_format.page_break_before = True


def _find_para(doc: Document, predicate, start: int = 0, end: int | None = None) -> int | None:
    for idx, para in enumerate(doc.paragraphs[start:end], start=start):
        if predicate(para):
            return idx
    return None


def _find_text(doc: Document, text: str, start: int = 0) -> int | None:
    compact = text.replace(" ", "")
    return _find_para(
        doc,
        lambda p: (p.text or "").replace(" ", "").strip() == compact,
        start=start,
    )


def _find_contains(doc: Document, text: str, start: int = 0) -> int | None:
    return _find_para(doc, lambda p: text in (p.text or ""), start=start)


def _remove_between(doc: Document, start_elem, stop_elem, keep_elems) -> None:
    """Remove direct body children from start to stop, excluding kept elements."""
    body = doc.element.body
    children = list(body)
    keep_ids = {id(elem) for elem in keep_elems}
    try:
        start = children.index(start_elem)
        stop = children.index(stop_elem)
    except ValueError as exc:
        raise RuntimeError("无法定位模板 XML 片段，源模板结构可能已变化") from exc
    for elem in children[start:stop]:
        if id(elem) in keep_ids:
            continue
        body.remove(elem)


def _replace_between(doc: Document, start_elem, stop_elem, new_elems) -> None:
    """Replace direct body children between two anchor elements."""
    body = doc.element.body
    children = list(body)
    try:
        start = children.index(start_elem)
        stop = children.index(stop_elem)
    except ValueError as exc:
        raise RuntimeError("无法定位模板 XML 片段，源模板结构可能已变化") from exc
    for elem in children[start + 1 : stop]:
        body.remove(elem)
    cursor = start_elem
    for elem in new_elems:
        cursor.addnext(elem)
        cursor = elem


def _remove_after(doc: Document, start_elem, keep_elems=()) -> None:
    body = doc.element.body
    keep_ids = {id(elem) for elem in keep_elems}
    children = list(body)
    try:
        start = children.index(start_elem)
    except ValueError as exc:
        raise RuntimeError("无法定位附录 XML 片段，源模板结构可能已变化") from exc
    for elem in children[start + 1 :]:
        if elem.tag == qn("w:sectPr"):
            continue
        if id(elem) in keep_ids:
            continue
        body.remove(elem)


def setup_cover(doc: Document) -> None:
    table = doc.tables[0]
    fields = {
        0: "{{ title_zh }}",
        1: "{{ name }}",
        2: "{{ student_id }}",
        3: "{{ class_name }}",
        4: "{{ advisor }}",
        5: "{{ finish_date }}",
    }
    for row_idx, placeholder in fields.items():
        _replace_cell_value(table.rows[row_idx].cells[1], placeholder)
    # 该批论文题目较长，略收小封面题目字号，避免顶出空白页。
    for para in table.rows[0].cells[1].paragraphs:
        for run in para.runs:
            run.font.size = Pt(14)

    for section in doc.sections:
        for header in (section.header, section.first_page_header):
            for para in header.paragraphs:
                if "直流电源设计" in (para.text or "") or "XX" in (para.text or ""):
                    _replace_para_text(para, "{{ title_zh }}")


def setup_toc(doc: Document) -> None:
    toc_idx = _find_text(doc, "目录")
    abstract_idx = _find_text(doc, "摘要", start=(toc_idx or 0) + 1)
    if toc_idx is None or abstract_idx is None:
        return

    abstract_title = doc.paragraphs[abstract_idx]._p
    sample_idx = toc_idx + 1
    sample = doc.paragraphs[sample_idx]
    _replace_para_text(sample, "{{ toc_placeholder }}")
    _remove_between(doc, sample._p, abstract_title, {sample._p})
    _insert_after(sample._p, _page_break_p())


def remove_redundant_cover_page_break(doc: Document) -> None:
    """The source has both a section break and a page break after the cover."""
    toc_idx = _find_text(doc, "目录")
    if toc_idx is None:
        return
    for para in doc.paragraphs[:toc_idx]:
        has_page_break = bool(para._p.findall(".//w:br[@w:type='page']", NS))
        if has_page_break and not (para.text or "").strip():
            para._p.getparent().remove(para._p)
            return


def setup_abstracts(doc: Document) -> None:
    zh_title = _find_text(doc, "摘要")
    en_title = _find_text(doc, "ABSTRACT", start=(zh_title or 0) + 1)
    body_idx = _find_para(doc, lambda p: p.style and p.style.name == "Heading 1", start=(en_title or 0) + 1)
    if zh_title is None or en_title is None or body_idx is None:
        return
    _page_break_before(doc.paragraphs[en_title])

    zh_kw = _find_contains(doc, "关键词", zh_title + 1)
    en_kw = _find_contains(doc, "Keywords", en_title + 1)

    zh_body_idx = next(
        i
        for i in range(zh_title + 1, en_title)
        if i != zh_kw and (doc.paragraphs[i].text or "").strip()
    )
    en_body_idx = next(
        i
        for i in range(en_title + 1, body_idx)
        if i != en_kw and (doc.paragraphs[i].text or "").strip()
    )

    zh_title_elem = doc.paragraphs[zh_title]._p
    en_title_elem = doc.paragraphs[en_title]._p
    body_elem = doc.paragraphs[body_idx]._p

    zh_kw_para = doc.paragraphs[zh_kw] if zh_kw is not None else doc.paragraphs[zh_body_idx]
    zh_body_block = _clone_para(doc.paragraphs[zh_body_idx], "{{ abs_p }}")
    zh_kw_block = _clone_para(zh_kw_para, "关键词：{{ keywords_zh }}")
    en_kw_para = doc.paragraphs[en_kw] if en_kw is not None else doc.paragraphs[en_body_idx]
    en_body_block = _clone_para(doc.paragraphs[en_body_idx], "{{ abs_p }}")
    en_kw_block = _clone_para(en_kw_para, "Keywords: {{ keywords_en }}")

    zh_blocks = [
        _p("{%p for abs_p in abstract_zh_list %}"),
        zh_body_block,
        _p("{%p endfor %}"),
        zh_kw_block,
    ]
    _replace_between(doc, zh_title_elem, en_title_elem, zh_blocks)

    en_blocks = [
        _p("{%p for abs_p in abstract_en_list %}"),
        en_body_block,
        _p("{%p endfor %}"),
        en_kw_block,
    ]
    _replace_between(doc, en_title_elem, body_elem, en_blocks)


def setup_body(doc: Document) -> None:
    paras = doc.paragraphs
    body_idx = _find_para(doc, lambda p: p.style and p.style.name == "Heading 1")
    refs_idx = _find_text(doc, "参考文献", start=(body_idx or 0) + 1)
    if body_idx is None or refs_idx is None:
        return

    h1_idx = body_idx
    h2_idx = _find_para(doc, lambda p: p.style and p.style.name == "Heading 2", start=h1_idx + 1, end=refs_idx)
    h3_idx = _find_para(doc, lambda p: p.style and p.style.name == "Heading 3", start=(h2_idx or h1_idx) + 1, end=refs_idx)
    if h2_idx is None or h3_idx is None:
        return

    ch_content_idx = _find_para(
        doc,
        lambda p: p.style and p.style.name == "Normal" and (p.text or "").strip(),
        start=h1_idx + 1,
        end=h2_idx,
    )
    sec_content_idx = _find_para(
        doc,
        lambda p: p.style and p.style.name == "Normal" and (p.text or "").strip(),
        start=h2_idx + 1,
        end=h3_idx,
    )
    sub_content_idx = _find_para(
        doc,
        lambda p: p.style and p.style.name == "Normal" and (p.text or "").strip(),
        start=h3_idx + 1,
        end=refs_idx,
    )
    if ch_content_idx is None or sec_content_idx is None or sub_content_idx is None:
        return

    h1 = paras[h1_idx]._p
    h2 = paras[h2_idx]._p
    h3 = paras[h3_idx]._p
    ch_content = paras[ch_content_idx]._p
    sec_content = paras[sec_content_idx]._p
    sub_content = paras[sub_content_idx]._p
    refs_title = paras[refs_idx]._p

    _replace_para_text(paras[h1_idx], "{{ ch.title }}")
    _page_break_before(paras[h1_idx])
    _replace_para_text(paras[ch_content_idx], "{{ para }}")
    _replace_para_text(paras[h2_idx], "{{ sec.title }}")
    _replace_para_text(paras[sec_content_idx], "{{ para }}")
    _replace_para_text(paras[h3_idx], "{{ sub.title }}")
    _replace_para_text(paras[sub_content_idx], "{{ para }}")

    controls = [
        _insert_before(h1, _p("{%p for ch in chapters %}")),
        _insert_after(h1, _p("{%p for para in ch.content %}")),
        _insert_after(ch_content, _p("{%p endfor %}")),
        _insert_before(h2, _p("{%p for sec in ch.sections %}")),
        _insert_after(h2, _p("{%p for para in sec.content %}")),
        _insert_after(sec_content, _p("{%p endfor %}")),
        _insert_before(h3, _p("{%p for sub in sec.subsections %}")),
        _insert_after(h3, _p("{%p for para in sub.content %}")),
        _insert_after(sub_content, _p("{%p endfor %}")),
    ]

    end_sub = _insert_after(controls[-1], _p("{%p endfor %}"))
    end_sec = _insert_after(end_sub, _p("{%p endfor %}"))
    end_ch = _insert_after(end_sec, _p("{%p endfor %}"))
    controls.extend([end_sub, end_sec, end_ch])

    keep = {h1, h2, h3, ch_content, sec_content, sub_content, *controls}
    _remove_between(doc, h1, refs_title, keep)


def setup_references(doc: Document) -> None:
    refs_idx = _find_text(doc, "参考文献")
    ack_idx = _find_text(doc, "致谢", start=(refs_idx or 0) + 1)
    if refs_idx is None or ack_idx is None:
        return
    _page_break_before(doc.paragraphs[refs_idx])
    ack_elem = doc.paragraphs[ack_idx]._p

    sample_idx = _find_para(
        doc,
        lambda p: p.style and p.style.name == "Normal" and len((p.text or "").strip()) > 5,
        start=refs_idx + 1,
        end=ack_idx,
    )
    if sample_idx is None:
        return

    sample = doc.paragraphs[sample_idx]
    _replace_para_text(sample, "{{ ref.text if ref.text is defined else ref }}")
    start_loop = _insert_before(sample._p, _p("{%p for ref in references %}"))
    end_loop = _insert_after(sample._p, _p("{%p endfor %}"))
    _remove_between(doc, sample._p, ack_elem, {start_loop, sample._p, end_loop})


def setup_ack_and_appendix(doc: Document) -> None:
    ack_idx = _find_text(doc, "致谢")
    appendix_idx = _find_text(doc, "附录", start=(ack_idx or 0) + 1)
    if ack_idx is not None:
        _page_break_before(doc.paragraphs[ack_idx])
    if ack_idx is not None and appendix_idx is not None:
        sample_idx = _find_para(
            doc,
            lambda p: p.style and p.style.name == "Normal" and (p.text or "").strip(),
            start=ack_idx + 1,
            end=appendix_idx,
        )
        if sample_idx is not None:
            sample = doc.paragraphs[sample_idx]
            _replace_para_text(sample, "{{ acknowledgement }}")
            _remove_between(doc, sample._p, doc.paragraphs[appendix_idx]._p, {sample._p})

    appendix_idx = _find_text(doc, "附录", start=(ack_idx or 0) + 1)
    if appendix_idx is None:
        return
    _page_break_before(doc.paragraphs[appendix_idx])
    appendix_elem = doc.paragraphs[appendix_idx]._p
    sample_idx = _find_para(
        doc,
        lambda p: p.style and p.style.name == "Normal" and (p.text or "").strip(),
        start=appendix_idx + 1,
    )
    if sample_idx is None:
        return
    sample = doc.paragraphs[sample_idx]
    _replace_para_text(sample, "{{ appendix }}")
    _replace_between(doc, appendix_elem, sample._p, [])
    _remove_after(doc, sample._p, {sample._p})


def _strip_revision_markup(docx_path: Path) -> None:
    unwrap = {
        f"{{{WNS}}}ins",
        f"{{{WNS}}}moveTo",
    }
    drop = {
        f"{{{WNS}}}del",
        f"{{{WNS}}}moveFrom",
        f"{{{WNS}}}commentRangeStart",
        f"{{{WNS}}}commentRangeEnd",
        f"{{{WNS}}}commentReference",
        f"{{{WNS}}}proofErr",
    }

    tmp = docx_path.with_suffix(docx_path.suffix + ".tmp")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename.endswith(".xml"):
                try:
                    root = etree.fromstring(data)
                except Exception:
                    pass
                else:
                    etree.strip_tags(root, *unwrap)
                    etree.strip_elements(root, *drop, with_tail=False)
                    data = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
            zout.writestr(info, data)
    tmp.replace(docx_path)


def make(src_path: str | os.PathLike[str], out_path: str | os.PathLike[str]) -> None:
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document(src_path)
    setup_cover(doc)
    remove_redundant_cover_page_break(doc)
    setup_toc(doc)
    setup_abstracts(doc)
    setup_body(doc)
    setup_references(doc)
    setup_ack_and_appendix(doc)
    doc.save(out)
    _strip_revision_markup(out)
    print(f"完成: {out}")


def main() -> int:
    if len(sys.argv) < 2:
        print("用法: python make.py <source.docx> [template.docx]")
        return 1
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "template.docx"
    make(src, out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
