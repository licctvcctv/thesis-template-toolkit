from __future__ import annotations

import csv
import json
import re
import subprocess
import tempfile
from copy import deepcopy
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from docxtpl import DocxTemplate


PAPER_DIR = Path(__file__).resolve().parent
PROJECT_DIR = PAPER_DIR.parents[1]
TEMPLATE_DOCX = PROJECT_DIR / "templates" / "uav_lightweight_detection" / "template.docx"
CONTENT_DIR = PAPER_DIR / "content"
META_PATH = CONTENT_DIR / "meta.json"
CHAPTERS_PATH = CONTENT_DIR / "chapters.json"
FIGURES_PATH = CONTENT_DIR / "figures.json"
TABLES_PATH = CONTENT_DIR / "tables.json"
METRICS_PATH = CONTENT_DIR / "metrics.json"
REFERENCES_PATH = CONTENT_DIR / "references.json"

FONT_CN = "宋体"
FONT_HEI = "黑体"
FONT_LATIN = "Times New Roman"


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_asset(value: str, generated_assets: dict[str, Path] | None = None) -> Path:
    if value.startswith("generated:"):
        if generated_assets is None:
            raise KeyError(f"Generated asset requested before creation: {value}")
        return generated_assets[value.split(":", 1)[1]]
    path = Path(value)
    return path if path.is_absolute() else PAPER_DIR / path


def set_run_font(run, size_pt: float | None = None, cn_font: str = FONT_CN, latin_font: str = FONT_LATIN, bold=None):
    run.font.name = latin_font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), cn_font)
    r_fonts.set(qn("w:ascii"), latin_font)
    r_fonts.set(qn("w:hAnsi"), latin_font)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def set_paragraph_style(paragraph, *, size_pt=11.5, first_line=True, line_spacing=1.35):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = line_spacing
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    for run in paragraph.runs:
        set_run_font(run, size_pt=size_pt)


def add_text_with_citations(paragraph, text: str, size_pt=11.5):
    parts = re.split(r"(\[\[\d+\]\])", text)
    for part in parts:
        if not part:
            continue
        match = re.fullmatch(r"\[\[(\d+)\]\]", part)
        if match:
            run = paragraph.add_run(f"[{match.group(1)}]")
            set_run_font(run, size_pt=size_pt)
            run.font.superscript = True
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size_pt=size_pt)


def add_body_paragraph(doc, text: str):
    paragraph = doc.add_paragraph()
    add_text_with_citations(paragraph, text)
    set_paragraph_style(paragraph)
    return paragraph


def add_heading(doc, text: str, level: int):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.space_before = Pt(6 if level == 1 else 3)
    paragraph.paragraph_format.space_after = Pt(6 if level == 1 else 3)
    paragraph.paragraph_format.line_spacing = 1.25
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.page_break_before = True
        size = 16
    else:
        size = 12
    run = paragraph.add_run(text)
    set_run_font(run, size_pt=size, cn_font=FONT_HEI, bold=True)
    return paragraph


def add_caption(doc, caption: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(caption)
    set_run_font(run, size_pt=10.5)
    return paragraph


def add_figure(doc, figure: dict, generated_assets: dict[str, Path]):
    image_path = resolve_asset(figure["image"], generated_assets)
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(figure.get("width_in", 6.0)))
    add_caption(doc, figure["caption"])


def add_table(doc, table_spec: dict):
    headers = table_spec["headers"]
    rows = table_spec["rows"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for col, value in enumerate(headers):
        table.cell(0, col).text = str(value)
    for row_data in rows:
        cells = table.add_row().cells
        for col, value in enumerate(row_data):
            cells[col].text = str(value)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size_pt=9.5)
    add_caption(doc, table_spec["caption"])


def hide_table_borders(table):
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.first_child_found_in("w:tblBorders")
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)


def set_fixed_table_widths(table, widths):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(1, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width.twips))
        tbl_grid.append(grid_col)

    for cell, width in zip(table.rows[0].cells, widths):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.first_child_found_in("w:tcW")
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.append(tc_w)
        tc_w.set(qn("w:type"), "dxa")
        tc_w.set(qn("w:w"), str(width.twips))


def converted_formula_math(latex: str):
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        md = tmp_path / "formula.md"
        out = tmp_path / "formula.docx"
        md.write_text(f"$$\n{latex}\n$$\n", encoding="utf-8")
        subprocess.run(
            ["pandoc", str(md), "-f", "markdown+tex_math_dollars", "-t", "docx", "-o", str(out)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        doc = Document(out)
        for paragraph in doc.paragraphs:
            xml = paragraph._p.xml
            if "m:oMath" in xml or "m:oMathPara" in xml:
                math_elements = paragraph._p.xpath(".//m:oMath")
                if math_elements:
                    return deepcopy(math_elements[0])
    raise ValueError(f"Formula conversion did not produce Word math: {latex}")


def add_formula(doc, block: dict):
    new_math = converted_formula_math(block["latex"])
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(7.9), alignment=WD_TAB_ALIGNMENT.CENTER)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(15.8), alignment=WD_TAB_ALIGNMENT.RIGHT)

    leading_tab = paragraph.add_run()
    leading_tab.add_tab()
    paragraph._p.append(new_math)
    number_run = paragraph.add_run()
    number_run.add_tab()
    number_run.add_text(block["number"])
    set_run_font(number_run, size_pt=10.5)


def configure_plot_fonts():
    plt.rcParams["font.sans-serif"] = ["PingFang SC", "Arial Unicode MS", "Heiti TC", "SimHei"]
    plt.rcParams["axes.unicode_minus"] = False


def save_plot(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    plt.tight_layout()
    plt.savefig(path, dpi=220, bbox_inches="tight")
    plt.close()


def generate_metrics_bar(metrics: dict, config: dict) -> Path:
    configure_plot_fonts()
    out = resolve_asset(config["path"])
    rows = metrics["validation_best_map50"]
    labels = ["mAP@50", "mAP@50-95", "精确率", "召回率"]
    keys = ["map50", "map50_95", "precision", "recall"]
    models = ["YOLOv8s基准", "Ghost轻量化", "Ghost+CBAM", "完整模型"]
    x = range(len(labels))
    width = 0.18
    plt.figure(figsize=(10.5, 5.8))
    for idx, row in enumerate(rows):
        vals = [float(row[k]) for k in keys]
        xs = [v + (idx - 1.5) * width for v in x]
        plt.bar(xs, vals, width=width, label=models[idx])
        for px, val in zip(xs, vals):
            plt.text(px, val + 0.008, f"{val:.4f}", ha="center", va="bottom", fontsize=7)
    plt.xticks(list(x), labels)
    plt.ylim(0, 0.62)
    plt.ylabel("指标值")
    plt.title(config.get("title", "性能指标对比"))
    plt.grid(axis="y", alpha=0.25)
    plt.legend(fontsize=8)
    save_plot(out)
    return out


def read_training_logs(logs: list[dict]):
    loaded = []
    for item in logs:
        rows = []
        with Path(item["path"]).open(encoding="utf-8") as f:
            for row in csv.DictReader(f):
                rows.append(row)
        loaded.append({"label": item["label"], "rows": rows})
    return loaded


def series(rows: list[dict], key: str):
    points = []
    for row in rows:
        raw = row.get(key, "").strip()
        epoch = row.get("epoch", "").strip()
        if not raw or not epoch:
            continue
        points.append((int(float(epoch)), float(raw)))
    return points


def plot_series(logs, key, title):
    for item in logs:
        pts = series(item["rows"], key)
        if not pts:
            continue
        xs, ys = zip(*pts)
        plt.plot(xs, ys, linewidth=1.4, label=item["label"])
    plt.title(title)
    plt.xlabel("训练轮次")
    plt.ylabel("损失值" if "loss" in key else "指标值")
    plt.grid(alpha=0.25)


def generate_loss_curves(figures: dict, config: dict) -> Path:
    configure_plot_fonts()
    out = resolve_asset(config["path"])
    logs = read_training_logs(figures["training_logs"])
    specs = [
        ("train/box_loss", "训练集边界框损失"),
        ("train/cls_loss", "训练集分类损失"),
        ("train/dfl_loss", "训练集DFL损失"),
        ("val/box_loss", "验证集边界框损失"),
        ("val/cls_loss", "验证集分类损失"),
        ("val/dfl_loss", "验证集DFL损失"),
    ]
    plt.figure(figsize=(12, 7))
    for i, (key, title) in enumerate(specs, start=1):
        plt.subplot(2, 3, i)
        plot_series(logs, key, title)
        if i == 1:
            plt.legend(fontsize=7)
    plt.suptitle(config.get("title", "损失曲线对比"), fontsize=14, y=1.02)
    save_plot(out)
    return out


def generate_map_curves(figures: dict, config: dict) -> Path:
    configure_plot_fonts()
    out = resolve_asset(config["path"])
    logs = read_training_logs(figures["training_logs"])
    specs = [("metrics/mAP50(B)", "mAP@50曲线"), ("metrics/mAP50-95(B)", "mAP@50-95曲线")]
    plt.figure(figsize=(10.5, 4.6))
    for i, (key, title) in enumerate(specs, start=1):
        plt.subplot(1, 2, i)
        plot_series(logs, key, title)
        if i == 1:
            plt.legend(fontsize=8)
    plt.suptitle(config.get("title", "mAP曲线对比"), fontsize=14, y=1.03)
    save_plot(out)
    return out


def generate_class_ap(config: dict) -> Path:
    configure_plot_fonts()
    out = resolve_asset(config["path"])
    csv_path = Path(config["evaluation_csv"])
    with csv_path.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    by_model = {row["model"]: row for row in rows}
    baseline = by_model["YOLOv8s"]
    improved = by_model["YOLOv8s-Improved"]
    classes = [
        ("pedestrian", "行人"),
        ("people", "人群"),
        ("bicycle", "自行车"),
        ("car", "轿车"),
        ("van", "面包车"),
        ("truck", "卡车"),
        ("tricycle", "三轮车"),
        ("awning-tricycle", "遮篷三轮"),
        ("bus", "公交车"),
        ("motor", "摩托车"),
    ]
    labels = [label for _, label in classes]
    base_vals = [float(baseline[f"AP50_{name}"]) for name, _ in classes]
    improved_vals = [float(improved[f"AP50_{name}"]) for name, _ in classes]
    x = range(len(labels))
    width = 0.38
    plt.figure(figsize=(11, 5.2))
    plt.bar([i - width / 2 for i in x], base_vals, width=width, label="YOLOv8s基准")
    plt.bar([i + width / 2 for i in x], improved_vals, width=width, label="完整模型")
    plt.xticks(list(x), labels, rotation=25)
    plt.ylabel("AP50")
    plt.ylim(0, max(base_vals + improved_vals) + 0.08)
    plt.title(config.get("title", "各类别AP50对比"))
    plt.grid(axis="y", alpha=0.25)
    plt.legend(fontsize=8)
    save_plot(out)
    return out


def generate_small_object_ap(config: dict) -> Path:
    configure_plot_fonts()
    out = resolve_asset(config["path"])
    csv_path = Path(config["evaluation_csv"])
    with csv_path.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    models = [row["model"] for row in rows]
    display_models = ["基准模型", "Ghost轻量化", "Ghost-CBAM", "完整模型"]
    classes = [
        ("pedestrian", "行人"),
        ("people", "人群"),
        ("bicycle", "自行车"),
        ("tricycle", "三轮车"),
        ("awning-tricycle", "遮篷三轮"),
        ("motor", "摩托车"),
    ]
    labels = [label for _, label in classes]
    x = range(len(labels))
    width = 0.18
    colors = ["#3B82F6", "#22A06B", "#F59E0B", "#EF4444"]
    plt.figure(figsize=(11.2, 5.2))
    for idx, (model_name, model_label) in enumerate(zip(models, display_models)):
        row = next(r for r in rows if r["model"] == model_name)
        values = [float(row[f"AP50_{name}"]) for name, _ in classes]
        xs = [i + (idx - 1.5) * width for i in x]
        plt.bar(xs, values, width=width, color=colors[idx], label=model_label)
    plt.xticks(list(x), labels)
    plt.ylabel("AP50")
    plt.ylim(0, 0.35)
    plt.title(config.get("title", "小目标类别AP50对比"))
    plt.grid(axis="y", alpha=0.25)
    plt.legend(fontsize=8, ncol=4)
    save_plot(out)
    return out


def generate_fps_bar(metrics: dict, config: dict) -> Path:
    configure_plot_fonts()
    out = resolve_asset(config["path"])
    rows = metrics["test_evaluation"]
    models = ["基准模型", "Ghost轻量化", "Ghost-CBAM", "完整模型"]
    fps = [float(row[5]) for row in rows]
    colors = ["#3B82F6", "#22A06B", "#F59E0B", "#EF4444"]
    plt.figure(figsize=(8, 4.8))
    bars = plt.bar(models, fps, color=colors)
    for bar, value in zip(bars, fps):
        plt.text(bar.get_x() + bar.get_width() / 2, value + 3, f"{value:.1f}", ha="center", fontsize=9)
    plt.ylabel("FPS")
    plt.title(config.get("title", "推理速度对比"))
    plt.ylim(0, max(fps) + 45)
    plt.grid(axis="y", alpha=0.25)
    save_plot(out)
    return out


def generate_module_delta_chart(metrics: dict, config: dict) -> Path:
    configure_plot_fonts()
    out = resolve_asset(config["path"])
    rows = metrics["validation_best_map50"]
    pairs = [
        (rows[0], rows[1], "Ghost替换"),
        (rows[1], rows[2], "加入CBAM"),
        (rows[2], rows[3], "加入Wise-IoU"),
    ]
    complexity_keys = [("params_m", "参数量"), ("gflops", "GFLOPs")]
    metric_keys = [("precision", "Precision"), ("recall", "Recall"), ("map50", "mAP@50"), ("map50_95", "mAP@50-95")]
    labels = [label for _, _, label in pairs]

    plt.figure(figsize=(11.2, 5.2))
    ax1 = plt.subplot(1, 2, 1)
    width = 0.34
    x = range(len(labels))
    for idx, (key, name) in enumerate(complexity_keys):
        values = []
        for before, after, _ in pairs:
            b = float(before[key])
            a = float(after[key])
            values.append((a - b) / b * 100 if b else 0)
        xs = [i + (idx - 0.5) * width for i in x]
        bars = ax1.bar(xs, values, width=width, label=name)
        for bar, value in zip(bars, values):
            ax1.text(bar.get_x() + bar.get_width() / 2, value + (0.6 if value >= 0 else 0.8), f"{value:+.2f}%", ha="center", fontsize=8)
    ax1.axhline(0, color="#333333", linewidth=0.8)
    ax1.set_xticks(list(x), labels)
    ax1.set_ylabel("变化率")
    ax1.set_title("复杂度变化")
    ax1.grid(axis="y", alpha=0.25)
    ax1.legend(fontsize=8)

    ax2 = plt.subplot(1, 2, 2)
    width = 0.18
    for idx, (key, name) in enumerate(metric_keys):
        values = [float(after[key]) - float(before[key]) for before, after, _ in pairs]
        xs = [i + (idx - 1.5) * width for i in x]
        bars = ax2.bar(xs, values, width=width, label=name)
        for bar, value in zip(bars, values):
            ax2.text(bar.get_x() + bar.get_width() / 2, value + (0.0012 if value >= 0 else -0.004), f"{value:+.4f}", ha="center", fontsize=7, rotation=90)
    ax2.axhline(0, color="#333333", linewidth=0.8)
    ax2.set_xticks(list(x), labels)
    ax2.set_ylabel("指标变化量")
    ax2.set_title("精度指标变化")
    ax2.grid(axis="y", alpha=0.25)
    ax2.legend(fontsize=7, ncol=2)
    plt.suptitle(config.get("title", "模块改进指标变化对比"), fontsize=14, y=1.03)
    save_plot(out)
    return out


def generate_final_loss_bar(figures: dict, config: dict) -> Path:
    configure_plot_fonts()
    out = resolve_asset(config["path"])
    logs = read_training_logs(figures["training_logs"])
    specs = [("val/box_loss", "Box Loss"), ("val/cls_loss", "Cls Loss"), ("val/dfl_loss", "DFL Loss")]
    labels = [item["label"] for item in logs]
    x = range(len(labels))
    width = 0.24
    plt.figure(figsize=(10.5, 4.8))
    for idx, (key, name) in enumerate(specs):
        values = []
        for item in logs:
            pts = series(item["rows"], key)
            values.append(pts[-1][1] if pts else 0)
        xs = [i + (idx - 1) * width for i in x]
        bars = plt.bar(xs, values, width=width, label=name)
        for bar, value in zip(bars, values):
            plt.text(bar.get_x() + bar.get_width() / 2, value + 0.02, f"{value:.3f}", ha="center", fontsize=8)
    plt.xticks(list(x), labels)
    plt.ylabel("末轮验证损失")
    plt.title(config.get("title", "四类模型末轮验证损失对比"))
    plt.grid(axis="y", alpha=0.25)
    plt.legend(fontsize=8)
    save_plot(out)
    return out


def generate_test_metrics_chart(metrics: dict, config: dict) -> Path:
    configure_plot_fonts()
    out = resolve_asset(config["path"])
    rows = metrics["test_evaluation"]
    labels = ["基准模型", "Ghost轻量化", "Ghost-CBAM", "完整模型"]
    map50 = [float(row[3]) for row in rows]
    map95 = [float(row[4]) for row in rows]
    fps = [float(row[5]) for row in rows]
    params = [float(row[1]) for row in rows]
    gflops = [float(row[2]) for row in rows]
    x = range(len(labels))

    plt.figure(figsize=(11.2, 5.2))
    ax1 = plt.subplot(1, 2, 1)
    width = 0.34
    bars1 = ax1.bar([i - width / 2 for i in x], map50, width=width, label="mAP@50")
    bars2 = ax1.bar([i + width / 2 for i in x], map95, width=width, label="mAP@50-95")
    for bars in (bars1, bars2):
        for bar in bars:
            value = bar.get_height()
            ax1.text(bar.get_x() + bar.get_width() / 2, value + 0.006, f"{value:.4f}", ha="center", fontsize=8, rotation=90)
    ax1.set_xticks(list(x), labels)
    ax1.set_ylim(0, max(map50) + 0.08)
    ax1.set_ylabel("测试集mAP")
    ax1.set_title("测试集精度")
    ax1.grid(axis="y", alpha=0.25)
    ax1.legend(fontsize=8)

    ax2 = plt.subplot(1, 2, 2)
    bars = ax2.bar([i - width / 2 for i in x], params, width=width, label="参数量(M)", color="#60A5FA")
    bars_g = ax2.bar([i + width / 2 for i in x], gflops, width=width, label="GFLOPs", color="#34D399")
    ax2.set_xticks(list(x), labels)
    ax2.set_ylabel("复杂度")
    ax2.grid(axis="y", alpha=0.25)
    ax2b = ax2.twinx()
    ax2b.plot(list(x), fps, color="#EF4444", marker="o", linewidth=2.0, label="FPS")
    ax2b.set_ylabel("FPS")
    ax2b.set_ylim(0, max(fps) + 45)
    for i, value in enumerate(fps):
        ax2b.text(i, value + 6, f"{value:.1f}", ha="center", fontsize=8, color="#B91C1C")
    handles = [bars, bars_g, ax2b.lines[0]]
    labels_legend = [h.get_label() for h in handles]
    ax2.legend(handles, labels_legend, fontsize=8, loc="upper right")
    ax2.set_title("复杂度与速度")
    plt.suptitle(config.get("title", "测试集综合评价图"), fontsize=14, y=1.03)
    save_plot(out)
    return out


def generate_charts(figures: dict, metrics: dict) -> dict[str, Path]:
    generated = {}
    for key, config in figures.get("generated_charts", {}).items():
        chart_type = config["type"]
        if chart_type == "metrics_bar":
            generated[key] = generate_metrics_bar(metrics, config)
        elif chart_type == "training_loss_curves":
            generated[key] = generate_loss_curves(figures, config)
        elif chart_type == "map_curves":
            generated[key] = generate_map_curves(figures, config)
        elif chart_type == "class_ap_bar":
            generated[key] = generate_class_ap(config)
        elif chart_type == "small_object_ap_bar":
            generated[key] = generate_small_object_ap(config)
        elif chart_type == "fps_bar":
            generated[key] = generate_fps_bar(metrics, config)
        elif chart_type == "module_delta_chart":
            generated[key] = generate_module_delta_chart(metrics, config)
        elif chart_type == "final_loss_bar":
            generated[key] = generate_final_loss_bar(figures, config)
        elif chart_type == "test_metrics_chart":
            generated[key] = generate_test_metrics_chart(metrics, config)
        else:
            raise ValueError(f"Unsupported chart type: {chart_type}")
    return generated


def marker_numbers(chapters: dict) -> list[int]:
    numbers = []
    for chapter in chapters["chapters"]:
        for block in chapter["blocks"]:
            if block.get("type") == "paragraph":
                numbers.extend(int(n) for n in re.findall(r"\[\[(\d+)\]\]", block["text"]))
    return numbers


def validate_citations(chapters: dict, references: list[str]):
    numbers = marker_numbers(chapters)
    expected = list(range(1, len(references) + 1))
    if sorted(set(numbers)) != expected:
        raise ValueError(f"Citation markers must cover references {expected}; found {sorted(set(numbers))}")
    first_seen = []
    for number in numbers:
        if number not in first_seen:
            first_seen.append(number)
    if first_seen != expected:
        raise ValueError(f"First citation order must be ascending {expected}; found {first_seen}")


def render_chapters(subdoc, chapters: dict, figures: dict, tables: dict, generated_assets: dict[str, Path]):
    figure_map = figures["figures"]
    table_map = tables["tables"]
    for chapter in chapters["chapters"]:
        add_heading(subdoc.docx, chapter["title"], 1)
        for block in chapter["blocks"]:
            block_type = block["type"]
            if block_type == "heading":
                add_heading(subdoc.docx, block["text"], block["level"])
            elif block_type == "paragraph":
                add_body_paragraph(subdoc.docx, block["text"])
            elif block_type == "figure":
                add_figure(subdoc.docx, figure_map[block["id"]], generated_assets)
            elif block_type == "table":
                add_table(subdoc.docx, table_map[block["id"]])
            elif block_type == "formula":
                add_formula(subdoc.docx, block)
            else:
                raise ValueError(f"Unsupported block type: {block_type}")


def add_references(doc, references: list[str]):
    add_heading(doc, "参考文献", 1)
    for idx, ref in enumerate(references, start=1):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.first_line_indent = Cm(-0.74)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.74))
        label = paragraph.add_run(f"[{idx}]")
        label.add_tab()
        body = paragraph.add_run(ref)
        set_run_font(label, size_pt=10.5)
        set_run_font(body, size_pt=10.5)


def add_acknowledgement(doc, acknowledgement: str):
    add_heading(doc, "致谢", 1)
    add_body_paragraph(doc, acknowledgement)


def main():
    meta = load_json(META_PATH)
    chapters = load_json(CHAPTERS_PATH)
    figures = load_json(FIGURES_PATH)
    tables = load_json(TABLES_PATH)
    metrics = load_json(METRICS_PATH)
    references = load_json(REFERENCES_PATH)
    validate_citations(chapters, references)

    generated_assets = generate_charts(figures, metrics)
    tpl = DocxTemplate(TEMPLATE_DOCX)
    chapters_subdoc = tpl.new_subdoc()
    render_chapters(chapters_subdoc, chapters, figures, tables, generated_assets)
    add_references(chapters_subdoc.docx, references)
    add_acknowledgement(chapters_subdoc.docx, meta["acknowledgement"])

    output_docx = PAPER_DIR / meta["output_filename"]
    tpl.render({"title_zh": meta["title_zh"], "chapters_subdoc": chapters_subdoc})
    tpl.save(output_docx)
    print(output_docx)


if __name__ == "__main__":
    main()
