#!/usr/bin/env python3
"""Build the thesis DOCX from JSON content and experiment outputs."""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


PAPER_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = PAPER_DIR.parents[1]
TEMPLATE_PATH = PROJECT_ROOT / "templates" / "smbu" / "source.docx"
CONTENT_DIR = PAPER_DIR / "content"
DATA_DIR = PAPER_DIR / "data"
OUTPUT_NAME = "基于数据库过滤的谱图聚类算法设计_论文初稿.docx"
OUTPUT_PATH = PAPER_DIR / OUTPUT_NAME


def load_json(name: str):
    return json.loads((CONTENT_DIR / name).read_text(encoding="utf-8"))


def set_run_font(run, size=12, bold=False, italic=False, font_cn="宋体", font_en="Times New Roman"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_cn)
    r_fonts.set(qn("w:ascii"), font_en)
    r_fonts.set(qn("w:hAnsi"), font_en)


def set_paragraph_format(paragraph, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5, before=0, after=0):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if first_line:
        fmt.first_line_indent = Cm(0.74)
    else:
        fmt.first_line_indent = None


def set_cell_text(cell, text, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    set_paragraph_format(p, first_line=False, align=align, line=1.2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def set_cell_shading(cell, fill="EAF2F8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size, bold in [("Heading 1", 16, True), ("Heading 2", 12, True), ("Heading 3", 12, True)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)


def add_page_number_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        for child in list(footer._element):
            footer._element.remove(child)
        p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prefix = p.add_run("第 ")
        set_run_font(prefix, size=10.5)
        run = p.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)
        set_run_font(run, size=10.5)
        suffix = p.add_run(" 页")
        set_run_font(suffix, size=10.5)


def add_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_outline_level(paragraph, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def clear_outline_level(paragraph) -> None:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is not None:
        p_pr.remove(outline)


def add_tc_field(paragraph, title: str, level: int = 1, identifier: str = "C") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f'TC "{title}" \\f {identifier} \\l "{level}"'
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_centered_title(doc, text, size=18, bold=True, before=0, after=8, font_cn="黑体"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, font_cn=font_cn)
    return p


def add_blank_paragraphs(doc, count=1):
    for _ in range(count):
        doc.add_paragraph()


def add_body_para(doc, text, size=12):
    p = doc.add_paragraph()
    set_paragraph_format(p)
    add_text_with_citations(p, text, size=size)
    return p


def add_text_with_citations(paragraph, text, size=12):
    parts = re.split(r"(\[\[\d+\]\])", text)
    for part in parts:
        if not part:
            continue
        match = re.fullmatch(r"\[\[(\d+)\]\]", part)
        if match:
            run = paragraph.add_run(f"[{match.group(1)}]")
            run.font.superscript = True
            set_run_font(run, size=size - 1)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size)


def paragraph_text_from_xml(paragraph_element) -> str:
    return "".join(node.text or "" for node in paragraph_element.iter(qn("w:t")))


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def set_template_paragraph(paragraph, text, size=12, bold=False, align=None, font_cn="宋体", font_en="Times New Roman"):
    clear_paragraph(paragraph)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, font_cn=font_cn, font_en=font_en)


def insert_paragraph_after(paragraph, text, size=12, bold=False, align=None, font_cn="宋体", font_en="Times New Roman"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    set_template_paragraph(inserted, text, size=size, bold=bold, align=align, font_cn=font_cn, font_en=font_en)
    return inserted


def set_template_paragraphs(paragraph, texts, size=12, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    set_template_paragraph(paragraph, texts[0], size=size, align=align)
    set_paragraph_format(paragraph, first_line=first_line, align=align)
    current = paragraph
    for text in texts[1:]:
        current = insert_paragraph_after(current, text, size=size, align=align)
        set_paragraph_format(current, first_line=first_line, align=align)


def replace_run_text(paragraph, old, new):
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)


def replace_text_across_runs(paragraph, old, new):
    full_text = "".join(run.text for run in paragraph.runs)
    if old not in full_text or not paragraph.runs:
        return False
    paragraph.runs[0].text = full_text.replace(old, new)
    for run in paragraph.runs[1:]:
        run.text = ""
    return True


def set_paragraph_text_preserving_format(paragraph, text):
    """Replace visible text while leaving the template paragraph formatting alone."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def clear_paragraph_runs(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def clone_run_format(source_run, target_run):
    r_pr = source_run._r.rPr
    if r_pr is not None:
        target_run._r.insert(0, deepcopy(r_pr))


def insert_paragraph_after_like(paragraph, text):
    new_p = OxmlElement("w:p")
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        new_p.append(deepcopy(p_pr))
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    run = inserted.add_run(text)
    if paragraph.runs:
        clone_run_format(paragraph.runs[0], run)
    return inserted


def set_template_paragraphs_preserving(paragraph, texts):
    set_paragraph_text_preserving_format(paragraph, texts[0])
    current = paragraph
    for text in texts[1:]:
        current = insert_paragraph_after_like(current, text)


def set_keywords_paragraph_preserving(paragraph, label, values, sep):
    clear_paragraph_runs(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if label.startswith("【关键词"):
        label_run = paragraph.add_run("【关键词】")
        set_run_font(label_run, size=12, bold=False, font_cn="黑体", font_en="黑体")
        colon_run = paragraph.add_run("：")
        set_run_font(colon_run, size=12, bold=True, font_cn="宋体", font_en="Times New Roman")
        value_run = paragraph.add_run(sep.join(values))
        set_run_font(value_run, size=12, bold=False, font_cn="宋体", font_en="Times New Roman")
    elif label.startswith("【Key Words"):
        label_run = paragraph.add_run("【Key Words】")
        set_run_font(label_run, size=12, bold=True, font_cn="Times New Roman", font_en="Times New Roman")
        colon_run = paragraph.add_run(": ")
        set_run_font(colon_run, size=12, bold=False, font_cn="Times New Roman", font_en="Times New Roman")
        value_run = paragraph.add_run(sep.join(values))
        set_run_font(value_run, size=12, bold=False, font_cn="Times New Roman", font_en="Times New Roman")
    elif label.startswith("【Ключевые"):
        label_run = paragraph.add_run("【Ключевые слова】: ")
        set_run_font(label_run, size=12, bold=True, font_cn="Times New Roman", font_en="Times New Roman")
        value_run = paragraph.add_run(sep.join(values))
        set_run_font(value_run, size=12, bold=False, font_cn="Times New Roman", font_en="Times New Roman")
    else:
        set_paragraph_text_preserving_format(paragraph, f"{label}{sep.join(values)}")


def set_cell_text_preserving(cell, text):
    if not cell.paragraphs:
        cell.add_paragraph()
    set_paragraph_text_preserving_format(cell.paragraphs[0], text)
    for paragraph in cell.paragraphs[1:]:
        set_paragraph_text_preserving_format(paragraph, "")


def set_cover_value_cell(cell, text, style_run=None):
    if not cell.paragraphs:
        cell.add_paragraph()
    paragraph = cell.paragraphs[0]
    clear_paragraph_runs(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    if style_run is not None:
        clone_run_format(style_run, run)
    else:
        set_run_font(run, size=16, bold=True, font_cn="楷体")
    r_pr = run._r.get_or_add_rPr()
    color = r_pr.find(qn("w:color"))
    if color is not None:
        r_pr.remove(color)
    for paragraph in cell.paragraphs[1:]:
        clear_paragraph_runs(paragraph)


def set_keywords_paragraph(paragraph, label, values, sep):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(paragraph, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=12, bold=True, font_cn="黑体")
    value_run = paragraph.add_run(sep.join(values))
    set_run_font(value_run, size=12)


def remove_body_from_template_anchor(doc: Document, anchor_text: str = "论文标题") -> None:
    body = doc._element.body
    children = list(body)
    start = None
    for idx, child in enumerate(children):
        if child.tag == qn("w:p") and paragraph_text_from_xml(child).strip() == anchor_text:
            start = idx
            break
    if start is None:
        raise RuntimeError(f"Template body anchor not found: {anchor_text}")
    for child in children[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def trim_template_stale_toc_and_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        text = paragraph_text_from_xml(child)
        if child.tag == qn("w:sdt") and ("XXXX" in text or "Оглавление" in text or "目    录" in text):
            body.remove(child)
    children = list(body)
    keyword_idx = None
    body_idx = None
    for idx, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = paragraph_text_from_xml(child).strip()
        if text.startswith("【Ключевые слова】"):
            keyword_idx = idx
        if text == "论文标题":
            body_idx = idx
            break
    if body_idx is None:
        raise RuntimeError("Template body anchor not found: 论文标题")
    # The SMBU template carries a cached Russian/placeholder TOC between the
    # abstract block and the body sample. Remove that cache and the sample body;
    # a fresh TOC and generated thesis body are inserted afterward.
    start = (keyword_idx + 1) if keyword_idx is not None else body_idx
    for child in children[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def tighten_cover_spacing(doc: Document, title: str) -> None:
    body = doc._element.body
    children = list(body)
    title_idx = None
    table_idx = None
    for idx, child in enumerate(children):
        if child.tag == qn("w:p") and paragraph_text_from_xml(child).strip() == title:
            title_idx = idx
        elif title_idx is not None and child.tag == qn("w:tbl"):
            table_idx = idx
            break
    if title_idx is None or table_idx is None:
        return
    blank_paragraphs = [
        child for child in children[title_idx + 1 : table_idx]
        if child.tag == qn("w:p") and not paragraph_text_from_xml(child).strip()
    ]
    for child in blank_paragraphs[1:]:
        body.remove(child)


def remove_blank_paragraphs_between(doc: Document, start_text: str, end_text: str, keep: int = 0) -> None:
    body = doc._element.body
    children = list(body)
    start_idx = None
    end_idx = None
    for idx, child in enumerate(children):
        text = paragraph_text_from_xml(child).strip() if child.tag == qn("w:p") else ""
        if start_idx is None and text.startswith(start_text):
            start_idx = idx
            continue
        if start_idx is not None and text.startswith(end_text):
            end_idx = idx
            break
    if start_idx is None or end_idx is None:
        return
    blank_children = [
        child for child in children[start_idx + 1 : end_idx]
        if child.tag == qn("w:p")
        and not paragraph_text_from_xml(child).strip()
        and not list(child.iter(qn("w:drawing")))
    ]
    for child in blank_children[keep:]:
        body.remove(child)


def tighten_russian_cover_spacing(doc: Document) -> None:
    remove_blank_paragraphs_between(doc, "Факультет", "Направление подготовки", keep=1)
    remove_blank_paragraphs_between(doc, "«Проектирование", "Работу выполнил:", keep=1)
    remove_blank_paragraphs_between(doc, "Научный руководитель:", "Шэньчжэнь", keep=1)


def tighten_chinese_declaration_spacing(doc: Document) -> None:
    remove_blank_paragraphs_between(doc, "本人郑重声明", "毕业论文（设计）作者签名", keep=0)
    remove_blank_paragraphs_between(doc, "日期：", "学位论文使用授权书", keep=0)
    remove_blank_paragraphs_between(doc, "学校可以向国家有关部门", "毕业论文（设计）作者签名", keep=0)


def tighten_russian_declaration_spacing(doc: Document) -> None:
    remove_blank_paragraphs_between(doc, "Подпись автора выпускной работы", "Доверенность", keep=0)
    remove_blank_paragraphs_between(doc, "ВУЗ может отправлять", "Подпись автора выпускной работы", keep=0)


def prepare_template_front_matter(doc: Document, meta) -> None:
    title_ru = meta.get("title_ru", "Проектирование алгоритма кластеризации масс-спектров на основе фильтрации базы данных")
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        normalized_title = re.sub(r"\s+", "", text).lower()
        clear_outline_level(paragraph)
        if normalized_title == "摘要":
            add_tc_field(paragraph, "摘 要", 1)
        elif normalized_title == "abstract":
            add_tc_field(paragraph, "Abstract", 1)
        elif normalized_title == "аннотация":
            add_tc_field(paragraph, "Аннотация", 1)

        if text == "论文标题中俄文":
            set_paragraph_text_preserving_format(paragraph, meta["title_zh"])
        elif text == "（二号，楷体，加粗，居中）":
            set_paragraph_text_preserving_format(paragraph, "")
        elif text == "Факультет ХХХ":
            set_paragraph_text_preserving_format(paragraph, f"Факультет {meta.get('college_ru', meta['college'])}")
        elif text.startswith("Направление подготовки"):
            set_paragraph_text_preserving_format(paragraph, f"Направление подготовки {meta.get('major_ru', meta['major'])}")
        elif text == "«Тема ВКР»":
            set_paragraph_text_preserving_format(paragraph, f"«{title_ru}»")
        elif text == "2023":
            set_paragraph_text_preserving_format(paragraph, meta["year"])
        elif text.startswith("这里是中文摘要"):
            set_template_paragraphs_preserving(paragraph, meta["abstract_zh"])
        elif text.startswith("This is abstract"):
            set_template_paragraphs_preserving(paragraph, meta["abstract_en"])
        elif text.startswith("Здесь располагается"):
            set_template_paragraphs_preserving(paragraph, meta["abstract_ru"])
        elif text.startswith("【关键词】"):
            set_keywords_paragraph_preserving(paragraph, "【关键词】：", meta["keywords_zh"], "；")
        elif text.startswith("【Key Words】"):
            set_keywords_paragraph_preserving(paragraph, "【Key Words】: ", meta["keywords_en"], "; ")
        elif text.startswith("【Ключевые слова】"):
            set_keywords_paragraph_preserving(paragraph, "【Ключевые слова】: ", meta["keywords_ru"], "; ")

        replace_run_text(paragraph, "双语", meta["title_zh"])
        replace_run_text(paragraph, "На двух языках", title_ru)
        replace_text_across_runs(paragraph, "双语", meta["title_zh"])
        replace_text_across_runs(paragraph, "На двух языках", title_ru)

    if doc.tables:
        cover_value_style_run = None
        if len(doc.tables[0].rows) > 2:
            sample_paragraphs = doc.tables[0].rows[2].cells[1].paragraphs
            if sample_paragraphs and sample_paragraphs[0].runs:
                cover_value_style_run = sample_paragraphs[0].runs[0]
        cover_rows = [
            ("姓名", meta["name"]),
            ("院系", meta["college"]),
            ("专业", meta["major"]),
            ("学号", meta["student_id"]),
            ("指导教师", meta["advisor"]),
            ("职称", meta.get("advisor_title", "")),
            ("提交日期", f"{meta['year']} 年 {meta['month']} 月"),
        ]
        for row, (_, value) in zip(doc.tables[0].rows, cover_rows):
            set_cover_value_cell(row.cells[1], value, cover_value_style_run)
    trim_template_stale_toc_and_body(doc)


def set_cover_cell_bottom_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for name in ["top", "left", "right", "insideH", "insideV"]:
        element = borders.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            borders.append(element)
        element.set(qn("w:val"), "nil")
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), "000000")


def add_cover(doc, meta):
    logo = PAPER_DIR / "images" / "template" / "image1.png"
    if logo.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        p.add_run().add_picture(str(logo), width=Cm(11.6))
    else:
        add_centered_title(doc, meta["school"], size=16, bold=True, before=4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    line = p.add_run("─" * 52)
    set_run_font(line, size=9)
    add_blank_paragraphs(doc, 5)
    add_centered_title(doc, "本  科  毕  业  论  文（设  计）", size=24, bold=True, before=6, after=58)
    add_centered_title(doc, meta["title_zh"], size=20, bold=True, before=0, after=6, font_cn="楷体")
    add_centered_title(doc, meta["title_en"], size=13, bold=True, after=34, font_cn="Times New Roman")
    table = doc.add_table(rows=7, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("姓名", meta["name"]),
        ("院系", meta["college"]),
        ("专业", meta["major"]),
        ("学号", meta["student_id"]),
        ("指导教师", meta["advisor"]),
        ("职称", meta.get("advisor_title", "")),
        ("提交日期", f"{meta['year']}年{meta['month']}月"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        row.cells[0].width = Cm(4.0)
        row.cells[1].width = Cm(10.5)
        set_cell_text(row.cells[0], "  ".join(label), size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], value, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cover_cell_bottom_border(row.cells[0])
        set_cover_cell_bottom_border(row.cells[1])
        for cell in row.cells:
            set_cell_margins(cell, top=130, bottom=130, start=80, end=80)
    remove_table_borders(table)
    for row in table.rows:
        set_cover_cell_bottom_border(row.cells[0])
        set_cover_cell_bottom_border(row.cells[1])
    doc.add_page_break()


def add_declaration(doc, meta):
    add_centered_title(doc, "深圳北理莫斯科大学", size=15, bold=True)
    add_centered_title(doc, "本科毕业论文（设计）诚信声明", size=16, bold=True)
    p = add_body_para(
        doc,
        f"本人郑重声明：所呈交的毕业论文（设计），题目《{meta['title_zh']}》是本人在指导教师的指导下，独立进行研究工作所取得的成果。对本文的研究做出重要贡献的个人和集体，均已在文中以明确方式注明。除此之外，本论文不包含任何其他个人或集体已经发表或撰写过的作品成果。本人完全意识到本声明的法律后果由本人承担。",
    )
    p.paragraph_format.space_after = Pt(24)
    for text in ["毕业论文（设计）作者签名：", "日期：      年    月    日"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(text)
        set_run_font(run)
    doc.add_page_break()

    add_centered_title(doc, "学位论文使用授权书", size=16, bold=True)
    for text in [
        "本人完全了解深圳北理莫斯科大学关于收集、保存、使用学位论文的规定，即：",
        "1. 按照学校要求提交学位论文的印刷本和电子版本；",
        "2. 学校可以采用影印、缩印或其他复制手段保存论文；",
        "3. 学校可以向国家有关部门或机构送交论文电子版，允许论文被查阅、抽检。",
    ]:
        add_body_para(doc, text)
    for _ in range(2):
        doc.add_paragraph()
    for text in ["毕业论文（设计）作者签名：", "指导教师签名：", "日期：      年    月    日"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(text)
        set_run_font(run)
    doc.add_page_break()


def add_abstracts(doc, meta):
    add_centered_title(doc, "摘 要", size=16, bold=True)
    for paragraph in meta["abstract_zh"]:
        add_body_para(doc, paragraph)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    label = p.add_run("【关键词】：")
    set_run_font(label, size=12, bold=True, font_cn="黑体")
    value = p.add_run("；".join(meta["keywords_zh"]))
    set_run_font(value, size=12)
    doc.add_page_break()

    add_centered_title(doc, "Abstract", size=16, bold=True, font_cn="Times New Roman")
    for paragraph in meta["abstract_en"]:
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        run = p.add_run(paragraph)
        set_run_font(run, size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    label = p.add_run("【Key Words】: ")
    set_run_font(label, size=12, bold=True, font_cn="Times New Roman")
    value = p.add_run("; ".join(meta["keywords_en"]))
    set_run_font(value, size=12)
    doc.add_page_break()


def format_toc_title(paragraph):
    p = paragraph
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.first_line_indent = None
    fmt.left_indent = Cm(0)
    fmt.right_indent = Cm(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = 1.5
    run = p.add_run("目    录")
    set_run_font(run, size=18, bold=True, font_cn="黑体", font_en="黑体")
    return p


def add_toc_title(doc):
    doc.add_page_break()
    return format_toc_title(doc.add_paragraph())


def format_toc_entry(paragraph, level):
    try:
        paragraph.style = f"toc {min(level, 3)}"
    except KeyError:
        pass
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = None
    indent_cm = 0.74 * (max(1, min(level, 3)) - 1)
    fmt.left_indent = Cm(indent_cm)
    fmt.right_indent = Cm(0.74 if level > 1 else 0)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.tab_stops.add_tab_stop(Cm(14.64), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def add_sdt_paragraph(doc, sdt_content):
    p = OxmlElement("w:p")
    sdt_content.append(p)
    return Paragraph(p, doc._body)


def insert_body_element(doc, element):
    body = doc._element.body
    sect_pr = body.sectPr
    if sect_pr is None:
        body.append(element)
    else:
        sect_pr.addprevious(element)


def copy_header_footer_refs(src_section, dst_section) -> None:
    """Make a new section carry explicit header/footer refs.

    Word can inherit refs across sections, but LibreOffice maps sections to page
    styles and is more predictable when the TOC section has concrete refs.
    """
    src = src_section._sectPr
    dst = dst_section._sectPr
    for child in list(dst):
        if child.tag in {qn("w:headerReference"), qn("w:footerReference")}:
            dst.remove(child)

    refs = [
        deepcopy(child)
        for child in src
        if child.tag in {qn("w:headerReference"), qn("w:footerReference")}
    ]
    for ref in reversed(refs):
        dst.insert(0, ref)


def materialize_linked_header_footer_refs(doc: Document) -> None:
    """Write inherited header/footer refs into each section explicitly."""
    ref_tags = {qn("w:headerReference"), qn("w:footerReference")}
    inherited_refs = []
    for sect_pr in doc._element.body.iter(qn("w:sectPr")):
        refs = [child for child in sect_pr if child.tag in ref_tags]
        if refs:
            inherited_refs = [deepcopy(child) for child in refs]
            continue
        if not inherited_refs:
            continue
        for ref in reversed(inherited_refs):
            sect_pr.insert(0, deepcopy(ref))


def create_toc_sdt():
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), "Table of Contents")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "TOC")
    doc_part = OxmlElement("w:docPartObj")
    gallery = OxmlElement("w:docPartGallery")
    gallery.set(qn("w:val"), "Table of Contents")
    unique = OxmlElement("w:docPartUnique")
    doc_part.append(gallery)
    doc_part.append(unique)
    sdt_pr.append(alias)
    sdt_pr.append(tag)
    sdt_pr.append(doc_part)
    sdt_content = OxmlElement("w:sdtContent")
    sdt.append(sdt_pr)
    sdt.append(sdt_content)
    return sdt, sdt_content


def add_toc(doc, toc_entries=None):
    # Keep the table of contents in its own Word section. LibreOffice converts
    # DOCX sections into Writer page styles, so a plain page break can make the
    # TOC inherit body-page header/footer state ambiguously in the UI.
    previous_section = doc.sections[-1]
    toc_section = doc.add_section(WD_SECTION.NEW_PAGE)
    copy_header_footer_refs(previous_section, toc_section)
    sdt, sdt_content = create_toc_sdt()
    insert_body_element(doc, sdt)
    format_toc_title(add_sdt_paragraph(doc, sdt_content))

    field_p = add_sdt_paragraph(doc, sdt_content)
    field_p.paragraph_format.space_before = Pt(0)
    field_p.paragraph_format.space_after = Pt(0)
    field_run = field_p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\f C '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    field_run._r.append(fld_begin)
    field_run._r.append(instr)
    field_run._r.append(fld_sep)

    if toc_entries:
        for entry in toc_entries:
            p = add_sdt_paragraph(doc, sdt_content)
            format_toc_entry(p, entry["level"])
            run = p.add_run(f"{entry['title']}\t{entry['page']}")
            set_run_font(run, size=14, bold=entry["level"] == 1)

    end_p = add_sdt_paragraph(doc, sdt_content)
    end_p.paragraph_format.space_before = Pt(0)
    end_p.paragraph_format.space_after = Pt(0)
    end_run = end_p.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_end)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    copy_header_footer_refs(toc_section, body_section)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    if level == 1:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 12, bold=True, font_cn="黑体")
    return p


def set_table_borders(table, mode="three_line"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    def border(name, val, sz="8"):
        element = borders.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")

    if mode == "grid":
        for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border(name, "single", "6")
        return

    border("top", "single", "12")
    border("bottom", "single", "12")
    border("left", "nil")
    border("right", "nil")
    border("insideV", "nil")
    border("insideH", "nil")
    # Header bottom line.
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.first_child_found_in("w:tcBorders")
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:color"), "000000")
        tc_borders.append(bottom)


def add_caption(doc, text, kind="图"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=False)


def add_table_from_data(doc, table_data, caption_no):
    add_caption(doc, f"表{caption_no} {table_data['caption']}", kind="表")
    headers = table_data["headers"]
    rows = table_data["rows"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = table_data.get("widths") or [15.5 / len(headers)] * len(headers)
    for idx, header in enumerate(headers):
        table.columns[idx].width = Cm(widths[idx])
        set_cell_text(table.rows[0].cells[idx], header, size=10.5, bold=True)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            table.columns[idx].width = Cm(widths[idx])
            set_cell_text(row.cells[idx], str(value), size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT if idx > 0 else WD_ALIGN_PARAGRAPH.CENTER)
    set_table_borders(table)
    doc.add_paragraph()


def make_experiment_results_table(metrics):
    rows = []
    for item in metrics["results"]:
        rows.append(
            [
                item["algorithm"],
                item["clusters"],
                f"{item['ari']:.4f}",
                f"{item['nmi']:.4f}",
                f"{item['silhouette']:.4f}",
                f"{item['runtime_sec']:.4f}",
                item.get("candidate_pairs", item.get("all_pairs", "")),
            ]
        )
    return {
        "caption": "聚类算法实验结果对比",
        "headers": ["算法", "聚类数", "ARI", "NMI", "轮廓系数", "时间/s", "候选对"],
        "rows": rows,
        "widths": [3.0, 2.0, 2.0, 2.0, 2.3, 2.1, 2.3],
    }


def add_figure(doc, figure, figure_no):
    if figure.get("kind", "image") != "image":
        raise ValueError(f"Figure {figure.get('id', figure_no)} must reference an image file, not a Word-built diagram.")
    image_path = PAPER_DIR / figure["path"]
    if not image_path.exists():
        raise FileNotFoundError(f"Figure asset not found: {image_path}")
    if figure.get("source") not in {"imagegen", "experiment", "runtime", "user_provided"}:
        raise ValueError(f"Figure {figure.get('id', figure_no)} must declare a valid source category.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(float(figure.get("width_cm", 14.0))))
    add_caption(doc, f"图{figure_no} {figure['caption']}", kind="图")


def add_code_block(doc, block, code_no):
    add_caption(doc, f"代码{code_no} {block['caption']}", kind="代码")
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.05
    for idx, line in enumerate(block["code"]):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, size=9, font_cn="Courier New", font_en="Courier New")
    set_table_borders(table, mode="grid")
    doc.add_paragraph()


def latex_to_omml(latex: str) -> str:
    def add_math_namespaces(xml: str) -> str:
        if "xmlns:m=" in xml:
            return xml
        ns = (
            ' xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
            ' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        )
        if xml.startswith("<m:oMathPara"):
            return xml.replace("<m:oMathPara", f"<m:oMathPara{ns}", 1)
        if xml.startswith("<m:oMath"):
            return xml.replace("<m:oMath", f"<m:oMath{ns}", 1)
        return xml

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        md = tmp_path / "formula.md"
        out = tmp_path / "formula.docx"
        md.write_text(f"$${latex}$$\n", encoding="utf-8")
        pandoc_bin = shutil.which("pandoc")
        if pandoc_bin is None:
            try:
                import pypandoc

                pandoc_bin = pypandoc.get_pandoc_path()
            except Exception as exc:
                raise RuntimeError("Pandoc is required to convert formula LaTeX to Word OMML.") from exc
        subprocess.run(
            [pandoc_bin, "-f", "markdown+tex_math_dollars", "-t", "docx", str(md), "-o", str(out)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        with zipfile.ZipFile(out) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        match = re.search(r"(<m:oMathPara[\s\S]*?</m:oMathPara>)", xml)
        if match:
            return add_math_namespaces(match.group(1))
        match = re.search(r"(<m:oMath[\s\S]*?</m:oMath>)", xml)
        if match:
            return add_math_namespaces(match.group(1))
    raise RuntimeError(f"Formula conversion failed: {latex}")


def add_formula(doc, block):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(7.4), WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Cm(15.2), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    omml_xml = latex_to_omml(block["latex"])
    p._p.append(parse_xml(omml_xml))
    run = p.add_run(f"\t{block['number']}")
    set_run_font(run, size=12)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = borders.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def patch_border_none():
    original = set_table_borders

    def wrapped(table, mode="three_line"):
        if mode == "none":
            remove_table_borders(table)
            return
        original(table, mode)

    return wrapped


set_table_borders = patch_border_none()


def validate_citations(chapters, ref_count):
    markers = []

    def visit_block(block):
        if isinstance(block, str):
            markers.extend(int(x) for x in re.findall(r"\[\[(\d+)\]\]", block))

    for chapter in chapters:
        for section in chapter["sections"]:
            for block in section["blocks"]:
                visit_block(block)
    expected = list(range(1, ref_count + 1))
    first_seen = []
    for marker in markers:
        if marker not in first_seen:
            first_seen.append(marker)
    if first_seen != expected:
        raise RuntimeError(f"Citation first-use order mismatch: {first_seen} != {expected}")
    if any(markers[i] > markers[i + 1] for i in range(len(markers) - 1)):
        raise RuntimeError(f"Citation sequence moves backward: {markers}")


def add_chapters(doc, chapters, tables, figures, code_blocks, metrics):
    table_no = 1
    figure_no = 1
    code_no = 1
    tables_by_id = {item["id"]: item for item in tables}
    figures_by_id = {item["id"]: item for item in figures}
    code_by_id = {item["id"]: item for item in code_blocks}

    for ch_idx, chapter in enumerate(chapters, start=1):
        add_heading(doc, f"第{ch_idx}章 {chapter['title']}", 1)
        for sec_idx, section in enumerate(chapter["sections"], start=1):
            add_heading(doc, f"{ch_idx}.{sec_idx} {section['title']}", 2)
            for block in section["blocks"]:
                if isinstance(block, str):
                    add_body_para(doc, block)
                elif block["type"] == "formula":
                    add_formula(doc, block)
                elif block["type"] == "table":
                    if block["id"] == "experiment_results":
                        table_data = make_experiment_results_table(metrics)
                    else:
                        table_data = tables_by_id[block["id"]]
                    add_table_from_data(doc, table_data, table_no)
                    table_no += 1
                elif block["type"] == "figure":
                    add_figure(doc, figures_by_id[block["id"]], figure_no)
                    figure_no += 1
                elif block["type"] == "code":
                    add_code_block(doc, code_by_id[block["id"]], code_no)
                    code_no += 1
                else:
                    raise ValueError(f"Unknown block type: {block['type']}")


def add_references(doc, references):
    add_heading(doc, "参考文献", 1)
    for idx, ref in enumerate(references, start=1):
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        run = p.add_run(f"[{idx}] {ref}")
        set_run_font(run, size=10.5)


def add_acknowledgement(doc, meta):
    add_heading(doc, "致谢", 1)
    for paragraph in meta["acknowledgement"]:
        add_body_para(doc, paragraph)


def add_appendix(doc, meta):
    add_heading(doc, meta["appendix_title"], 1)
    for paragraph in meta["appendix_content"]:
        add_body_para(doc, paragraph)


def count_formula_blocks(chapters) -> int:
    total = 0
    for chapter in chapters:
        for section in chapter["sections"]:
            total += sum(1 for block in section["blocks"] if isinstance(block, dict) and block.get("type") == "formula")
    return total


def count_docx_math(docx_path: Path) -> int:
    with zipfile.ZipFile(docx_path) as zf:
        total = 0
        for name in zf.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                xml = zf.read(name).decode("utf-8", errors="ignore")
                total += xml.count("<m:oMath")
        return total


def assert_no_placeholders(docx_path: Path) -> None:
    with zipfile.ZipFile(docx_path) as zf:
        xml = "\n".join(zf.read(name).decode("utf-8", errors="ignore") for name in zf.namelist() if name.startswith("word/") and name.endswith(".xml"))
    forbidden = ["__TABLE:", "__CODE:", "{{", "}}", "Тема ВКР", "论文标题中俄文", "здесь вводится", "[此处键入正文]"]
    found = [item for item in forbidden if item in xml]
    if found:
        raise RuntimeError(f"Unresolved or stale template text remains: {found}")


def build() -> Path:
    meta = load_json("meta.json")
    chapters = load_json("chapters.json")
    references = load_json("references.json")
    figures = load_json("figures.json")
    tables = load_json("tables.json")
    code_blocks = load_json("code_blocks.json")
    metrics = json.loads((DATA_DIR / "experiment_metrics.json").read_text(encoding="utf-8"))
    # Keep the TOC as a Word field, but allow a rendered-page cache when
    # LibreOffice does not populate a newly inserted TOC field during export.
    # The cache is regenerated from render output instead of hand-written.
    toc_entries_path = CONTENT_DIR / "toc_pages.json"
    toc_entries = load_json("toc_pages.json") if toc_entries_path.exists() else None

    validate_citations(chapters, len(references))

    doc = Document(str(TEMPLATE_PATH))
    prepare_template_front_matter(doc, meta)
    add_toc(doc, toc_entries)
    add_chapters(doc, chapters, tables, figures, code_blocks, metrics)
    add_references(doc, references)
    add_acknowledgement(doc, meta)
    materialize_linked_header_footer_refs(doc)
    doc.save(OUTPUT_PATH)

    expected_math = count_formula_blocks(chapters)
    actual_math = count_docx_math(OUTPUT_PATH)
    if actual_math < expected_math:
        raise RuntimeError(f"Formula count mismatch: expected at least {expected_math}, found {actual_math}")
    assert_no_placeholders(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build()
    print(path)
