#!/usr/bin/env python3
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate
from PIL import Image


PROJECT_ROOT = Path(__file__).resolve().parent
CONTENT_DIR = PROJECT_ROOT / "content"
FIGURE_DIR = PROJECT_ROOT / "images"
TEMPLATE_PATH = Path("/Users/a136/vs/45425/thesis_project/templates/nylg/template.docx")
OUTPUT_PATH = PROJECT_ROOT / "基于Spark的豆瓣影视数据分析系统的设计与实现_论文初稿.docx"
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def load_json(name: str):
    return json.loads((CONTENT_DIR / name).read_text(encoding="utf-8"))


def build_context() -> dict:
    meta = load_json("meta.json")
    references = load_json("references.json")
    return {
        **meta,
        "chapters": load_json("chapters.json"),
        "references": references,
    }


def set_run_font(run, name: str = "宋体", size_pt: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:ascii"), "Times New Roman" if name == "宋体" else name)
    r_fonts.set(qn("w:hAnsi"), "Times New Roman" if name == "宋体" else name)


def clear_paragraph(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        run.text = ""


def replace_paragraph_text(paragraph: Paragraph, text: str, size_pt: float = 12) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", size_pt)
    paragraph.paragraph_format.line_spacing = Pt(20)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def has_visible_content(paragraph: Paragraph) -> bool:
    if paragraph.text.strip():
        return True
    return any(list(paragraph._element.iter(qn(tag))) for tag in ("w:drawing", "w:pict", "w:object"))


def has_section_break(paragraph: Paragraph) -> bool:
    p_pr = paragraph._element.pPr
    return p_pr is not None and p_pr.find(qn("w:sectPr")) is not None


def fit_image_width_cm(image_path: Path, requested_width: float, max_height_cm: float = 10.8) -> float:
    width = min(float(requested_width), 14.2)
    with Image.open(image_path) as image:
        px_w, px_h = image.size
    if px_w and px_h:
        height = width * px_h / px_w
        if height > max_height_cm:
            width = max_height_cm * px_w / px_h
    return max(5.0, round(width, 1))


def apply_picture_crop(inline_shape, crop: dict | None) -> None:
    if not crop:
        return
    blip_fills = inline_shape._inline.xpath(".//pic:blipFill")
    if not blip_fills:
        return
    blip_fill = blip_fills[0]
    src_rect = blip_fill.find(qn("a:srcRect"))
    if src_rect is None:
        src_rect = OxmlElement("a:srcRect")
        blip_fill.append(src_rect)
    mapping = {"left": "l", "right": "r", "top": "t", "bottom": "b"}
    for key, attr in mapping.items():
        if key in crop and crop[key]:
            src_rect.set(attr, str(int(float(crop[key]) * 1000)))


def insert_figure_after(anchor: Paragraph, image_path: Path, caption: str, requested_width: float, crop: dict | None = None) -> Paragraph:
    if not image_path.exists():
        raise FileNotFoundError(f"缺少论文插图: {image_path}")

    image_para = insert_paragraph_after(anchor)
    image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_para.paragraph_format.space_before = Pt(4)
    image_para.paragraph_format.space_after = Pt(2)
    image_para.paragraph_format.keep_with_next = True
    width = fit_image_width_cm(image_path, requested_width)
    inline_shape = image_para.add_run().add_picture(str(image_path), width=Cm(width))
    apply_picture_crop(inline_shape, crop)

    caption_para = insert_paragraph_after(image_para)
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_para.paragraph_format.space_before = Pt(0)
    caption_para.paragraph_format.space_after = Pt(6)
    run = caption_para.add_run(caption)
    set_run_font(run, "黑体", 10.5, bold=False)
    return caption_para


def inject_figures(docx_path: Path) -> None:
    document = Document(str(docx_path))
    figures = load_json("figures.json")
    inserted: set[str] = set()

    for paragraph in list(document.paragraphs):
        text = paragraph.text
        for group in figures:
            anchor_text = group["anchor"]
            if anchor_text not in text or anchor_text in inserted:
                continue
            replace_paragraph_text(paragraph, group["replacement"])
            anchor = paragraph
            for item in group["items"]:
                file_name, caption, width = item[:3]
                crop = item[3] if len(item) > 3 else None
                anchor = insert_figure_after(anchor, FIGURE_DIR / file_name, caption, width, crop)
            inserted.add(anchor_text)
            break

    missing = {group["anchor"] for group in figures} - inserted
    if missing:
        raise RuntimeError("未找到插图锚点段落: " + "、".join(sorted(missing)))
    document.save(str(docx_path))


def set_cell_border(cell, top=None, bottom=None, left=None, right=None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.first_child_found_in("w:tcBorders")
    if old is not None:
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        if spec is None:
            continue
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), spec.get("val", "single"))
        border.set(qn("w:sz"), spec.get("sz", "4"))
        border.set(qn("w:color"), "000000")
        border.set(qn("w:space"), "0")
        borders.append(border)
    tc_pr.append(borders)


def format_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = Pt(16)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                for run in paragraph.runs:
                    set_run_font(run, "宋体", 9.5, bold=False)
            set_cell_border(cell, left={"val": "nil"}, right={"val": "nil"})
            if row_idx == 0:
                set_cell_border(cell, top={"val": "single", "sz": "12"}, bottom={"val": "single", "sz": "6"}, left={"val": "nil"}, right={"val": "nil"})
            elif row_idx == len(table.rows) - 1:
                set_cell_border(cell, bottom={"val": "single", "sz": "12"}, left={"val": "nil"}, right={"val": "nil"})


def set_table_column_widths(table, widths: list[float] | None) -> None:
    if not widths:
        return
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(float(width))


def insert_tables(docx_path: Path) -> None:
    document = Document(str(docx_path))
    tables = load_json("tables.json")
    pattern = re.compile(r"__TABLE:([A-Za-z0-9_]+)__")
    inserted: set[str] = set()

    for paragraph in list(document.paragraphs):
        match = pattern.search(paragraph.text)
        if not match:
            continue
        table_id = match.group(1)
        data = tables[table_id]
        clear_paragraph(paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(3)
        caption_run = paragraph.add_run(data["caption"])
        set_run_font(caption_run, "黑体", 10.5)

        table = paragraph._parent.add_table(rows=1, cols=len(data["headers"]), width=Cm(15.0))
        for col_idx, value in enumerate(data["headers"]):
            table.rows[0].cells[col_idx].text = str(value)
        for row_values in data["rows"]:
            cells = table.add_row().cells
            for col_idx, value in enumerate(row_values):
                cells[col_idx].text = str(value)
        set_table_column_widths(table, data.get("widths"))
        format_table(table)
        paragraph._p.addnext(table._tbl)
        inserted.add(table_id)

    missing = set(tables) - inserted
    if missing:
        raise RuntimeError("未插入表格占位符: " + "、".join(sorted(missing)))
    document.save(str(docx_path))


def shade_cell(cell, fill: str = "F5F5F5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def clear_cell_borders_and_shading(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("w:shd", "w:tcBorders"):
        node = tc_pr.find(qn(tag))
        if node is not None:
            tc_pr.remove(node)


def insert_code_blocks(docx_path: Path) -> None:
    document = Document(str(docx_path))
    code_blocks = load_json("code_blocks.json")
    pattern = re.compile(r"__CODE:([A-Za-z0-9_]+)__")
    inserted: set[str] = set()

    for paragraph in list(document.paragraphs):
        match = pattern.search(paragraph.text)
        if not match:
            continue
        block_id = match.group(1)
        data = code_blocks[block_id]

        table = paragraph._parent.add_table(rows=1, cols=1, width=Cm(14.2))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.rows[0].cells[0]
        clear_cell_borders_and_shading(cell)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "6"},
            bottom={"val": "single", "sz": "6"},
            left={"val": "single", "sz": "6"},
            right={"val": "single", "sz": "6"},
        )
        p = cell.paragraphs[0]
        clear_paragraph(p)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(12)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        for line_idx, line in enumerate(data["lines"]):
            if line_idx:
                p.add_run().add_break()
            run = p.add_run(line)
            set_run_font(run, "Times New Roman", 9.5)
        paragraph._p.addnext(table._tbl)
        remove_paragraph(paragraph)
        inserted.add(block_id)

    missing = set(code_blocks) - inserted
    if missing:
        raise RuntimeError("未插入代码占位符: " + "、".join(sorted(missing)))
    document.save(str(docx_path))


def remove_numpr(paragraph_or_style) -> None:
    element = paragraph_or_style._element if hasattr(paragraph_or_style, "_element") else paragraph_or_style.element
    p_pr = element.get_or_add_pPr()
    for node in list(p_pr.findall(qn("w:numPr"))):
        p_pr.remove(node)


def normalize_keywords(docx_path: Path, data: dict) -> None:
    document = Document(str(docx_path))
    replacements = {
        "摘要中文关键词": (data["keywords_zh"], "宋体"),
        "摘要英文关键词": (data["keywords_en"], "Times New Roman"),
    }
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name not in replacements:
            continue
        text, font_name = replacements[style_name]
        clear_paragraph(paragraph)
        run = paragraph.add_run()
        run.add_tab()
        run.add_text(text)
        set_run_font(run, font_name, 12)
    document.save(str(docx_path))


def normalize_abstract_identity(docx_path: Path, data: dict) -> None:
    document = Document(str(docx_path))
    source = f"{data['major']} {data['name']}"
    target = f"{data['major']}  {data['name']}"
    source_en = f"{data['major_en']} Major {data['name_en']}"
    for paragraph in document.paragraphs:
        if paragraph.text.strip() != source:
            if paragraph.text.strip() == source_en:
                clear_paragraph(paragraph)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(source_en)
                set_run_font(run, "Times New Roman", 12, bold=True)
            continue
        clear_paragraph(paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(target)
        set_run_font(run, "黑体", 14, bold=False)
    document.save(str(docx_path))


def normalize_front_matter_titles(docx_path: Path, data: dict) -> None:
    document = Document(str(docx_path))
    title = data["title_zh"]
    title_paragraphs = [p for p in document.paragraphs if p.text.strip() == title]
    title_specs = [
        (22, False),  # 内封统计页中文题名
        (16, True),   # 内封信息页中文题名
        (16, False),  # 中文摘要页中文题名
    ]
    for paragraph, (size, bold) in zip(title_paragraphs, title_specs):
        clear_paragraph(paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        set_run_font(run, "黑体", size, bold=bold)

    toc_title = next((p for p in document.paragraphs if p.text.strip() == "目    录"), None)
    if toc_title is not None:
        clear_paragraph(toc_title)
        toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = toc_title.add_run("目    录")
        set_run_font(run, "黑体", 16, bold=True)
    document.save(str(docx_path))


def normalize_body(docx_path: Path) -> None:
    document = Document(str(docx_path))
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name == "论文正文":
            paragraph.paragraph_format.line_spacing = Pt(20)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p_pr = paragraph._p.get_or_add_pPr()
            for tag in ("w:keepNext", "w:keepLines"):
                for node in list(p_pr.findall(qn(tag))):
                    p_pr.remove(node)
            text = paragraph.text.strip()
            is_caption = re.match(r"^(图|表)\d+-\d+", text)
            for run in paragraph.runs:
                if is_caption:
                    set_run_font(run, "黑体", 10.5, bold=False)
                else:
                    set_run_font(run, "宋体", 12)
        if re.match(r"^图\d+-\d+|^表\d+-\d+", paragraph.text.strip()):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.save(str(docx_path))


def superscript_citations(docx_path: Path) -> None:
    document = Document(str(docx_path))
    citation_re = re.compile(r"(\[[0-9,\-—–，、 ]+\])")
    for paragraph in document.paragraphs:
        raw = "".join(run.text for run in paragraph.runs)
        text = raw.strip()
        if not citation_re.search(raw) or text.startswith("["):
            continue
        clear_paragraph(paragraph)
        for part in citation_re.split(raw):
            if not part:
                continue
            run = paragraph.add_run(part)
            set_run_font(run, "宋体", 12)
            if citation_re.fullmatch(part):
                run.font.superscript = True
    document.save(str(docx_path))


def fix_headers(docx_path: Path, title: str) -> None:
    document = Document(str(docx_path))
    for section in document.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            for paragraph in header.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.left_indent = Pt(0)
                paragraph.paragraph_format.right_indent = Pt(0)
                paragraph.paragraph_format.first_line_indent = Pt(0)
                if "实时电商数据分析平台" in paragraph.text or title in paragraph.text:
                    clear_paragraph(paragraph)
                    run = paragraph.add_run(title)
                    set_run_font(run, "宋体", 10.5, bold=False)
    document.save(str(docx_path))


def remove_empty_artifacts_around_references(docx_path: Path) -> None:
    document = Document(str(docx_path))
    paragraphs = list(document.paragraphs)
    ref_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "参考文献")
    for paragraph in reversed(paragraphs[:ref_idx]):
        if has_visible_content(paragraph) or has_section_break(paragraph):
            break
        remove_paragraph(paragraph)

    paragraphs = list(document.paragraphs)
    ref_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "参考文献")
    ack_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "致谢")
    for paragraph in reversed(paragraphs[ref_idx + 1:ack_idx]):
        if not has_visible_content(paragraph) and not has_section_break(paragraph):
            remove_paragraph(paragraph)
    document.save(str(docx_path))


def collect_toc_entries(data: dict) -> list[tuple[int, str]]:
    entries: list[tuple[int, str]] = []
    for ch in data["chapters"]:
        entries.append((1, ch["title"]))
        for section in ch["sections"]:
            entries.append((2, section["title"]))
            for subsection in section.get("subsections", []):
                entries.append((3, subsection["title"]))
    entries.extend([(1, "参考文献"), (1, "致谢")])
    return entries


def estimate_toc_page(title: str) -> int:
    toc_pages_path = CONTENT_DIR / "toc_pages.json"
    if toc_pages_path.exists():
        toc_pages = json.loads(toc_pages_path.read_text(encoding="utf-8"))
        if title in toc_pages:
            return int(toc_pages[title])
        compact_title = re.sub(r"\s+", "", title)
        for toc_title, page in toc_pages.items():
            if re.sub(r"\s+", "", toc_title) == compact_title:
                return int(page)

    hints = {
        "1": 1, "2": 5, "3": 10, "4": 15, "5": 24, "6": 42, "7": 46,
        "参考文献": 48, "致谢": 49,
    }
    prefix = title.split()[0]
    while prefix:
        if prefix in hints:
            return hints[prefix]
        if "." not in prefix:
            break
        prefix = prefix.rsplit(".", 1)[0]
    return 1


def add_toc_field_begin(paragraph: Paragraph) -> None:
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    fld_begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\\\o "1-3" \\\\h \\\\z \\\\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run = paragraph.add_run()
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)


def add_toc_field_end(paragraph: Paragraph) -> None:
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(fld_end)


def rebuild_toc(docx_path: Path, data: dict) -> None:
    document = Document(str(docx_path))
    paragraphs = list(document.paragraphs)
    toc_title_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "目    录")
    body_start_idx = next(
        i for i, p in enumerate(paragraphs)
        if p.text.strip() == "1 绪论" and p.style is not None and p.style.name == "Heading 1"
    )
    toc_title = paragraphs[toc_title_idx]
    for paragraph in reversed(paragraphs[toc_title_idx + 1: body_start_idx]):
        if has_section_break(paragraph):
            clear_paragraph(paragraph)
            continue
        remove_paragraph(paragraph)

    anchor = toc_title
    first = None
    last = None
    for level, title in collect_toc_entries(data):
        toc_para = insert_paragraph_after(anchor)
        try:
            toc_para.style = f"toc {level}"
        except KeyError:
            pass
        toc_para.paragraph_format.left_indent = Cm({1: 0, 2: 0.35, 3: 0.70}.get(level, 0))
        toc_para.paragraph_format.space_before = Pt(0)
        toc_para.paragraph_format.space_after = Pt(0)
        toc_para.paragraph_format.line_spacing = 1.25
        toc_para.paragraph_format.tab_stops.add_tab_stop(Cm(15.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        toc_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = toc_para.add_run(f"{title}\t{estimate_toc_page(title)}")
        set_run_font(run, "宋体", 12, bold=False)
        first = first or toc_para
        last = toc_para
        anchor = toc_para
    if first and last:
        add_toc_field_begin(first)
        add_toc_field_end(last)
    document.save(str(docx_path))


def enforce_major_page_breaks(docx_path: Path) -> None:
    document = Document(str(docx_path))
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if text and (style_name == "Heading 1" or text in {"参考文献", "致谢"}):
            paragraph.paragraph_format.page_break_before = True
    document.save(str(docx_path))


def set_update_fields_on_open(docx_path: Path) -> None:
    with zipfile.ZipFile(docx_path, "r") as zf:
        files = {name: zf.read(name) for name in zf.namelist()}
    xml = files.get(
        "word/settings.xml",
        (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:settings xmlns:w="{WNS}"></w:settings>'
        ).encode("utf-8"),
    ).decode("utf-8")
    if "w:updateFields" not in xml:
        xml = xml.replace("</w:settings>", '<w:updateFields w:val="true"/></w:settings>')
    files["word/settings.xml"] = xml.encode("utf-8")
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    tmp.replace(docx_path)


def repair_docx_compatibility(docx_path: Path) -> None:
    soffice = shutil.which("soffice")
    if not soffice:
        return
    with tempfile.TemporaryDirectory(prefix="nylg_docx_repair_") as temp_dir:
        temp_path = Path(temp_dir)
        profile_dir = temp_path / "lo_profile"
        out_dir = temp_path / "out"
        profile_dir.mkdir(parents=True, exist_ok=True)
        out_dir.mkdir(parents=True, exist_ok=True)
        env = os.environ.copy()
        env.setdefault("TMPDIR", "/private/tmp")
        env["HOME"] = str(profile_dir)
        command = [
            soffice,
            f"-env:UserInstallation=file://{profile_dir}",
            "--invisible",
            "--headless",
            "--norestore",
            "--convert-to",
            "docx",
            "--outdir",
            str(out_dir),
            str(docx_path),
        ]
        result = subprocess.run(command, check=False, capture_output=True, text=True, env=env)
        repaired_path = out_dir / docx_path.name
        if result.returncode != 0 or not repaired_path.exists():
            detail = (result.stderr or result.stdout or "未知错误").strip()
            raise RuntimeError(f"DOCX兼容性修复失败: {detail}")
        shutil.copy2(repaired_path, docx_path)


def main() -> None:
    data = build_context()
    doc = DocxTemplate(str(TEMPLATE_PATH))
    doc.render(data)
    doc.save(str(OUTPUT_PATH))
    fix_headers(OUTPUT_PATH, data["title_zh"])
    normalize_body(OUTPUT_PATH)
    inject_figures(OUTPUT_PATH)
    insert_tables(OUTPUT_PATH)
    insert_code_blocks(OUTPUT_PATH)
    normalize_keywords(OUTPUT_PATH, data)
    normalize_abstract_identity(OUTPUT_PATH, data)
    superscript_citations(OUTPUT_PATH)
    normalize_body(OUTPUT_PATH)
    normalize_front_matter_titles(OUTPUT_PATH, data)
    remove_empty_artifacts_around_references(OUTPUT_PATH)
    enforce_major_page_breaks(OUTPUT_PATH)
    rebuild_toc(OUTPUT_PATH, data)
    normalize_front_matter_titles(OUTPUT_PATH, data)
    repair_docx_compatibility(OUTPUT_PATH)
    print(f"已生成: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
