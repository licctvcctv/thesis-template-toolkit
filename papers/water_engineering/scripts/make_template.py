from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parents[1]
SOURCE_UPLOAD = PROJECT / "_source_materials" / "original_uploads" / "water_engineering" / "2026本科毕业设计.docx"
TEMPLATE_DIR = PROJECT / "templates" / "water_engineering"
SOURCE = TEMPLATE_DIR / "source.docx"
TEMPLATE = TEMPLATE_DIR / "template.docx"
DIAGNOSTICS = PROJECT / "_source_materials" / "diagnostics" / "water_engineering"


def text_of(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs)


def replace_paragraph(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
    else:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""


def main() -> None:
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    DIAGNOSTICS.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE_UPLOAD, SOURCE)
    shutil.copy2(SOURCE_UPLOAD, TEMPLATE)

    doc = Document(TEMPLATE)
    paragraphs = doc.paragraphs
    for paragraph in paragraphs:
        text = text_of(paragraph).strip()
        if text.startswith("题目："):
            replace_paragraph(paragraph, "题目：{{ title_zh }}")
        elif text == "学 院":
            replace_paragraph(paragraph, "学 院 {{ college }}")
        elif text == "专 业":
            replace_paragraph(paragraph, "专 业 {{ major }}")
        elif text == "年 级":
            replace_paragraph(paragraph, "年 级 {{ grade }}")
        elif text == "姓 名":
            replace_paragraph(paragraph, "姓 名 {{ student_name }}")
        elif text == "学 号":
            replace_paragraph(paragraph, "学 号 {{ student_id }}")
        elif text == "指导教师":
            replace_paragraph(paragraph, "指导教师 {{ advisor }}")

    if len(paragraphs) > 34:
        replace_paragraph(paragraphs[34], "{{ abstract_zh }}")
    if len(paragraphs) > 36:
        replace_paragraph(paragraphs[36], "关键词：{{ keywords_zh }}")
    if len(paragraphs) > 39:
        replace_paragraph(paragraphs[39], "{{ abstract_en }}")
    if len(paragraphs) > 41:
        replace_paragraph(paragraphs[41], "KEY WORDS: {{ keywords_en }}")

    manifest = {
        "source": str(SOURCE),
        "template": str(TEMPLATE),
        "content_source": str(ROOT / "content"),
        "dynamic_regions": [
            "cover fields",
            "Chinese abstract and keywords",
            "English abstract and keywords",
            "body chapters generated from content/chapters.json",
            "references generated from content/references.json",
            "acknowledgement generated from content/meta.json"
        ],
        "preserved_regions": [
            "school cover layout",
            "declaration page wording",
            "TOC field/cache region",
            "heading styles and page/section setup"
        ]
    }
    (DIAGNOSTICS / "placeholder_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    doc.save(TEMPLATE)


if __name__ == "__main__":
    main()
