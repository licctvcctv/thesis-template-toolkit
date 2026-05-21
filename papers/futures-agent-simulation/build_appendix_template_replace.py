from __future__ import annotations

import json
import re
from xml.sax.saxutils import escape
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.oxml.ns import nsdecls, qn


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
SOURCE_DIR = PROJECT / "_source_materials" / "futures-agent-simulation" / "appendix_inputs"
TEMPLATE = SOURCE_DIR / "appendix_template_sample.docx"
OPENING = SOURCE_DIR / "opening_report_confirmed.docx"
CONTENT = ROOT / "content" / "appendix.json"
OUTPUT = ROOT / "附录_基于自进化LLM的期货市场多智能体决策模拟系统.docx"


def set_run_font(run, size: float | None = None, font: str | None = None) -> None:
    if font:
        run.font.name = font
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if size:
        run.font.size = Pt(size)


def replace_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_cell_text(cell, text: str) -> None:
    p = cell.paragraphs[0]
    replace_paragraph_text(p, text)
    for paragraph in cell.paragraphs[1:]:
        replace_paragraph_text(paragraph, "")


def replace_paragraph_with_math(paragraph, text: str) -> None:
    replace_paragraph_text(paragraph, "")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_xml = (
        f'<m:oMathPara {nsdecls("m")}>'
        f"<m:oMath><m:r><m:t>{escape(text)}</m:t></m:r></m:oMath>"
        f"</m:oMathPara>"
    )
    paragraph._p.append(parse_xml(math_xml))


def iter_body_paragraphs(doc: Document):
    for child in doc.element.body:
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)


def find_paragraph(doc: Document, exact_text: str) -> Paragraph:
    for paragraph in iter_body_paragraphs(doc):
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise ValueError(f"未找到段落：{exact_text}")


def clone_paragraph_after(anchor, template, text: str = "") -> Paragraph:
    new_p = deepcopy(template._p)
    anchor._p.addnext(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    replace_paragraph_text(paragraph, text)
    return paragraph


def insert_after_paragraph(anchor, template, texts: list[str]) -> Paragraph:
    current = anchor
    for text in texts:
        current = clone_paragraph_after(current, template, text)
    return current


def delete_between(start_para: Paragraph, end_para: Paragraph | None) -> None:
    body = start_para._p.getparent()
    start_idx = list(body).index(start_para._p)
    end_idx = list(body).index(end_para._p) if end_para is not None else len(body)
    for child in list(body)[start_idx + 1:end_idx]:
        body.remove(child)


def copy_docx_body_into_section(doc: Document, marker: str, next_marker: str | None, source_path: Path) -> None:
    start = find_paragraph(doc, marker)
    if marker.startswith("附件"):
        start.paragraph_format.page_break_before = True
    end = find_paragraph(doc, next_marker) if next_marker else None
    delete_between(start, end)

    current = start._p
    source = Document(str(source_path))
    for child in source.element.body:
        if child.tag not in {qn("w:p"), qn("w:tbl")}:
            continue
        copied = deepcopy(child)
        current.addnext(copied)
        current = copied


def build_cover(doc: Document, meta: dict) -> None:
    table = doc.tables[0]
    replace_cell_text(table.cell(0, 1), meta.get("cover_title", meta["title"]))
    replace_cell_text(table.cell(1, 1), meta.get("college", ""))
    replace_cell_text(table.cell(2, 1), meta.get("major", ""))
    replace_cell_text(table.cell(3, 1), meta.get("name", ""))
    class_id = "/".join(part for part in [meta.get("class_name", ""), meta.get("student_id", "")] if part)
    replace_cell_text(table.cell(3, 3), class_id)
    replace_cell_text(table.cell(4, 2), meta.get("advisor", ""))
    date_range = " 至 ".join(part for part in [meta.get("start_date", ""), meta.get("end_date", "")] if part)
    replace_cell_text(table.cell(5, 1), date_range)


def build_directory(doc: Document, directory: list[str]) -> None:
    for old, new in zip(
        [
            "附件1：开题报告………………………………………………共 4页",
            "附件2：主要源代码……………………………………………共67页",
            "附件3：外文文献译文…………………………………………共18页",
            "附件4：外文文献原文…………………………………………共20页",
        ],
        directory,
    ):
        replace_paragraph_text(find_paragraph(doc, old), new)


def opening_lines(meta: dict) -> list[str]:
    doc = Document(str(OPENING))
    raw = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    lines = [
        meta["title"],
        "开题报告",
    ]
    skip = {meta["title"], "开题报告"}
    lines.extend(text for text in raw if text not in skip)
    return lines


def read_code_block(spec: dict) -> str:
    path = ROOT / spec["path"]
    lines = path.read_text(encoding="utf-8").splitlines()
    start = int(spec.get("start", 1))
    end = int(spec.get("end", len(lines)))
    selected = lines[start - 1: end]
    header = f"// 文件：{spec['path']}（第 {start}-{end} 行，共 {len(selected)} 行）"
    return "\n".join([header, *selected])


def code_lines(data: dict) -> list[str]:
    lines = ["以下代码选自期货市场多智能体决策模拟系统真实项目文件，覆盖智能体画像、历史行情数据加载、世界状态建模、LLM 交易决策、收益评分、后端服务、状态广播和前端交易面板等关键功能。"]
    for spec in data["code_blocks"]:
        lines.extend(["", spec["title"]])
        lines.extend(read_code_block(spec).splitlines())
    return lines


def translation_lines(data: dict) -> list[dict]:
    lit = data["foreign_literature"]
    items: list[dict] = [
        {"type": "text", "text": lit["translation_title"]},
    ]
    paragraphs = lit["translation_paragraphs"]
    figures = lit.get("figures", [])
    figure_slots = {
        2: 0,
        max(5, len(paragraphs) // 3): 1,
        max(8, (len(paragraphs) * 2) // 3): 2,
        max(10, len(paragraphs) - 4): 3,
    }
    for idx, text in enumerate(paragraphs):
        items.append({"type": "text", "text": text})
        fig_idx = figure_slots.get(idx)
        if fig_idx is not None and fig_idx < len(figures):
            items.append({"type": "image", **figures[fig_idx]})
    return items


def is_original_heading(text: str) -> bool:
    return bool(
        re.match(r"^(\d+\.|S\d+\.|[A-Z][A-Za-z ]{2,35}$)", text)
        or text in {"Abstract", "References", "Appendix", "Acknowledgements"}
        or text.startswith(("Fig. ", "Table "))
    )


def original_paragraphs(lit: dict) -> list[str]:
    text_path = ROOT / lit.get("original_text_file", "references/foreign_literature/tradingagents_2412.20138.raw.txt")
    raw = text_path.read_text(encoding="utf-8", errors="ignore").replace("\f", "\n")
    paragraphs: list[str] = []
    current = ""

    def flush() -> None:
        nonlocal current
        clean = re.sub(r"\s+", " ", current).strip()
        if clean:
            paragraphs.append(clean)
        current = ""

    for raw_line in raw.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        if not line:
            flush()
            continue
        if line in {"TradingAgents: Multi-Agents LLM Financial Trading Framework", "TAURIC RESEARCH"}:
            continue
        if re.fullmatch(r"\d{1,2}", line):
            continue
        if is_original_heading(line) or line.startswith("• "):
            flush()
            paragraphs.append(line)
            continue
        if not current:
            current = line
        elif current.endswith("-"):
            current = current[:-1] + line
        else:
            current += " " + line
        if len(current) > 850 and re.search(r"[.!?)]$", line):
            flush()
    flush()
    return paragraphs


def original_lines(data: dict) -> list[dict]:
    lit = data["foreign_literature"]
    items: list[dict] = []
    json_path = ROOT / lit.get("original_json", "content/foreign_original.json")
    if json_path.exists():
        original = json.loads(json_path.read_text(encoding="utf-8"))
        for item in original["items"]:
            item_type = item["type"]
            if item_type in {"title", "authors", "paragraph", "reference"}:
                items.append({"type": "text", "text": item["text"], "role": item_type})
            elif item_type == "heading":
                items.append({"type": "text", "text": item["text"], "role": f"heading{item.get('level', 1)}"})
            elif item_type == "list_item":
                items.append({"type": "text", "text": f"• {item['text']}", "role": "list_item"})
            elif item_type == "equation":
                items.append({"type": "text", "text": item["text"], "role": "equation"})
            elif item_type == "figure" and item.get("src"):
                items.append({"type": "image", "image": item["src"], "caption": item.get("caption", ""), "width_cm": 13.2})
            elif item_type == "table":
                if item.get("caption"):
                    items.append({"type": "text", "text": "Table: " + item["caption"], "role": "table_caption"})
                for row in item.get("rows", [])[:18]:
                    items.append({"type": "text", "text": " | ".join(row), "role": "table_row"})
            elif item_type == "code":
                if item.get("title"):
                    items.append({"type": "text", "text": item["title"], "role": "code_title"})
                for line in item.get("text", "").splitlines()[:90]:
                    if line.strip():
                        items.append({"type": "text", "text": line, "role": "code"})
    else:
        items.extend({"type": "text", "text": text} for text in original_paragraphs(lit))
    return items


def style_inserted_opening(paragraph) -> None:
    text = paragraph.text.strip()
    if text in {"开题报告"}:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if re.match(r"^[一二三四五]、", text):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.bold = True


def style_inserted_original(paragraph) -> None:
    text = paragraph.text.strip()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    role = getattr(paragraph, "_appendix_role", "")
    for run in paragraph.runs:
        if role in {"code", "table_row"}:
            set_run_font(run, size=8.0, font="Consolas")
        elif role == "equation":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(run, size=10.0, font="Times New Roman")
        elif role.startswith("heading") or role in {"title", "code_title", "table_caption"}:
            set_run_font(run, size=10.5, font="Times New Roman")
        else:
            set_run_font(run, size=9.5, font="Times New Roman")
        if role.startswith("heading") or role in {"title", "code_title", "table_caption"} or is_original_heading(text) or text.startswith(("外文文献题名：", "作者：", "来源：", "获取地址：")):
            run.bold = True


def fill_section(doc: Document, marker: str, next_marker: str | None, items, template_para, mode: str = "text") -> None:
    start = find_paragraph(doc, marker)
    if marker.startswith("附件"):
        start.paragraph_format.page_break_before = True
    end = find_paragraph(doc, next_marker) if next_marker else None
    delete_between(start, end)
    current = start
    for item in items:
        if isinstance(item, dict) and item.get("type") == "image":
            current = clone_paragraph_after(current, template_para, "")
            if item.get("page_break_before"):
                current.paragraph_format.page_break_before = True
            if mode == "original":
                current.paragraph_format.keep_with_next = True
            current.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = current.runs[0] if current.runs else current.add_run()
            run.add_picture(str(ROOT / item["image"]), width=Cm(float(item.get("width_cm", 13.2))))
            if item.get("caption"):
                current = clone_paragraph_after(current, template_para, item["caption"])
                current.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if mode == "original":
                    setattr(current, "_appendix_role", "figure_caption")
                    current.paragraph_format.keep_together = True
                    style_inserted_original(current)
            continue
        text = item["text"] if isinstance(item, dict) else item
        current = clone_paragraph_after(current, template_para, text)
        if isinstance(item, dict) and item.get("role"):
            setattr(current, "_appendix_role", item["role"])
        if mode == "code":
            for run in current.runs:
                set_run_font(run, size=8.0, font="Consolas")
        elif mode == "opening":
            style_inserted_opening(current)
        elif mode == "original":
            if isinstance(item, dict) and item.get("role") == "equation":
                replace_paragraph_with_math(current, text)
            else:
                style_inserted_original(current)


def main() -> None:
    data = json.loads(CONTENT.read_text(encoding="utf-8"))
    doc = Document(str(TEMPLATE))
    meta = data["meta"]

    # Cache source paragraph formatting before deleting old appendix content.
    opening_template = doc.paragraphs[19]
    code_template = doc.paragraphs[102]
    trans_template = doc.paragraphs[4035]

    build_cover(doc, meta)
    build_directory(doc, data["directory"])

    copy_docx_body_into_section(doc, "附件1：开题报告", "附件2：主要源代码", OPENING)
    fill_section(doc, "附件2：主要源代码", "附件3：外文文献译文", code_lines(data), code_template, mode="code")
    fill_section(doc, "附件3：外文文献译文", "附件4：外文文献原文", translation_lines(data), trans_template)
    fill_section(doc, "附件4：外文文献原文", None, original_lines(data), trans_template, mode="original")

    doc.save(str(OUTPUT))
    print(f"完成: {OUTPUT}")


if __name__ == "__main__":
    main()
