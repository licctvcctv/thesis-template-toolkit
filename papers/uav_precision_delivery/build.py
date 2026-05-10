from __future__ import annotations

from copy import deepcopy
from functools import lru_cache
import json
import re
import shutil
import socket
import subprocess
import tempfile
import textwrap
import time
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree


ROOT = Path(__file__).resolve().parent
CONTENT = ROOT / "content"
TEMPLATE = ROOT.parents[1] / "templates" / "scu_eie_2026" / "template.docx"
OUTPUT = ROOT / "无人机精准投放技术-游忠锦.docx"
AUTO_TOC_PLACEHOLDER = "__AUTO_TOC_GENERATED_BY_OFFICE__"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
INLINE_MARKUP_RE = re.compile(r"(\[\[m:(.*?)\]\]|\[\[\d+\]\])")


def load_json(name: str) -> Any:
    return json.loads((CONTENT / name).read_text(encoding="utf-8"))


def set_run_font(run, size: int | None = 12, bold: bool | None = None, name: str = "宋体") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_run_mixed_font(
    run,
    size: int | None = 12,
    bold: bool | None = None,
    east_asia: str = "宋体",
    latin: str = "Times New Roman",
) -> None:
    """Match thesis reference typography: Chinese Songti, Latin Times New Roman."""
    run.font.name = latin
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), latin)


def set_paragraph_format(paragraph, first_line: bool = False, line_spacing: float = 1.5) -> None:
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if first_line:
        pf.first_line_indent = Pt(24)


def clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)


def add_cover_line(paragraph, label: str, value: str, pad_left: int = 2, pad_right: int = 10) -> None:
    clear_paragraph(paragraph)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, 14, True)
    value_run = paragraph.add_run(f"{' ' * pad_left}{value}{' ' * pad_right}")
    set_run_font(value_run, 14, True)
    value_run.underline = True
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def add_cover_runs(paragraph, parts: list[tuple[str, bool]]) -> None:
    """Rewrite a cover row while preserving the template's paragraph geometry."""
    clear_paragraph(paragraph)
    for text, underline in parts:
        run = paragraph.add_run(text)
        set_run_font(run, 14, True)
        run.underline = underline
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def cover_field(value: str, left_spaces: int, right_spaces: int) -> str:
    return f"{' ' * left_spaces}{value}{' ' * right_spaces}"


def paragraph_text_from_element(element) -> str:
    texts = []
    for node in element.iter(qn("w:t")):
        if node.text:
            texts.append(node.text)
    return "".join(texts)


def remove_from_anchor(doc: Document, anchor_text: str) -> None:
    body = doc.element.body
    children = list(body.iterchildren())
    start = None
    for idx, child in enumerate(children):
        if child.tag == qn("w:p") and paragraph_text_from_element(child).strip() == anchor_text:
            start = idx
            break
    if start is None:
        raise RuntimeError(f"未找到模板正文锚点：{anchor_text}")
    for child in children[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def replace_template_cover(doc: Document, meta: dict[str, Any]) -> None:
    for paragraph in doc.paragraphs:
        compact = re.sub(r"\s+", "", paragraph.text)
        if compact.startswith("题目"):
            add_cover_runs(paragraph, [
                ("题    目 ", False),
                (cover_field(meta["title_zh"], 8, 8), True),
            ])
        elif compact.startswith("学院"):
            add_cover_runs(paragraph, [
                ("学    院 ", False),
                (cover_field(meta["college"], 12, 12), True),
            ])
        elif compact.startswith("专业"):
            add_cover_runs(paragraph, [
                ("专    业 ", False),
                (cover_field(meta["major"], 14, 14), True),
            ])
        elif compact.startswith("学生姓名"):
            add_cover_runs(paragraph, [
                ("学生姓名 ", False),
                (cover_field(meta["student_name"], 14, 14), True),
            ])
        elif compact.startswith("学号"):
            add_cover_runs(paragraph, [
                ("学    号 ", False),
                (f"  {meta['student_id']}    ", True),
                (" 年级", False),
                (f"  {meta['grade']}     ", True),
            ])
        elif compact.startswith("指导教师"):
            add_cover_runs(paragraph, [
                ("指导教师 ", False),
                (cover_field(meta["advisor"], 14, 14), True),
            ])
        if "教务处制表" in paragraph.text:
            clear_paragraph(paragraph)
            run = paragraph.add_run("教务处制表")
            set_run_font(run, 16, True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if "二〇二四年五月二十五日" in paragraph.text:
            clear_paragraph(paragraph)
            run = paragraph.add_run(meta["date_full_cn"])
            set_run_font(run, 15, True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def update_headers(doc: Document, title: str) -> None:
    header_text = f"四川大学本科毕业论文 \t{title}"
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                clear_paragraph(paragraph)
                run = paragraph.add_run(header_text)
                set_run_font(run, 9)


def add_center(doc: Document, text: str, size: int = 14, bold: bool = True, style: str | None = None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, False, 1.5)
    run = p.add_run(text)
    set_run_font(run, size, bold)
    return p


def add_body_paragraph(doc: Document, text: str, first_line: bool = True):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line, 1.5)
    append_text_with_citations(p, text)
    return p


@lru_cache(maxsize=128)
def latex_to_omath(latex: str):
    """Convert a LaTeX display formula to a Word OMML inline math element."""
    pandoc = shutil.which("pandoc")
    if not pandoc:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        md_path = tmp_dir / "formula.md"
        docx_path = tmp_dir / "formula.docx"
        md_path.write_text(f"$$\n{latex}\n$$\n", encoding="utf-8")
        result = subprocess.run(
            [pandoc, "-f", "markdown+tex_math_dollars", "-t", "docx", str(md_path), "-o", str(docx_path)],
            text=True,
            capture_output=True,
            timeout=20,
        )
        if result.returncode != 0 or not docx_path.exists():
            return None
        with zipfile.ZipFile(docx_path, "r") as zf:
            xml = zf.read("word/document.xml")
    root = etree.fromstring(xml)
    items = root.xpath(".//m:oMath", namespaces={"m": M_NS})
    return deepcopy(items[0]) if items else None


def iter_formula_blocks(nodes: list[dict[str, Any]], chapter_index: int | None = None):
    for idx, node in enumerate(nodes, start=1):
        current_chapter = chapter_index if chapter_index is not None else idx
        if node.get("type") == "formula":
            yield current_chapter, node
        sections = node.get("sections")
        if isinstance(sections, list):
            yield from iter_formula_blocks(sections, current_chapter)
        blocks = node.get("blocks")
        if isinstance(blocks, list):
            for block in blocks:
                if block.get("type") == "formula":
                    yield current_chapter, block


def validate_formula_blocks(chapters: list[dict[str, Any]]) -> list[str]:
    expected_by_chapter: dict[int, int] = {}
    numbers: list[str] = []
    for chapter_index, block in iter_formula_blocks(chapters):
        number = block.get("number")
        latex = block.get("latex")
        if not isinstance(number, str) or not re.fullmatch(r"（\d+\.\d+）", number):
            raise RuntimeError(f"公式编号格式错误：{number!r}，应使用类似（2.1）的全角括号编号")
        if not isinstance(latex, str) or not latex.strip():
            raise RuntimeError(f"公式 {number} 缺少 latex 字段，不能退回普通文本公式")
        major, minor = (int(part) for part in re.fullmatch(r"（(\d+)\.(\d+)）", number).groups())
        if major != chapter_index:
            raise RuntimeError(f"公式 {number} 所属章节不匹配，应位于第 {chapter_index} 章编号体系内")
        expected_minor = expected_by_chapter.get(major, 0) + 1
        if minor != expected_minor:
            raise RuntimeError(f"公式编号不连续：第 {major} 章当前为 {number}，应为（{major}.{expected_minor}）")
        expected_by_chapter[major] = minor
        numbers.append(number)
    return numbers


def add_formula(doc: Document, text: str, number: str, latex: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.0), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(5.8), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    if not latex:
        raise RuntimeError(f"公式 {number} 缺少 latex 字段，不能写成普通文本：{text}")
    omath = latex_to_omath(latex)
    if omath is None:
        raise RuntimeError(f"公式 {number} 转换 Word 数学对象失败：{latex}")
    p._p.append(deepcopy(omath))
    p.add_run("\t")
    num_run = p.add_run(number)
    set_run_font(num_run, 11, False, "Times New Roman")


def append_text_with_citations(paragraph, text: str) -> None:
    pos = 0
    for match in INLINE_MARKUP_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, 12)
        part = match.group(0)
        citation = re.fullmatch(r"\[\[(\d+)\]\]", part)
        if citation:
            run = paragraph.add_run(f"[{citation.group(1)}]")
            set_run_font(run, 10)
            run.font.superscript = True
        else:
            math_marker = re.fullmatch(r"\[\[m:(.*)\]\]", part)
            if not math_marker:
                raise RuntimeError(f"无法识别的行内标记：{part}")
            latex = math_marker.group(1)
            omath = latex_to_omath(latex)
            if omath is None:
                raise RuntimeError(f"行内公式转换失败：{latex}")
            paragraph._p.append(deepcopy(omath))
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, 12)


def iter_text_blocks(nodes: list[dict[str, Any]]):
    for node in nodes:
        if isinstance(node.get("text"), str):
            yield node["text"]
        for key in ("blocks", "sections"):
            children = node.get(key)
            if isinstance(children, list):
                yield from iter_text_blocks(children)


def validate_citations(chapters: list[dict[str, Any]], references: list[str]) -> None:
    citations: list[int] = []
    adjacent_groups: list[str] = []
    for text in iter_text_blocks(chapters):
        adjacent_groups.extend(re.findall(r"(?:\[\[\d+\]\]){2,}", text))
        citations.extend(int(match) for match in re.findall(r"\[\[(\d+)\]\]", text))
    if not citations:
        raise RuntimeError("正文中没有找到参考文献引用")
    if adjacent_groups:
        raise RuntimeError(
            "正文中存在一次打多个参考文献标的情况："
            f"{adjacent_groups}。请拆成单句单标，例如一个论断只跟一个 [[n]]。"
        )

    ref_count = len(references)
    out_of_range = sorted({num for num in citations if num < 1 or num > ref_count})
    if out_of_range:
        raise RuntimeError(f"正文引用编号超出参考文献范围：{out_of_range}")
    if citations != sorted(citations):
        raise RuntimeError(
            "正文引用编号顺序存在回退："
            f"当前序列为 {citations}。请按 [1]、[2]、[3] 的顺序逐个打标，不要在后文回到更小编号。"
        )

    first_seen: list[int] = []
    seen: set[int] = set()
    for num in citations:
        if num not in seen:
            seen.add(num)
            first_seen.append(num)

    expected = list(range(1, len(first_seen) + 1))
    if first_seen != expected:
        raise RuntimeError(
            "参考文献首次引用顺序错误："
            f"首次出现为 {first_seen}，应为 {expected}。请按正文首次出现顺序重排 references.json 并同步正文编号。"
        )

    missing = [num for num in range(1, ref_count + 1) if num not in seen]
    if missing:
        raise RuntimeError(f"参考文献列表中存在未被正文引用的条目：{missing}")


def validate_output_formulas(docx_path: Path, expected_numbers: list[str]) -> None:
    with zipfile.ZipFile(docx_path, "r") as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    paragraphs = root.xpath(".//w:body/w:p", namespaces={"w": W_NS, "m": M_NS})
    actual_numbers: list[str] = []
    omath_count = 0
    for paragraph in paragraphs:
        omaths = paragraph.xpath(".//m:oMath", namespaces={"m": M_NS})
        if not omaths:
            continue
        visible_text = "".join(paragraph.xpath("./w:r/w:t/text()", namespaces={"w": W_NS}))
        numbers = re.findall(r"（\d+\.\d+）", visible_text)
        if len(numbers) != 1:
            continue
        if visible_text.strip() != numbers[0]:
            continue
        omath_count += len(omaths)
        if len(omaths) != 1:
            raise RuntimeError(
                "公式段落结构异常："
                f"数学对象 {len(omaths)} 个，编号 {numbers}，段落文本 {visible_text!r}"
            )
        leaked_text = "".join(paragraph.xpath(".//w:t/text()", namespaces={"w": W_NS})).strip()
        if leaked_text != numbers[0]:
            raise RuntimeError(f"公式 {numbers[0]} 出现普通文本公式残留：{leaked_text!r}")
        actual_numbers.append(numbers[0])
    if omath_count != len(expected_numbers):
        raise RuntimeError(f"公式数量不匹配：JSON 中 {len(expected_numbers)} 个，DOCX 中 {omath_count} 个 Word 数学对象")
    if actual_numbers != expected_numbers:
        raise RuntimeError(f"公式编号顺序不匹配：DOCX 为 {actual_numbers}，JSON 为 {expected_numbers}")


def add_keywords(doc: Document, label: str, words: list[str], english: bool = False) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, False, 1.5)
    label_run = p.add_run(label)
    set_run_font(label_run, 12, True, "Times New Roman" if english else "宋体")
    body_run = p.add_run(("; " if english else "；").join(words) + ("." if english else ""))
    set_run_font(body_run, 12, False, "Times New Roman" if english else "宋体")


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def collect_toc_entries(chapters: list[dict[str, Any]]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []

    def visit(node: dict[str, Any], level: int) -> None:
        title = f"{node['number']} {node['title']}"
        entries.append({"level": level, "text": title})
        for child in node.get("sections", []):
            visit(child, min(level + 1, 3))

    for chapter in chapters:
        visit(chapter, 1)
    entries.extend([{"level": 1, "text": "参考文献"}, {"level": 1, "text": "致谢"}])
    return entries


def add_toc(doc: Document, chapters: list[dict[str, Any]], toc_pages: dict[str, Any]) -> None:
    title_p = add_center(doc, "目录", 22, True)
    title_p.paragraph_format.line_spacing = 1.0
    title_p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(p, False, 1.0)
    run = p.add_run(AUTO_TOC_PLACEHOLDER)
    set_run_font(run, 10, False)


def find_free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def refresh_toc_with_libreoffice(docx_path: Path) -> None:
    soffice = shutil.which("soffice")
    if not soffice:
        candidate = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        if candidate.exists():
            soffice = str(candidate)
    lo_python = Path("/Applications/LibreOffice.app/Contents/Resources/python")
    if not soffice or not lo_python.exists():
        raise RuntimeError("未找到 LibreOffice，无法生成自动目录")

    port = find_free_port()
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        profile_dir = tmp_dir / "lo-profile"
        profile_dir.mkdir()
        script_path = tmp_dir / "refresh_toc.py"
        script_path.write_text(textwrap.dedent(f"""
            import sys
            import time
            import uno
            from com.sun.star.beans import PropertyValue

            docx_path = sys.argv[1]
            port = sys.argv[2]
            placeholder = sys.argv[3]

            def prop(name, value):
                item = PropertyValue()
                item.Name = name
                item.Value = value
                return item

            local_ctx = uno.getComponentContext()
            resolver = local_ctx.ServiceManager.createInstanceWithContext(
                "com.sun.star.bridge.UnoUrlResolver", local_ctx
            )
            ctx = None
            last_error = None
            for _ in range(50):
                try:
                    ctx = resolver.resolve(
                        f"uno:socket,host=127.0.0.1,port={{port}};urp;StarOffice.ComponentContext"
                    )
                    break
                except Exception as exc:
                    last_error = exc
                    time.sleep(0.2)
            if ctx is None:
                raise RuntimeError(f"LibreOffice UNO 连接失败: {{last_error}}")

            desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
            url = uno.systemPathToFileUrl(docx_path)
            doc = desktop.loadComponentFromURL(
                url,
                "_blank",
                0,
                (prop("Hidden", True), prop("UpdateDocMode", 3)),
            )
            if doc is None:
                raise RuntimeError("LibreOffice 未能打开 DOCX")

            search = doc.createSearchDescriptor()
            search.SearchString = placeholder
            found = doc.findFirst(search)
            if found:
                text = found.getText()
                cursor = text.createTextCursorByRange(found)
                found.setString("")
                toc = doc.createInstance("com.sun.star.text.ContentIndex")
                for name, value in (
                    ("Title", ""),
                    ("CreateFromOutline", True),
                    ("Level", 3),
                    ("UseHyperlinks", True),
                ):
                    try:
                        setattr(toc, name, value)
                    except Exception:
                        pass
                text.insertTextContent(cursor, toc, False)

            indexes = doc.getDocumentIndexes()
            for i in range(indexes.getCount()):
                indexes.getByIndex(i).update()
            doc.getTextFields().refresh()
            doc.store()
            doc.close(True)
        """), encoding="utf-8")

        accept = f"socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext"
        proc = subprocess.Popen([
            soffice,
            "--headless",
            "--invisible",
            "--nodefault",
            "--nofirststartwizard",
            "--norestore",
            f"-env:UserInstallation={profile_dir.as_uri()}",
            f"--accept={accept}",
        ], stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        try:
            result = subprocess.run(
                [str(lo_python), str(script_path), str(docx_path.resolve()), str(port), AUTO_TOC_PLACEHOLDER],
                text=True,
                capture_output=True,
                timeout=60,
            )
            if result.returncode != 0:
                raise RuntimeError(result.stderr or result.stdout)
        finally:
            proc.terminate()
            try:
                proc.wait(timeout=5)
            except subprocess.TimeoutExpired:
                proc.kill()


def upsert_xml_child(parent, tag: str):
    child = parent.find(tag)
    if child is None:
        child = etree.SubElement(parent, tag)
    return child


def compact_generated_toc_styles(docx_path: Path) -> None:
    """Keep the Office-generated TOC while matching the school template typography."""
    ns = {"w": W_NS}
    style_tokens = {
        "TOC1": {"line": "400", "size": "28", "left": None, "left_chars": None, "first_line": "420", "first_line_chars": "200"},
        "TOC2": {"line": "400", "size": "28", "left": "420", "left_chars": "200", "first_line": "420", "first_line_chars": "200"},
        "TOC3": {"line": "400", "size": "24", "left": "840", "left_chars": "400", "first_line": "420", "first_line_chars": "200"},
    }
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / "compact.docx"
        with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/styles.xml":
                    root = etree.fromstring(data)
                    for style_id, token in style_tokens.items():
                        styles = root.xpath(f'.//w:style[@w:styleId="{style_id}"]', namespaces=ns)
                        if not styles:
                            continue
                        style = styles[0]
                        p_pr = upsert_xml_child(style, f"{{{W_NS}}}pPr")
                        spacing = upsert_xml_child(p_pr, f"{{{W_NS}}}spacing")
                        spacing.set(qn("w:lineRule"), "exact")
                        spacing.set(qn("w:line"), token["line"])
                        spacing.set(qn("w:before"), "0")
                        spacing.set(qn("w:after"), "0")
                        ind = upsert_xml_child(p_pr, f"{{{W_NS}}}ind")
                        for attr in (
                            "left",
                            "leftChars",
                            "start",
                            "startChars",
                            "firstLine",
                            "firstLineChars",
                            "hanging",
                            "hangingChars",
                        ):
                            ind.attrib.pop(qn(f"w:{attr}"), None)
                        if token["left"] is not None:
                            ind.set(qn("w:left"), token["left"])
                        if token["left_chars"] is not None:
                            ind.set(qn("w:leftChars"), token["left_chars"])
                        ind.set(qn("w:firstLine"), token["first_line"])
                        ind.set(qn("w:firstLineChars"), token["first_line_chars"])
                        r_pr = upsert_xml_child(style, f"{{{W_NS}}}rPr")
                        fonts = upsert_xml_child(r_pr, f"{{{W_NS}}}rFonts")
                        fonts.set(qn("w:ascii"), "Times New Roman")
                        fonts.set(qn("w:hAnsi"), "Times New Roman")
                        fonts.set(qn("w:eastAsia"), "宋体")
                        fonts.set(qn("w:cs"), "Times New Roman")
                        for tag in ("sz", "szCs"):
                            size = upsert_xml_child(r_pr, f"{{{W_NS}}}{tag}")
                            size.set(qn("w:val"), token["size"])
                        for tag in ("b", "bCs"):
                            node = r_pr.find(f"{{{W_NS}}}{tag}")
                            if node is not None:
                                r_pr.remove(node)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                zout.writestr(item, data)
        shutil.move(str(tmp_path), docx_path)


def add_heading(doc: Document, number: str, title: str, level: int) -> None:
    text = f"{number} {title}"
    style = f"Heading {level}"
    try:
        p = doc.add_paragraph(style=style)
    except KeyError:
        p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level <= 2 else WD_ALIGN_PARAGRAPH.LEFT
    if level == 1:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_run_font(run, 15 if level == 1 else 13 if level == 2 else 12, True)


def add_back_matter_heading(doc: Document, title: str) -> None:
    """Create an unnumbered major heading that is still captured by the Word TOC."""
    try:
        p = doc.add_paragraph(style="Heading 1")
    except KeyError:
        p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(title)
    set_run_font(run, 15, True)


def table_lookup(tables: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {table["id"]: table for table in tables}


def figure_lookup(figures: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {figure["id"]: figure for figure in figures}


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is not None:
        tbl_pr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "12")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "000000")
        borders.append(tag)
    for edge in ("left", "right", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    inside_h = OxmlElement("w:insideH")
    inside_h.set(qn("w:val"), "single")
    inside_h.set(qn("w:sz"), "4")
    inside_h.set(qn("w:color"), "BFBFBF")
    borders.append(inside_h)
    tbl_pr.append(borders)


def add_table(doc: Document, table_data: dict[str, Any]) -> None:
    add_center(doc, table_data["caption"], 11, False)
    rows = table_data["rows"]
    headers = table_data["headers"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    set_table_borders(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        append_text_with_citations(p, header)
        for run in p.runs:
            run.bold = True
            set_run_font(run, 10, True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_format(p, False, 1.15)
            append_text_with_citations(p, value)
            for run in p.runs:
                set_run_font(run, 10)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_figure(doc: Document, figure_data: dict[str, Any]) -> None:
    image_path = ROOT / figure_data["path"]
    if not image_path.exists():
        raise FileNotFoundError(f"缺少图片：{image_path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(float(figure_data.get("width_inches", 5.5))))
    add_center(doc, figure_data["caption"], 11, False)


def add_block(doc: Document, block: dict[str, Any], figures: dict[str, dict[str, Any]], tables: dict[str, dict[str, Any]]) -> None:
    if block["type"] == "paragraph":
        add_body_paragraph(doc, block["text"])
    elif block["type"] == "formula":
        add_formula(doc, block["text"], block["number"], block.get("latex"))
    elif block["type"] == "figure":
        add_figure(doc, figures[block["id"]])
    elif block["type"] == "table":
        add_table(doc, tables[block["id"]])
    else:
        raise ValueError(f"未知正文块类型：{block['type']}")


def add_chapter(doc: Document, node: dict[str, Any], figures: dict[str, dict[str, Any]], tables: dict[str, dict[str, Any]], level: int = 1) -> None:
    add_heading(doc, node["number"], node["title"], level)
    for block in node.get("blocks", []):
        add_block(doc, block, figures, tables)
    for child in node.get("sections", []):
        add_chapter(doc, child, figures, tables, min(level + 1, 3))


def add_front_matter(doc: Document, meta: dict[str, Any], chapters: list[dict[str, Any]], toc_pages: dict[str, Any]) -> None:
    add_center(doc, meta["title_zh"], 16, True)
    doc.add_paragraph()
    add_center(doc, f"专业 {meta['major']}", 12, False)
    add_center(doc, f"学生 {meta['student_name']}  指导老师 {meta['advisor']}", 12, False)
    doc.add_paragraph()
    p = doc.add_paragraph()
    set_paragraph_format(p, True, 1.5)
    label = p.add_run("摘要：")
    set_run_font(label, 12, True)
    body = p.add_run(meta["abstract_zh"])
    set_run_font(body, 12)
    add_keywords(doc, "关键词：", meta["keywords_zh"])

    add_page_break(doc)
    add_center(doc, meta["title_en"], 15, True)
    doc.add_paragraph()
    add_center(doc, f"Major {meta['major']}", 12, False)
    add_center(doc, f"Student {meta['student_name']}  Supervisor {meta['advisor']}", 12, False)
    doc.add_paragraph()
    p = doc.add_paragraph()
    set_paragraph_format(p, True, 1.5)
    label = p.add_run("Abstract. ")
    set_run_font(label, 12, True, "Times New Roman")
    body = p.add_run(meta["abstract_en"])
    set_run_font(body, 12, False, "Times New Roman")
    add_keywords(doc, "Keywords: ", meta["keywords_en"], english=True)

    add_page_break(doc)
    add_toc(doc, chapters, toc_pages)


def add_references(doc: Document, references: list[str]) -> None:
    add_back_matter_heading(doc, "参考文献")
    for idx, ref in enumerate(references, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(-30)
        p.paragraph_format.left_indent = Pt(30)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(20)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(f"[{idx}] {ref}")
        set_run_mixed_font(run, 12, False)


def add_acknowledgement(doc: Document, meta: dict[str, Any]) -> None:
    add_back_matter_heading(doc, "致谢")
    for text in meta["acknowledgement"]:
        add_body_paragraph(doc, text)


def add_ai_statement(doc: Document, meta: dict[str, Any]) -> None:
    add_back_matter_heading(doc, "四川大学本科毕业论文 AI 工具使用声明")
    statement = meta["ai_statement"]
    lines = [
        f"本人声明，在撰写题目为《{meta['title_zh']}》的本科毕业论文过程中，对 AI 工具的使用情况如下：",
        f"一、使用的 AI 工具：{'、'.join(statement['tools'])}。",
        f"二、使用目的：{statement['purpose']}。",
        f"三、原创性与责任声明：{statement['human_review']}本人对论文内容的真实性、准确性和学术规范性承担责任。",
    ]
    for text in lines:
        add_body_paragraph(doc, text)


def strip_comments(docx_path: Path) -> None:
    """Remove template reviewer comments and right-margin annotation parts."""
    comment_parts = {
        "word/comments.xml",
        "word/commentsExtended.xml",
        "word/commentsIds.xml",
        "word/people.xml",
    }
    comment_rel_types = (
        "comments",
        "commentsExtended",
        "commentsIds",
        "people",
    )
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / "clean.docx"
        with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename in comment_parts:
                    continue
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    text = data.decode("utf-8")
                    for part in comment_parts:
                        name = "/" + part
                        text = re.sub(rf"<Override[^>]+PartName=\"{re.escape(name)}\"[^>]*/>", "", text)
                    data = text.encode("utf-8")
                elif item.filename.endswith(".rels"):
                    text = data.decode("utf-8")
                    for rel_type in comment_rel_types:
                        text = re.sub(
                            rf"<Relationship[^>]+Type=\"[^\"]*/{rel_type}\"[^>]*/>",
                            "",
                            text,
                        )
                    data = text.encode("utf-8")
                elif item.filename.startswith("word/") and item.filename.endswith(".xml"):
                    text = data.decode("utf-8")
                    text = re.sub(r"<w:commentRangeStart\b[^>]*/>", "", text)
                    text = re.sub(r"<w:commentRangeEnd\b[^>]*/>", "", text)
                    text = re.sub(r"<w:commentReference\b[^>]*/>", "", text)
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        shutil.move(str(tmp_path), docx_path)


def build() -> Path:
    meta = load_json("meta.json")
    chapters = load_json("chapters.json")
    figures = figure_lookup(load_json("figures.json"))
    tables = table_lookup(load_json("tables.json"))
    references = load_json("references.json")
    toc_pages = load_json("toc_pages.json")
    validate_citations(chapters, references)
    expected_formula_numbers = validate_formula_blocks(chapters)

    if not TEMPLATE.exists():
        raise FileNotFoundError(f"模板不存在：{TEMPLATE}")

    working = ROOT / "_working.docx"
    shutil.copy2(TEMPLATE, working)
    doc = Document(working)

    replace_template_cover(doc, meta)
    update_headers(doc, meta["title_zh"])
    remove_from_anchor(doc, "论文题目")

    # The retained template already ends the authorization page with a section break.
    add_front_matter(doc, meta, chapters, toc_pages)
    for chapter in chapters:
        add_chapter(doc, chapter, figures, tables, 1)
    add_references(doc, references)
    add_acknowledgement(doc, meta)
    if meta.get("include_ai_statement") is True:
        add_ai_statement(doc, meta)

    doc.save(OUTPUT)
    strip_comments(OUTPUT)
    refresh_toc_with_libreoffice(OUTPUT)
    compact_generated_toc_styles(OUTPUT)
    validate_output_formulas(OUTPUT, expected_formula_numbers)
    working.unlink(missing_ok=True)
    return OUTPUT


if __name__ == "__main__":
    out = build()
    print(out)
