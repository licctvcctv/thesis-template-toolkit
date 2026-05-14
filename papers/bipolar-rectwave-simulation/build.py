#!/usr/bin/env python3
"""
基于SPWM的375Hz与900Hz双频高质量发射优化设计论文组装器。

用法:
    cd /Users/a136/vs/45425/thesis_project
    python papers/spwm-dual-frequency/build.py [输出.docx]
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate, InlineImage
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
IMG_DIR = HERE / "images"
TPL = ROOT / "templates" / "hnwlxy" / "template.docx"
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MNS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
APPENDIX_PLACEHOLDER = "__APPENDIX_PLACEHOLDER__"

_pending_tables: list[dict] = []
_pending_formulas: list[dict] = []
CHAPTER_PREFIX_RE = re.compile(r"^(第[一二三四五六七八九十]+章)\s*(.+)$")


def load_json(name: str):
    return json.loads((HERE / name).read_text(encoding="utf-8"))


def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip()


def default_output_path() -> Path:
    title = load_json("meta.json").get("title_zh", "论文")
    return HERE / f"{safe_filename(title)}_论文初稿.docx"


def fit_image_width_mm(img_path: Path, requested_width: float = 120) -> float:
    max_width = 145.0
    max_height = 118.0
    width = min(float(requested_width), max_width)
    with Image.open(img_path) as im:
        px_w, px_h = im.size
    if px_w and px_h:
        height = width * px_h / px_w
        if height > max_height:
            width = max_height * px_w / px_h
    return max(42.0, round(width, 1))


def process_content(items, doc):
    result = []
    for item in items:
        if isinstance(item, str):
            result.append(item)
            continue
        if item.get("type") == "image":
            img = IMG_DIR / item["path"]
            width = fit_image_width_mm(img, item.get("width", 120))
            result.append(InlineImage(doc, str(img), width=Mm(width)))
            if item.get("caption"):
                result.append(item["caption"])
            continue
        if item.get("type") == "formula":
            _pending_formulas.append(item)
            idx = len(_pending_formulas) - 1
            result.append(f"__FORMULA_PLACEHOLDER_{idx}__")
            continue
        if item.get("type") == "table":
            _pending_tables.append(item)
            idx = len(_pending_tables) - 1
            if item.get("caption"):
                result.append(item["caption"])
            result.append(f"__TABLE_PLACEHOLDER_{idx}__")
            continue
        if item.get("text"):
            result.append(item["text"])
    return result


def normalize_chapters(chapters, doc):
    normalized = []
    for ch in chapters:
        ch = dict(ch)
        title = ch.get("title", "")
        match = CHAPTER_PREFIX_RE.match(title)
        if match:
            ch["toc_title"] = title
            # The template Heading 1 style already supplies the automatic
            # chapter number. Feeding only the chapter subject keeps WPS
            # outline/TOC recognition intact and avoids duplicated numbers.
            ch["title"] = " " + match.group(2)
        else:
            ch["toc_title"] = title
        ch["content"] = process_content(ch.get("content", []), doc)
        sections = []
        for sec in ch.get("sections", []):
            sec = dict(sec)
            sec["content"] = process_content(sec.get("content", []), doc)
            subsections = []
            for sub in sec.get("subsections", []):
                sub = dict(sub)
                sub["content"] = process_content(sub.get("content", []), doc)
                subsections.append(sub)
            sec["subsections"] = subsections
            sections.append(sec)
        ch["sections"] = sections
        normalized.append(ch)
    return normalized


def collect_toc(chapters):
    fixed_toc = HERE / "toc_pages.json"
    if fixed_toc.exists():
        return load_json("toc_pages.json")

    entries = [
        ("摘  要", "1", 1),
        ("ABSTRACT", "2", 1),
    ]
    page = 3
    for ch in chapters:
        entries.append((ch.get("toc_title") or ch["title"], str(page), 1))
        page += max(2, len(ch.get("sections", [])) + 1)
        for sec in ch.get("sections", []):
            entries.append((sec["title"], str(page), 2))
            for sub in sec.get("subsections", []):
                entries.append((sub["title"], str(page), 3))
    entries.extend([
        ("参考文献", str(page + 1), 1),
        ("致  谢", str(page + 2), 1),
        ("附  录", str(page + 3), 1),
    ])
    return entries


def build_context(doc):
    meta = load_json("meta.json")
    chapters = []
    for idx in range(1, 8):
        path = HERE / f"ch{idx}.json"
        if path.exists():
            chapters.append(load_json(f"ch{idx}.json"))
    chapters = normalize_chapters(chapters, doc)
    refs = load_json("references.json")
    ack = load_json("acknowledgement.json")
    appendix = load_json("appendix.json")
    appendix_text = APPENDIX_PLACEHOLDER if appendix.get("items") else appendix.get("text", "")
    return {
        **meta,
        "chapters": chapters,
        "references": [{"text": r if re.match(r"^\[[0-9]+\]", r) else f"[{idx}] {r}"} for idx, r in enumerate(refs, 1)],
        "acknowledgement": ack["text"],
        "appendix": appendix_text,
        "_appendix_data": appendix,
        "toc_placeholder": "__TOC_PLACEHOLDER__",
        "toc_entries": collect_toc(chapters),
    }


def set_run_font(run, size=12, bold=None):
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), "宋体")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")


def set_code_run_font(run, size=8):
    run.font.size = Pt(size)
    run.font.name = "Courier New"
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), "宋体")
    r_fonts.set(qn("w:ascii"), "Courier New")
    r_fonts.set(qn("w:hAnsi"), "Courier New")


def clear_paragraph(para):
    for run in para.runs:
        run.text = ""


def clear_paragraph_body(para):
    for child in list(para._p):
        if child.tag != qn("w:pPr"):
            para._p.remove(child)


def replace_paragraph_text(para, text, size=12, bold=None):
    clear_paragraph(para)
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold)


def insert_toc(doc: Document, toc_entries):
    placeholder = None
    for para in doc.paragraphs:
        if "__TOC_PLACEHOLDER__" in (para.text or ""):
            placeholder = para
            break
    if placeholder is None:
        return

    parent = placeholder._p.getparent()
    idx = parent.index(placeholder._p)
    parent.remove(placeholder._p)

    inserted = []
    for offset, (title, page, level) in enumerate(toc_entries):
        p = OxmlElement("w:p")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), "9000")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str({1: 0, 2: 420, 3: 780}.get(level, 0)))
        p_pr.append(ind)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "left")
        p_pr.append(jc)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:line"), "360")
        spacing.set(qn("w:lineRule"), "exact")
        p_pr.append(spacing)
        p.append(p_pr)
        r = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:eastAsia"), "宋体")
        r_fonts.set(qn("w:ascii"), "Times New Roman")
        r_fonts.set(qn("w:hAnsi"), "Times New Roman")
        r_pr.append(r_fonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "21")
        r_pr.append(sz)
        if level == 1:
            b = OxmlElement("w:b")
            b.set(qn("w:val"), "true")
            r_pr.append(b)
        r.append(r_pr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = f"{title}\t{page}"
        r.append(t)
        p.append(r)
        parent.insert(idx + offset, p)
        inserted.append(p)

    # Keep the visible current entries, but wrap them in a real Word/WPS TOC
    # field so the目录 can be refreshed from heading styles instead of being
    # only static text.
    if inserted:
        first_run = inserted[0].find(qn("w:r"))
        last_run = inserted[-1].find(qn("w:r"))
        if first_run is not None and last_run is not None:
            fld_begin = OxmlElement("w:fldChar")
            fld_begin.set(qn("w:fldCharType"), "begin")
            fld_begin.set(qn("w:dirty"), "true")
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = r'TOC \o "1-3" \h \z \u'
            fld_sep = OxmlElement("w:fldChar")
            fld_sep.set(qn("w:fldCharType"), "separate")
            first_run.insert(0, fld_sep)
            first_run.insert(0, instr)
            first_run.insert(0, fld_begin)
            fld_end = OxmlElement("w:fldChar")
            fld_end.set(qn("w:fldCharType"), "end")
            last_run.append(fld_end)


def set_outline_level(para, level: int):
    p_pr = para._p.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def split_keyword_line(para, label: str, value: str, english=False):
    clear_paragraph(para)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.first_line_indent = None
    label_run = para.add_run(label)
    if english:
        set_run_font(label_run, size=12, bold=True)
    else:
        set_run_font(label_run, size=12, bold=True)
        label_run.font.name = "黑体"
        label_run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    value_run = para.add_run(value)
    set_run_font(value_run, size=12, bold=False)


def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.first_child_found_in("w:tcBorders")
    if old is not None:
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        if spec is None:
            continue
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), spec.get("val", "single"))
        border.set(qn("w:sz"), spec.get("sz", "4"))
        border.set(qn("w:color"), "000000")
        border.set(qn("w:space"), "0")
        borders.append(border)
    tc_pr.append(borders)


def set_cell_shading(cell, fill="F7F7F7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:shd"))
    if old is not None:
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_nowrap(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    if tc_pr.find(qn("w:noWrap")) is None:
        no_wrap = OxmlElement("w:noWrap")
        no_wrap.set(qn("w:val"), "true")
        tc_pr.append(no_wrap)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    if not widths:
        return
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue
            cell = row.cells[idx]
            cell.width = Mm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 56.7)))
            tc_w.set(qn("w:type"), "dxa")


def format_table(table, font_size=10.5, nowrap=False):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    line_spacing = max(13, font_size + 6)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if nowrap:
                set_cell_nowrap(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.line_spacing = Pt(line_spacing)
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                for run in para.runs:
                    set_run_font(run, size=font_size, bold=(r_idx == 0))
            set_cell_border(cell, left={"val": "nil"}, right={"val": "nil"})
            if r_idx == 0:
                set_cell_border(
                    cell,
                    top={"val": "single", "sz": "12"},
                    bottom={"val": "single", "sz": "6"},
                    left={"val": "nil"},
                    right={"val": "nil"},
                )
            elif r_idx == len(table.rows) - 1:
                set_cell_border(
                    cell,
                    bottom={"val": "single", "sz": "12"},
                    left={"val": "nil"},
                    right={"val": "nil"},
                )


def insert_table_after(paragraph, table_data):
    headers = table_data["headers"]
    rows = table_data.get("rows", [])
    widths = table_data.get("widths")
    table_width = sum(widths) if widths else table_data.get("width", 150)
    table = paragraph._parent.add_table(rows=1, cols=len(headers), width=Mm(table_width))
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = str(value)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = str(value)
    set_table_widths(table, widths)
    format_table(table, table_data.get("font_size", 10.5), table_data.get("nowrap", False))
    paragraph._p.addnext(table._tbl)


def insert_pending_tables(doc: Document):
    for para in list(doc.paragraphs):
        match = re.search(r"__TABLE_PLACEHOLDER_(\d+)__", para.text or "")
        if not match:
            continue
        idx = int(match.group(1))
        clear_paragraph(para)
        insert_table_after(para, _pending_tables[idx])


def latex_to_omathpara(latex: str):
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        md_path = tmp_dir / "formula.md"
        docx_path = tmp_dir / "formula.docx"
        md_path.write_text(f"$${latex}$$\n", encoding="utf-8")
        subprocess.run(
            ["pandoc", str(md_path), "-o", str(docx_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        with zipfile.ZipFile(docx_path, "r") as zf:
            xml = zf.read("word/document.xml")
        root = ET.fromstring(xml)
        omath_para = root.find(f".//{{{MNS}}}oMathPara")
        if omath_para is None:
            raise RuntimeError(f"公式转换失败: {latex}")
        return parse_xml(ET.tostring(omath_para, encoding="unicode"))


def latex_to_omath(latex: str):
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        md_path = tmp_dir / "formula.md"
        docx_path = tmp_dir / "formula.docx"
        md_path.write_text(f"$${latex}$$\n", encoding="utf-8")
        subprocess.run(
            ["pandoc", str(md_path), "-o", str(docx_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        with zipfile.ZipFile(docx_path, "r") as zf:
            xml = zf.read("word/document.xml")
        root = ET.fromstring(xml)
        omath = root.find(f".//{{{MNS}}}oMath")
        if omath is None:
            raise RuntimeError(f"公式转换失败: {latex}")
        return parse_xml(ET.tostring(omath, encoding="unicode"))


def set_formula_tabs(para):
    p_pr = para._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is not None:
        p_pr.remove(tabs)
    tabs = OxmlElement("w:tabs")
    for val, pos in (("center", "4500"), ("right", "9000")):
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), val)
        tab.set(qn("w:pos"), pos)
        tabs.append(tab)
    p_pr.append(tabs)


def set_omathpara_justification(omath_para, val="center"):
    omath_pr = omath_para.find(f"{{{MNS}}}oMathParaPr")
    if omath_pr is None:
        omath_pr = OxmlElement("m:oMathParaPr")
        omath_para.insert(0, omath_pr)
    jc = omath_pr.find(f"{{{MNS}}}jc")
    if jc is None:
        jc = OxmlElement("m:jc")
        omath_pr.append(jc)
    jc.set(qn("m:val"), val)
    return omath_para


def set_formula_number_tabs(para):
    p_pr = para._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is not None:
        p_pr.remove(tabs)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), "9000")
    tabs.append(tab)
    p_pr.append(tabs)


def paragraph_after(paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def format_formula_paragraph(para, line_spacing=32):
    para.paragraph_format.first_line_indent = None
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = Pt(line_spacing)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY


def insert_formula(para, formula_data):
    number = formula_data.get("number")
    clear_paragraph_body(para)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_formula_paragraph(para)
    para._p.append(set_omathpara_justification(latex_to_omathpara(formula_data["latex"])))

    if number:
        number_para = paragraph_after(para)
        number_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        number_para.paragraph_format.first_line_indent = None
        number_para.paragraph_format.space_before = Pt(0)
        number_para.paragraph_format.space_after = Pt(6)
        number_para.paragraph_format.line_spacing = Pt(12)
        number_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        set_formula_number_tabs(number_para)
        run = number_para.add_run()
        run.add_tab()
        run.add_text(number)
        set_run_font(run, size=12)


def insert_formula_inline(para, formula_data):
    clear_paragraph_body(para)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    format_formula_paragraph(para)
    set_formula_tabs(para)
    first = para.add_run()
    first.add_tab()
    para._p.append(latex_to_omath(formula_data["latex"]))
    number = formula_data.get("number")
    if number:
        run = para.add_run()
        run.add_tab()
        run.add_text(number)
        set_run_font(run, size=12)


def insert_formula_table(para, formula_data):
    table = para._parent.add_table(rows=1, cols=3, width=Mm(145))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_widths(table, [20, 105, 20])
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=70, start=0, bottom=70, end=0)
        set_cell_border(
            cell,
            top={"val": "nil"},
            bottom={"val": "nil"},
            left={"val": "nil"},
            right={"val": "nil"},
        )
        for cell_para in cell.paragraphs:
            clear_paragraph_body(cell_para)
            cell_para.paragraph_format.first_line_indent = None
            cell_para.paragraph_format.space_before = Pt(0)
            cell_para.paragraph_format.space_after = Pt(0)
            cell_para.paragraph_format.line_spacing = Pt(30)
            cell_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY

    formula_para = table.rows[0].cells[1].paragraphs[0]
    formula_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula_para._p.append(latex_to_omath(formula_data["latex"]))

    number = formula_data.get("number")
    if number:
        number_para = table.rows[0].cells[2].paragraphs[0]
        number_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = number_para.add_run(number)
        set_run_font(run, size=12)

    para._p.addnext(table._tbl)
    remove_paragraph(para)


def insert_pending_formulas(doc: Document):
    for para in list(doc.paragraphs):
        match = re.search(r"__FORMULA_PLACEHOLDER_(\d+)__", para.text or "")
        if not match:
            continue
        idx = int(match.group(1))
        insert_formula_inline(para, _pending_formulas[idx])


def insert_paragraph_after(paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def remove_paragraph(paragraph):
    paragraph._p.getparent().remove(paragraph._p)


def format_appendix_paragraph(para, first_line=True):
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.line_spacing = Pt(22)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.first_line_indent = Pt(24) if first_line else None
    for run in para.runs:
        set_run_font(run, size=12)


def insert_appendix_heading(anchor, text):
    para = insert_paragraph_after(anchor)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing = Pt(22)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(text)
    set_run_font(run, size=12, bold=True)
    return para


def insert_appendix_subheading(anchor, text):
    para = insert_paragraph_after(anchor)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.line_spacing = Pt(20)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    para.paragraph_format.space_before = Pt(5)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.first_line_indent = None
    run = para.add_run(text)
    set_run_font(run, size=12, bold=True)
    return para


def insert_appendix_image(anchor, item):
    img = IMG_DIR / item["path"]
    width = fit_image_width_mm(img, item.get("width", 145))
    img_para = insert_paragraph_after(anchor)
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_para.paragraph_format.space_before = Pt(4)
    img_para.paragraph_format.space_after = Pt(4)
    img_para.add_run().add_picture(str(img), width=Mm(width))
    anchor = img_para
    caption = item.get("caption")
    if caption:
        cap_para = insert_paragraph_after(anchor)
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_para.paragraph_format.line_spacing = Pt(18)
        cap_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        cap_para.paragraph_format.space_before = Pt(0)
        cap_para.paragraph_format.space_after = Pt(6)
        run = cap_para.add_run(caption)
        set_run_font(run, size=10.5, bold=True)
        anchor = cap_para
    return anchor


def insert_appendix_code(anchor, item):
    caption = item.get("caption")
    if caption:
        cap_para = insert_paragraph_after(anchor)
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_para.paragraph_format.line_spacing = Pt(18)
        cap_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        cap_para.paragraph_format.space_before = Pt(4)
        cap_para.paragraph_format.space_after = Pt(3)
        run = cap_para.add_run(caption)
        set_run_font(run, size=10.5, bold=True)
        anchor = cap_para

    lines = item.get("lines", [])
    table = anchor._parent.add_table(rows=max(1, len(lines)), cols=1, width=Mm(item.get("width", 150)))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    if not lines:
        lines = [""]
    for idx, line in enumerate(lines):
        cell = table.rows[idx].cells[0]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell)
        set_cell_margins(cell)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "2"},
            bottom={"val": "single", "sz": "2"},
            left={"val": "single", "sz": "2"},
            right={"val": "single", "sz": "2"},
        )
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.line_spacing = Pt(10)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(line)
        set_code_run_font(run, size=item.get("font_size", 8))
    anchor._p.addnext(table._tbl)
    tail = OxmlElement("w:p")
    table._tbl.addnext(tail)
    tail_para = Paragraph(tail, anchor._parent)
    tail_para.paragraph_format.space_after = Pt(3)
    return tail_para


def insert_appendix_code_plain(anchor, item):
    caption = item.get("caption")
    if caption:
        align = item.get("caption_align", "left")
        cap_para = insert_paragraph_after(anchor)
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
        cap_para.paragraph_format.line_spacing = Pt(20)
        cap_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        cap_para.paragraph_format.space_before = Pt(5)
        cap_para.paragraph_format.space_after = Pt(2)
        cap_para.paragraph_format.first_line_indent = None
        run = cap_para.add_run(caption)
        set_run_font(run, size=12, bold=True)
        anchor = cap_para

    for line in item.get("lines", []):
        para = insert_paragraph_after(anchor)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.first_line_indent = None
        para.paragraph_format.line_spacing = Pt(item.get("line_spacing", 13))
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(line)
        set_code_run_font(run, size=item.get("font_size", 9))
        anchor = para
    return anchor


def insert_appendix_content(doc: Document, appendix_data):
    if not appendix_data or not appendix_data.get("items"):
        return
    placeholder = None
    for para in doc.paragraphs:
        if APPENDIX_PLACEHOLDER in (para.text or ""):
            placeholder = para
            break
    if placeholder is None:
        return

    anchor = placeholder
    for item in appendix_data.get("items", []):
        item_type = item.get("type")
        if item_type == "heading":
            anchor = insert_appendix_heading(anchor, item.get("text", ""))
        elif item_type == "subheading":
            anchor = insert_appendix_subheading(anchor, item.get("text", ""))
        elif item_type == "paragraph":
            para = insert_paragraph_after(anchor)
            para.add_run(item.get("text", ""))
            format_appendix_paragraph(para, first_line=item.get("first_line", True))
            anchor = para
        elif item_type == "image":
            anchor = insert_appendix_image(anchor, item)
        elif item_type == "code":
            anchor = insert_appendix_code(anchor, item)
        elif item_type == "code_plain":
            anchor = insert_appendix_code_plain(anchor, item)
    remove_paragraph(placeholder)


def style_paragraphs(doc: Document):
    citation_re = re.compile(r"(\[[0-9,\-—–，、 ]+\])")
    formula_re = re.compile(r"^.+（[0-9]+-[0-9]+）$")
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text.startswith("关键词：") or text.startswith("关键字："):
            label = "关键词：" if text.startswith("关键词：") else "关键字："
            split_keyword_line(para, label, text[len(label):], english=False)
            continue
        if text.startswith("Keywords:"):
            split_keyword_line(para, "Keywords: ", text[len("Keywords:"):].strip(), english=True)
            continue
        if text in {"摘 要", "摘  要", "ABSTRACT", "参考文献", "致 谢", "致  谢", "附 录", "附  录"}:
            set_outline_level(para, 0)
        has_drawing = bool(para._p.findall(f".//{{{WNS}}}drawing"))
        if has_drawing:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if formula_re.match(text) and not text.startswith("图"):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(6)
            for run in para.runs:
                set_run_font(run, size=12)
        if re.match(r"^[图表]\d+-\d+\s+", text):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_run_font(run, size=10.5, bold=True)
        elif text.startswith("参考文献") or text.startswith("致") or text.startswith("附"):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if re.match(r"^\[[0-9]+\]", text):
            try:
                para.style = doc.styles["Normal"]
            except KeyError:
                pass
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.first_line_indent = None
            p_pr = para._p.get_or_add_pPr()
            num_pr = p_pr.find(qn("w:numPr"))
            if num_pr is not None:
                p_pr.remove(num_pr)
            for run in para.runs:
                set_run_font(run, size=10.5)
        if "[" in text and "]" in text and not text.startswith("["):
            runs_text = "".join(run.text for run in para.runs)
            if citation_re.search(runs_text):
                clear_paragraph(para)
                for part in citation_re.split(runs_text):
                    if not part:
                        continue
                    run = para.add_run(part)
                    set_run_font(run, size=12)
                    if citation_re.fullmatch(part):
                        run.font.superscript = True


def post_process(docx_path: Path, toc_entries, appendix_data=None):
    doc = Document(docx_path)
    insert_toc(doc, toc_entries)
    insert_pending_formulas(doc)
    insert_pending_tables(doc)
    style_paragraphs(doc)
    insert_appendix_content(doc, appendix_data)
    doc.save(docx_path)
    patch_update_fields(docx_path)


def patch_update_fields(docx_path: Path):
    with zipfile.ZipFile(docx_path, "r") as zf:
        files = {name: zf.read(name) for name in zf.namelist()}
    xml = files.get(
        "word/settings.xml",
        (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:settings xmlns:w="{WNS}"></w:settings>'
        ).encode("utf-8"),
    ).decode("utf-8")
    if "w:updateFields" not in xml:
        xml = xml.replace("</w:settings>", '<w:updateFields w:val="true"/></w:settings>')
    files["word/settings.xml"] = xml.encode("utf-8")

    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    tmp.replace(docx_path)


def main() -> int:
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else default_output_path()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = DocxTemplate(str(TPL))
    context = build_context(doc)
    doc.render(context)
    doc.save(output)
    post_process(output, context["toc_entries"], context.get("_appendix_data"))
    print(f"完成: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
