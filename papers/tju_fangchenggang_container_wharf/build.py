from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any

from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
CONTENT = ROOT / "content"
TEMPLATE = PROJECT / "templates" / "tianjin_university_2026" / "template.docx"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
CITATION_RE = re.compile(r"\[\[(\d+)\]\]")
CHAPTER_PREFIX_RE = re.compile(r"^第[一二三四五六七八九十]+章\s*")
NUMBER_PREFIX_RE = re.compile(r"^\d+(?:\.\d+)*\s*")
_OMML_CACHE: dict[str, Any] = {}


def load_json(name: str) -> Any:
    return json.loads((CONTENT / name).read_text(encoding="utf-8"))


def p_text(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs)


def clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)


def set_mixed_font(run, size: int = 12, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_para_body(paragraph) -> None:
    pf = paragraph.paragraph_format
    pf.first_line_indent = Pt(24)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def set_outline_level(paragraph, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:outlineLvl")):
        p_pr.remove(old)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


def set_section_header(section, text: str) -> None:
    section.header.is_linked_to_previous = False
    for paragraph in section.header.paragraphs:
        clear_paragraph(paragraph)
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    p_pr = p._p.get_or_add_pPr()
    old_border = p_pr.find(qn("w:pBdr"))
    if old_border is not None:
        p_pr.remove(old_border)
    p_border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    p_border.append(bottom)
    p_pr.append(p_border)
    run = p.add_run(text)
    set_mixed_font(run, 10)


def apply_header_after_cover(doc: Document, text: str) -> None:
    for section in list(doc.sections)[1:]:
        set_section_header(section, text)


def replace_runs(paragraph, text: str) -> None:
    if not paragraph.runs:
        run = paragraph.add_run(text)
        set_mixed_font(run)
        return
    paragraph.runs[0].text = text
    set_mixed_font(paragraph.runs[0], paragraph.runs[0].font.size.pt if paragraph.runs[0].font.size else 12)
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_runs_preserve(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_keyword_line(paragraph, label: str, keywords: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(label)
        paragraph.add_run(keywords)
        return
    paragraph.runs[0].text = label
    if len(paragraph.runs) == 1:
        paragraph.add_run(keywords)
    else:
        paragraph.runs[1].text = keywords
    paragraph.runs[1].bold = False
    for run in paragraph.runs[2:]:
        run.text = ""


def render_template_text(text: str, meta: dict[str, Any]) -> str:
    def repl(match: re.Match[str]) -> str:
        key = match.group(1).strip()
        value = meta.get(key, "")
        if isinstance(value, list):
            return "\n".join(value)
        return str(value)

    return re.sub(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}", repl, text)


def apply_front_matter(doc: Document, meta: dict[str, Any]) -> None:
    for paragraph in doc.paragraphs:
        text = p_text(paragraph).strip()
        compact = re.sub(r"\s+", "", text)
        if "{{ keywords_zh }}" in text:
            replace_keyword_line(paragraph, "关键词：", str(meta["keywords_zh"]))
        elif "{{ keywords_en }}" in text:
            replace_keyword_line(paragraph, "KEY WORDS:", " " + str(meta["keywords_en"]))
        elif "{{" in text:
            replace_runs_preserve(paragraph, render_template_text(text, meta))
        elif text.startswith("题目："):
            replace_runs_preserve(paragraph, f"题目：{meta['title_zh']}")
        elif compact == "学院":
            replace_runs_preserve(paragraph, f"学 院 {meta['college']}")
        elif compact == "专业":
            replace_runs_preserve(paragraph, f"专 业 {meta['major']}")
        elif compact == "年级":
            replace_runs_preserve(paragraph, f"年 级 {meta['grade']}")
        elif compact == "姓名":
            replace_runs_preserve(paragraph, f"姓 名 {meta['student_name']}")
        elif compact == "学号":
            replace_runs_preserve(paragraph, f"学 号 {meta['student_id']}")
        elif compact == "指导教师":
            replace_runs_preserve(paragraph, f"指导教师 {meta['advisor']}")


def remove_body_from_anchor(doc: Document, anchor_text: str = "绪论") -> None:
    body = doc.element.body
    children = list(body.iterchildren())
    start = None
    for idx, child in enumerate(children):
        if child.tag == qn("w:p"):
            text = "".join(t.text or "" for t in child.iter(qn("w:t"))).strip()
            if text == anchor_text:
                start = idx
                break
    if start is None:
        raise RuntimeError(f"未找到正文起始锚点：{anchor_text}")
    for child in children[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def insert_paragraph_before(paragraph, text: str = ""):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph

    p = Paragraph(new_p, paragraph._parent)
    if text:
        p.add_run(text)
    return p


def add_text_with_citations(paragraph, text: str) -> None:
    pos = 0
    for match in CITATION_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_mixed_font(run)
        run = paragraph.add_run(f"[{match.group(1)}]")
        set_mixed_font(run)
        run.font.superscript = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_mixed_font(run)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_mixed_font(run)


def set_body_section_numbering(section) -> None:
    section.footer.is_linked_to_previous = False
    for paragraph in section.footer.paragraphs:
        clear_paragraph(paragraph)
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(p)
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), "1")
    pg_num.set(qn("w:fmt"), "decimal")
    sect_pr.append(pg_num)


def set_front_matter_toc_numbering(section, start: int = 3) -> None:
    section.footer.is_linked_to_previous = False
    for paragraph in section.footer.paragraphs:
        clear_paragraph(paragraph)
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(p)
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), str(start))
    pg_num.set(qn("w:fmt"), "upperRoman")
    sect_pr.append(pg_num)


def add_heading(doc: Document, text: str, level: int, page_break: bool = True) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1 and page_break:
        p.paragraph_format.page_break_before = True
    if level == 1:
        text = CHAPTER_PREFIX_RE.sub("", text)
    elif level in {2, 3}:
        text = NUMBER_PREFIX_RE.sub("", text)
    text = f" {text}"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_mixed_font(run, 15 if level == 1 else 12, True)


def add_body_p(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_para_body(p)
    add_text_with_citations(p, text)


def latex_to_omml(latex: str):
    if latex in _OMML_CACHE:
        return deepcopy(_OMML_CACHE[latex])

    if not shutil.which("pandoc"):
        raise RuntimeError("缺少 pandoc，无法把 LaTeX 公式转换为 Word/WPS 原生 OMML 公式")

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        md_path = tmp_path / "formula.md"
        docx_path = tmp_path / "formula.docx"
        md_path.write_text(f"$$\n{latex}\n$$\n", encoding="utf-8")
        subprocess.run(
            ["pandoc", str(md_path), "-f", "markdown+tex_math_dollars", "-o", str(docx_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        with zipfile.ZipFile(docx_path) as zf:
            xml = zf.read("word/document.xml")
        root = etree.fromstring(xml)
        ns = {"m": MATH_NS}
        omath = root.find(".//m:oMath", ns)
        if omath is None:
            raise RuntimeError(f"公式转换失败: {latex}")
        _OMML_CACHE[latex] = deepcopy(omath)
        return deepcopy(omath)


def add_formula(doc: Document, block: dict[str, str]) -> None:
    if not block.get("latex"):
        raise RuntimeError(f"公式缺少 latex 字段: {block}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.tab_stops.add_tab_stop(Pt(205), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Pt(410), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    p._p.append(latex_to_omml(block["latex"]))
    run = p.add_run("\t" + block["number"])
    set_mixed_font(run)


def add_caption(doc: Document, text: str, keep_with_next: bool = False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.25
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.keep_with_next = keep_with_next
    run = p.add_run(text)
    set_mixed_font(run, 10, False)
    return p


def add_figure(doc: Document, block: dict[str, Any]) -> None:
    image_path = (ROOT / block["path"]).resolve()
    if not image_path.exists():
        raise RuntimeError(f"missing figure image: {image_path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(float(block.get("width_in", 5.8))))
    caption = add_caption(doc, block["caption"])
    caption.paragraph_format.keep_together = True


def add_table_block(doc: Document, block: dict[str, Any]) -> None:
    caption = add_caption(doc, block["caption"], keep_with_next=True)
    if block.get("page_break_before"):
        caption.paragraph_format.page_break_before = True
    headers = block["headers"]
    rows = block["rows"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = block.get("widths_in") or [5.8 / len(headers)] * len(headers)

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, size, val in [
        ("top", "12", "single"),
        ("bottom", "12", "single"),
        ("left", "0", "nil"),
        ("right", "0", "nil"),
        ("insideH", "0", "nil"),
        ("insideV", "0", "nil"),
    ]:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)

    def set_header_bottom(cell) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.first_child_found_in("w:tcBorders")
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "0")
        bottom.set(qn("w:color"), "000000")
        tc_borders.append(bottom)

    def fill_cell(cell, text: Any, bold: bool = False) -> None:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            clear_paragraph(paragraph)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(str(text))
            set_mixed_font(run, 9, bold)

    for idx, text in enumerate(headers):
        table.columns[idx].width = Inches(float(widths[idx]))
        set_header_bottom(table.rows[0].cells[idx])
        fill_cell(table.rows[0].cells[idx], text, True)
    for row_data in rows:
        row = table.add_row()
        for idx, text in enumerate(row_data):
            table.columns[idx].width = Inches(float(widths[idx]))
            fill_cell(row.cells[idx], text)
    doc.add_paragraph()


def render_blocks(doc: Document, blocks: list[dict[str, Any]]) -> None:
    for block in blocks:
        if block["type"] == "p":
            add_body_p(doc, block["text"])
        elif block["type"] == "formula":
            add_formula(doc, block)
        elif block["type"] == "figure":
            add_figure(doc, block)
        elif block["type"] == "table":
            add_table_block(doc, block)
        else:
            raise RuntimeError(f"unknown block type: {block['type']}")


def render_chapters(doc: Document, chapters: dict[str, Any]) -> None:
    for chapter_index, chapter in enumerate(chapters["chapters"]):
        add_heading(doc, chapter["title"], 1, page_break=chapter_index != 0)
        if "blocks" in chapter:
            render_blocks(doc, chapter["blocks"])
        for section in chapter.get("sections", []):
            add_heading(doc, section["title"], 2)
            if "blocks" in section:
                render_blocks(doc, section["blocks"])
            for subsection in section.get("subsections", []):
                add_heading(doc, subsection["title"], 3)
                render_blocks(doc, subsection.get("blocks", []))


def add_references(doc: Document, references: list[str]) -> None:
    p = doc.add_paragraph()
    set_outline_level(p, 0)
    p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("参考文献")
    set_mixed_font(run, 15, True)
    for idx, ref in enumerate(references, start=1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Pt(-18)
        pf.left_indent = Pt(18)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.25
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.keep_together = True
        run = p.add_run(f"[{idx}] {ref}")
        set_mixed_font(run, 10)


def add_ack(doc: Document, meta: dict[str, Any]) -> None:
    p = doc.add_paragraph()
    set_outline_level(p, 0)
    p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("致 谢")
    set_mixed_font(run, 15, True)
    for para in meta.get("acknowledgement", []):
        add_body_p(doc, para)


def collect_toc_entries(chapters: dict[str, Any]) -> list[tuple[int, str]]:
    entries: list[tuple[int, str]] = []
    for chapter in chapters["chapters"]:
        entries.append((1, chapter["title"]))
        for section in chapter.get("sections", []):
            entries.append((2, section["title"]))
            for subsection in section.get("subsections", []):
                entries.append((3, subsection["title"]))
    entries.append((1, "参考文献"))
    entries.append((1, "致 谢"))
    return entries


def rebuild_toc(doc: Document, chapters: dict[str, Any], toc_pages: dict[str, int]) -> None:
    paragraphs = doc.paragraphs
    toc_title = None
    anchor = None
    for paragraph in paragraphs:
        text = p_text(paragraph).strip()
        if re.sub(r"\s+", "", text) == "目录":
            toc_title = paragraph
        if text == "绪论":
            anchor = paragraph
            break
    if toc_title is None or anchor is None:
        return

    deleting = False
    for paragraph in list(doc.paragraphs):
        if paragraph._p is toc_title._p:
            deleting = True
            continue
        if paragraph._p is anchor._p:
            break
        if deleting:
            paragraph._element.getparent().remove(paragraph._element)

    toc_paragraphs = []
    for level, title in collect_toc_entries(chapters):
        page = str(toc_pages.get(title, ""))
        p = insert_paragraph_before(anchor)
        toc_paragraphs.append(p)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        left = {1: 0, 2: 18, 3: 36}[level]
        p.paragraph_format.left_indent = Pt(left)
        p.paragraph_format.tab_stops.add_tab_stop(Pt(410), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(f"{title}\t{page}")
        set_mixed_font(run, 12 if level == 1 else 10, level == 1)

    if toc_paragraphs:
        begin_run = toc_paragraphs[0].insert_paragraph_before().add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        fld_begin.set(qn("w:dirty"), "true")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = r'TOC \o "1-3" \h \z \u'
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        begin_run._r.append(fld_begin)
        begin_run._r.append(instr)
        begin_run._r.append(fld_sep)
        begin_run.font.hidden = True

        end_run = toc_paragraphs[-1].add_run()
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        end_run._r.append(fld_end)
        end_run.font.hidden = True


def set_update_fields(docx_path: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        with zipfile.ZipFile(docx_path) as zf:
            zf.extractall(tmp_path)
        settings = tmp_path / "word" / "settings.xml"
        if settings.exists():
            text = settings.read_text(encoding="utf-8")
            if "w:updateFields" not in text:
                text = text.replace("</w:settings>", '<w:updateFields w:val="true"/></w:settings>')
            else:
                text = re.sub(r'<w:updateFields[^>]*/>', '<w:updateFields w:val="true"/>', text)
            settings.write_text(text, encoding="utf-8")
        rebuilt = docx_path.with_suffix(".tmp.docx")
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as zf:
            for file in tmp_path.rglob("*"):
                if file.is_file():
                    zf.write(file, file.relative_to(tmp_path).as_posix())
        rebuilt.replace(docx_path)


def validate_citations(chapters: dict[str, Any], ref_count: int) -> None:
    seen: list[int] = []

    def walk_blocks(blocks):
        for block in blocks:
            if block.get("type") == "p":
                seen.extend(int(n) for n in CITATION_RE.findall(block["text"]))

    for chapter in chapters["chapters"]:
        walk_blocks(chapter.get("blocks", []))
        for section in chapter.get("sections", []):
            walk_blocks(section.get("blocks", []))
            for subsection in section.get("subsections", []):
                walk_blocks(subsection.get("blocks", []))
    if seen != sorted(seen):
        raise RuntimeError(f"citation order moves backward: {seen}")
    if seen and max(seen) > ref_count:
        raise RuntimeError("citation marker exceeds reference count")
    expected = list(range(1, ref_count + 1))
    if seen != expected:
        raise RuntimeError(f"citation markers must cover references once in order: {seen} != {expected}")


def main() -> None:
    meta = load_json("meta.json")
    chapters = load_json("chapters.json")
    references = load_json("references.json")
    toc_pages_path = CONTENT / "toc_pages.json"
    toc_pages = json.loads(toc_pages_path.read_text(encoding="utf-8")) if toc_pages_path.exists() else {}
    validate_citations(chapters, len(references))

    doc = Document(TEMPLATE)
    apply_front_matter(doc, meta)
    rebuild_toc(doc, chapters, toc_pages)
    remove_body_from_anchor(doc)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    if len(doc.sections) >= 2:
        set_front_matter_toc_numbering(doc.sections[-2])
    set_body_section_numbering(body_section)
    apply_header_after_cover(doc, f"天津大学{meta.get('year', '2026')}届本科生毕业设计")
    render_chapters(doc, chapters)
    add_references(doc, references)
    add_ack(doc, meta)

    output = ROOT / meta["output_name"]
    doc.save(output)
    set_update_fields(output)
    print(output)


if __name__ == "__main__":
    main()
