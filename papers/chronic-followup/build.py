"""Build the revised chronic follow-up thesis DOCX from JSON and assets."""

from __future__ import annotations

import json
import shutil
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
CONTENT = HERE / "content"
TEMPLATE = ROOT / "templates" / "chronic-followup" / "template.docx"
FORMAT_REFERENCE = ROOT / "templates" / "chronic-followup" / "基于Spring Boot的门诊预约挂号系统设计与实现_修改版(1).docx"
DEFAULT_OUTPUT = HERE / "社区卫生服务中心慢性病随访系统设计与实现_论文修改稿.docx"


def load_json(name: str):
    with (CONTENT / name).open("r", encoding="utf-8") as f:
        return json.load(f)


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs).strip()


def compact_text(text: str) -> str:
    return "".join(str(text).split())


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag == qn("w:pPr"):
            continue
        p.remove(child)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    source_rpr = None
    for existing_run in paragraph.runs:
        if existing_run.text:
            source_rpr = deepcopy(existing_run._element.rPr)
            break
    clear_paragraph(paragraph)
    if text:
        run = paragraph.add_run(text)
        if source_rpr is not None:
            run._element.get_or_add_rPr().getparent().replace(run._element.rPr, deepcopy(source_rpr))


def add_run_with_rpr(paragraph: Paragraph, text: str, source_run=None):
    run = paragraph.add_run(text)
    if source_run is not None and source_run._element.rPr is not None:
        run._element.get_or_add_rPr().getparent().replace(run._element.rPr, deepcopy(source_run._element.rPr))
    return run


def set_keyword_paragraph_text(paragraph: Paragraph, text: str) -> None:
    original_runs = list(paragraph.runs)
    clear_paragraph(paragraph)
    if text.startswith("关键词："):
        add_run_with_rpr(paragraph, "关键词", original_runs[0] if len(original_runs) > 0 else None)
        add_run_with_rpr(paragraph, "：", original_runs[1] if len(original_runs) > 1 else None)
        add_run_with_rpr(paragraph, text.removeprefix("关键词："), original_runs[2] if len(original_runs) > 2 else None)
    elif text.startswith("Keywords:"):
        add_run_with_rpr(paragraph, "Keywords", original_runs[0] if len(original_runs) > 0 else None)
        add_run_with_rpr(paragraph, ": ", original_runs[1] if len(original_runs) > 1 else None)
        add_run_with_rpr(paragraph, text.removeprefix("Keywords:").strip(), original_runs[2] if len(original_runs) > 2 else None)
    else:
        set_paragraph_text(paragraph, text)


def apply_run_font(run, size: float | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if size:
        run.font.size = Pt(size)


def apply_code_run_font(run) -> None:
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(8)
    color = run._element.get_or_add_rPr().find(qn("w:color"))
    if color is not None:
        run._element.rPr.remove(color)


def apply_run_typography(run, *, east_asia: str, ascii_font: str = "Times New Roman", size: float = 12, bold: bool | None = None) -> None:
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), east_asia)
    rpr.rFonts.set(qn("w:ascii"), ascii_font)
    rpr.rFonts.set(qn("w:hAnsi"), ascii_font)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def find_paragraph(doc: Document, anchor: str) -> Paragraph | None:
    anchor = anchor.strip()
    compact_anchor = compact_text(anchor)
    for paragraph in doc.paragraphs:
        text = paragraph_text(paragraph)
        if not text:
            continue
        compact = compact_text(text)
        if text == anchor or compact == compact_anchor:
            return paragraph
        if len(anchor) > 8 and anchor in text:
            return paragraph
        if len(text) > 20 and text in anchor:
            return paragraph
        if len(anchor) > 16 and text.startswith(anchor[:16]):
            return paragraph
        if len(text) > 16 and anchor.startswith(text[:16]):
            return paragraph
    return None


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style is not None:
        new_paragraph.style = style
    if text:
        run = new_paragraph.add_run(text)
        apply_run_font(run)
    return new_paragraph


def insert_paragraph_before(paragraph: Paragraph, text: str = "", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style is not None:
        new_paragraph.style = style
    if text:
        run = new_paragraph.add_run(text)
        apply_run_font(run)
    return new_paragraph


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def find_paragraph_index(paragraphs: list[Paragraph], anchor: str, start: int = 0) -> int | None:
    anchor = anchor.strip()
    compact_anchor = compact_text(anchor)
    for i in range(start, len(paragraphs)):
        text = paragraph_text(paragraphs[i])
        if not text:
            continue
        compact = compact_text(text)
        if text == anchor or compact == compact_anchor:
            return i
        if len(anchor) > 8 and anchor in text:
            return i
        if len(text) > 20 and text in anchor:
            return i
        if len(anchor) > 16 and text.startswith(anchor[:16]):
            return i
        if len(text) > 16 and anchor.startswith(text[:16]):
            return i
    return None


def apply_trim_sections(doc: Document, trim_sections: list[dict]) -> list[str]:
    missing = []
    for item in trim_sections:
        paragraphs = list(doc.paragraphs)
        start_idx = find_paragraph_index(paragraphs, item["start_anchor"])
        if start_idx is None:
            missing.append(item["start_anchor"][:60])
            continue
        end_idx = find_paragraph_index(paragraphs, item["end_anchor"], start_idx + 1)
        if end_idx is None or end_idx <= start_idx:
            missing.append(item["end_anchor"][:60])
            continue

        start_para = paragraphs[start_idx]
        style = start_para.style
        for text in item.get("replacement_paragraphs", []):
            new_para = insert_paragraph_before(start_para, text, style=style)
            new_para.paragraph_format.first_line_indent = Mm(7.4)

        remove_end = item.get("remove_end", False)
        end_exclusive = end_idx + 1 if remove_end else end_idx
        for paragraph in paragraphs[start_idx:end_exclusive]:
            remove_paragraph(paragraph)
    return missing


def text_width_twips(doc: Document) -> int:
    section = doc.sections[0]
    try:
        return int(round((section.page_width - section.left_margin - section.right_margin) / 635))
    except Exception:
        return 9070


def set_right_dot_leader_tab(paragraph: Paragraph, pos_twips: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    for tabs in list(ppr.findall(qn("w:tabs"))):
        ppr.remove(tabs)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), str(pos_twips))
    tabs.append(tab)
    ppr.append(tabs)


def set_style_right_dot_leader_tab(style, pos_twips: int) -> None:
    ppr = style._element.pPr
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        style._element.append(ppr)
    for tabs in list(ppr.findall(qn("w:tabs"))):
        ppr.remove(tabs)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), str(pos_twips))
    tabs.append(tab)
    ppr.append(tabs)


def toc_right_dot_leader_tabs(doc: Document) -> dict[int, int]:
    position = text_width_twips(doc)
    for level in (1, 2, 3):
        try:
            style = doc.styles[f"toc {level}"]
        except Exception:
            continue
        set_style_right_dot_leader_tab(style, position)
    return {level: position for level in (1, 2, 3)}


def update_toc_cache(doc: Document, toc_entries: list[dict]) -> list[str]:
    toc_title_idx = None
    body_abstract_idx = None
    paragraphs = list(doc.paragraphs)
    for i, paragraph in enumerate(paragraphs):
        if paragraph_text(paragraph) == "目录":
            toc_title_idx = i
        elif toc_title_idx is not None and paragraph_text(paragraph) == "摘要" and paragraph.style.name == "Title":
            body_abstract_idx = i
            break
    if toc_title_idx is None or body_abstract_idx is None:
        return ["目录/摘要"]

    toc_paragraphs = [
        paragraph
        for paragraph in paragraphs[toc_title_idx + 1 : body_abstract_idx]
        if paragraph.style and paragraph.style.name.lower().startswith("toc")
    ]
    if not toc_paragraphs:
        return ["toc paragraphs"]

    styles_by_level = {}
    for paragraph in toc_paragraphs:
        style_name = paragraph.style.name.lower()
        if style_name.endswith("1"):
            styles_by_level[1] = paragraph.style
        elif style_name.endswith("2"):
            styles_by_level[2] = paragraph.style
        elif style_name.endswith("3"):
            styles_by_level[3] = paragraph.style

    first_toc = toc_paragraphs[0]
    if len(toc_paragraphs) < len(toc_entries):
        current = toc_paragraphs[-1]
        for _ in range(len(toc_entries) - len(toc_paragraphs)):
            current = insert_paragraph_after(current, "")
            toc_paragraphs.append(current)

    toc_right_tabs = toc_right_dot_leader_tabs(doc)
    for paragraph, entry in zip(toc_paragraphs, toc_entries):
        level = int(entry.get("level", 1))
        if level in styles_by_level:
            paragraph.style = styles_by_level[level]
        set_paragraph_text(paragraph, f"{entry['title']}\t{entry['page']}")
        set_right_dot_leader_tab(paragraph, toc_right_tabs.get(level, text_width_twips(doc)))

    for paragraph in toc_paragraphs[len(toc_entries) :]:
        remove_paragraph(paragraph)

    # Keep the first cached TOC paragraph close to the original TOC block.
    first_toc.paragraph_format.space_before = Pt(0)
    return []


def paragraph_has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def remove_figure_near_caption(doc: Document, caption_anchor: str) -> bool:
    paragraphs = list(doc.paragraphs)
    caption_index = find_paragraph_index(paragraphs, caption_anchor)
    if caption_index is None:
        return False

    caption_para = paragraphs[caption_index]
    to_remove = [caption_para]
    search_start = max(-1, caption_index - 7)
    for j in range(caption_index - 1, search_start, -1):
        text = paragraph_text(paragraphs[j])
        if paragraph_has_drawing(paragraphs[j]) or not text:
            to_remove.append(paragraphs[j])
            continue
        if to_remove:
            break

    seen = set()
    for paragraph in to_remove:
        key = id(paragraph._p)
        if key in seen:
            continue
        seen.add(key)
        remove_paragraph(paragraph)
    return True


def replace_image_near_caption(doc: Document, caption_anchor: str, image_path: Path, width_mm: float, caption_text: str) -> bool:
    paragraphs = list(doc.paragraphs)
    caption_para = None
    caption_index = -1
    compact_anchor = compact_text(caption_anchor)
    for i, paragraph in enumerate(paragraphs):
        text = paragraph_text(paragraph)
        compact = compact_text(text)
        if text == caption_anchor or caption_anchor in text or compact == compact_anchor:
            caption_para = paragraph
            caption_index = i
            break
    if caption_para is None:
        return False

    target_para = None
    blank_candidate = None
    search_start = max(-1, caption_index - 7)
    drawing_candidates = []
    for j in range(caption_index - 1, search_start, -1):
        text = paragraph_text(paragraphs[j])
        if paragraph_has_drawing(paragraphs[j]):
            drawing_candidates.append(paragraphs[j])
            if target_para is None:
                target_para = paragraphs[j]
            continue
        if blank_candidate is None and not text:
            blank_candidate = paragraphs[j]
            continue
        if text:
            break
    if target_para is None:
        target_para = blank_candidate
    if target_para is None:
        target_para = insert_paragraph_after(paragraphs[caption_index - 1])

    clear_paragraph(target_para)
    target_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = target_para.add_run()
    run.add_picture(str(image_path), width=Mm(width_mm))

    for paragraph in drawing_candidates[1:]:
        clear_paragraph(paragraph)

    set_paragraph_text(caption_para, caption_text)
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return True


def insert_figure_after(doc: Document, after_anchor: str, intro: str, image_path: Path, caption: str, width_mm: float) -> bool:
    anchor_para = find_paragraph(doc, after_anchor)
    if anchor_para is None:
        return False

    current = anchor_para
    if intro:
        current = insert_paragraph_after(current, intro, style=anchor_para.style)
        current.paragraph_format.first_line_indent = Mm(7.4)

    image_para = insert_paragraph_after(current)
    image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_run = image_para.add_run()
    image_run.add_picture(str(image_path), width=Mm(width_mm))

    caption_para = insert_paragraph_after(image_para, caption)
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return True


def insert_code_blocks(doc: Document, code_blocks: list[dict]) -> list[str]:
    missing = []
    for item in code_blocks:
        anchor_para = find_paragraph(doc, item["after_anchor"])
        if anchor_para is None:
            missing.append(item.get("caption", item["after_anchor"][:40]))
            continue

        current = anchor_para
        intros = item.get("paragraphs", item.get("intro", ""))
        if isinstance(intros, str):
            intros = [intros] if intros else []
        for intro in intros:
            intro_para = insert_paragraph_after(current, intro, style=anchor_para.style)
            intro_para.paragraph_format.first_line_indent = Mm(7.4)
            current = intro_para

        if item.get("type") == "note":
            continue

        caption_para = insert_paragraph_after(current, item["caption"])
        caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        current = caption_para

        code_value = item.get("code", "")
        lines = code_value if isinstance(code_value, list) else str(code_value).splitlines()

        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Mm(145)
        table._tbl.getparent().remove(table._tbl)
        current._p.addnext(table._tbl)

        cell = table.cell(0, 0)
        cell.width = Mm(145)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "4", "color": "000000"},
            bottom={"val": "single", "sz": "4", "color": "000000"},
            left={"val": "single", "sz": "4", "color": "000000"},
            right={"val": "single", "sz": "4", "color": "000000"},
        )
        tc_pr = cell._tc.get_or_add_tcPr()
        shade = tc_pr.find(qn("w:shd"))
        if shade is not None:
            tc_pr.remove(shade)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Mm(1.5)
        p.paragraph_format.right_indent = Mm(1.5)
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        for line_index, line in enumerate(lines):
            if line_index:
                p.add_run().add_break()
            run = p.add_run(line if line else " ")
            apply_code_run_font(run)
    return missing


def retitle_system_test_section(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    chapter_idx = next((i for i, p in enumerate(paragraphs) if paragraph_text(p) == "系统测试" and p.style.name == "Heading 1"), None)
    if chapter_idx is None:
        return

    end_idx = next(
        (
            i
            for i, p in enumerate(paragraphs[chapter_idx + 1 :], start=chapter_idx + 1)
            if paragraph_text(p) and p.style.name == "Heading 1"
        ),
        len(paragraphs),
    )
    section_paragraphs = paragraphs[chapter_idx + 1 : end_idx]
    intro = next((p for p in section_paragraphs if paragraph_text(p)), None)
    has_purpose = any(paragraph_text(p) == "测试目的" for p in section_paragraphs)
    if intro is not None and not has_purpose:
        purpose_heading = insert_paragraph_after(intro, "测试目的", style=doc.styles["Heading 2"])
        purpose_body = insert_paragraph_after(
            purpose_heading,
            "本次测试以本地部署后的系统为对象，检查管理员端和社区医生端能否完成日常随访业务。测试重点包括医生与患者资料维护、患者风险等级设置、诊疗记录录入、随访任务分配、随访记录保存以及统计页面展示等内容。",
            style=doc.styles["Normal"],
        )
        purpose_body.paragraph_format.first_line_indent = Mm(7.4)

    for paragraph in section_paragraphs:
        text = paragraph_text(paragraph)
        if text in {"测试目标", "测试计划"}:
            set_paragraph_text(paragraph, "测试设计过程")
            black_box = insert_paragraph_after(
                paragraph,
                "黑盒测试从页面操作入手。测试人员按照实际使用顺序填写、查询、修改和删除数据，记录页面提示、列表变化和状态变化，并将这些结果与业务需求进行对照。医生管理、患者管理、药品管理、诊疗记录、随访任务和随访记录均选取常用场景进行测试。",
                style=doc.styles["Normal"],
            )
            black_box.paragraph_format.first_line_indent = Mm(7.4)
            white_box = insert_paragraph_after(
                black_box,
                "白盒检查主要配合查看浏览器控制台、后端运行日志和数据库记录，确认保存或更新操作是否进入相应数据表，异常输入是否被拦截，关键状态字段是否按流程发生变化。",
                style=doc.styles["Normal"],
            )
            white_box.paragraph_format.first_line_indent = Mm(7.4)
        elif text == "本章小结":
            set_paragraph_text(paragraph, "测试结论")


def insert_object_model_section(doc: Document) -> list[str]:
    anchor_text = "系统统计分析模块面向管理员，主要对患者数量、疾病类型、风险等级、随访任务状态、随访完成率和医生工作量进行汇总展示，使管理者能够快速了解慢病随访工作进展。"
    anchor = find_paragraph(doc, anchor_text)
    if anchor is None:
        return ["对象模型设计"]

    heading = insert_paragraph_after(anchor, "对象模型设计", style=doc.styles["Heading 2"])
    current = heading
    paragraphs = [
        "对象模型用于说明系统中主要业务对象的属性、职责和关联关系。结合源码中的实体类与数据库表，系统可抽象出User、Doctor、Patient、Medicine、MedicalRecord、FollowupTask、FollowupRecord、RiskAssessment等对象，其中Patient对象是随访业务的核心，Doctor对象承担诊疗与随访执行职责，FollowupTask对象用于连接患者、医生与随访记录。",
        "从对象协作关系看，管理员主要维护用户、医生、患者、药品和随访任务等对象，社区医生主要围绕患者、随访任务和随访记录对象完成业务处理。各对象之间通过用户编号、医生编号、患者编号和任务编号建立关联，既对应后端实体类关系，也为数据库表设计提供依据。系统对象模型如图4-3所示。",
    ]
    for text in paragraphs:
        current = insert_paragraph_after(current, text, style=doc.styles["Normal"])
        current.paragraph_format.first_line_indent = Mm(7.4)
    return []


def rewrite_function_test_notes(doc: Document) -> list[str]:
    replacements = {
        "（1）医生管理功能测试": "医生管理测试主要检查医生资料维护、账号状态变更和列表反馈情况，重点核对保存后的医生记录是否能继续用于患者分配和随访任务下达，测试情况见表6-2。",
        "（2）患者管理功能测试": "患者管理测试围绕患者建档、信息修改、风险等级设置和导入导出展开，重点查看保存后的患者档案、风险状态和负责医生是否在页面与数据库中保持一致，测试情况见表6-3。",
        "（3）药品管理功能测试": "药品管理测试主要检查药品基础资料的新增、修改、停用和查询结果，重点确认药品编码、库存、状态等字段保存后能在后续诊疗记录中正常选择，测试情况见表6-4。",
        "（4）诊疗记录录入功能测试": "诊疗记录录入测试主要检查患者就诊信息和用药明细的保存情况，重点核对患者、医生、诊疗日期、诊断结果等字段是否能够按业务要求写入并回显，测试情况见表6-5。",
        "（5）随访任务分配功能测试": "随访任务分配测试主要检查任务创建、计划日期、负责医生和任务状态流转情况，重点查看任务保存后医生端是否能够查询到对应待办任务，测试情况见表6-6。",
        "（6）随访记录功能测试": "随访记录测试主要检查医生填写随访结果后的保存、查询和任务状态更新情况，重点核对血压、血糖、症状描述和下次随访日期等字段是否能够形成完整记录，测试情况见表6-7。",
    }
    missing = []
    label_prefixes = ("用例名称：", "用例描述：", "详细用例：")
    for heading_text, replacement in replacements.items():
        paragraphs = list(doc.paragraphs)
        heading = find_paragraph(doc, heading_text)
        if heading is None:
            missing.append(heading_text)
            continue
        idx = next((i for i, p in enumerate(paragraphs) if p._p is heading._p), None)
        if idx is None:
            missing.append(heading_text)
            continue
        to_remove = []
        for paragraph in paragraphs[idx + 1 : idx + 5]:
            if paragraph_text(paragraph).startswith(label_prefixes):
                to_remove.append(paragraph)
            elif paragraph_text(paragraph):
                break
        for paragraph in to_remove:
            remove_paragraph(paragraph)
        note = insert_paragraph_after(heading, replacement, style=doc.styles["Normal"])
        note.paragraph_format.first_line_indent = Mm(7.4)
    return missing


def apply_paragraph_replacements(doc: Document, replacements: list[dict]) -> list[str]:
    missing = []
    for item in replacements:
        paragraph = find_paragraph(doc, item["anchor"])
        if paragraph is None:
            missing.append(item["anchor"][:60])
            continue
        replacement = item["replacement"]
        if replacement.startswith("关键词：") or replacement.startswith("Keywords:"):
            set_keyword_paragraph_text(paragraph, replacement)
        else:
            set_paragraph_text(paragraph, replacement)
    return missing


def set_cell_text(cell, text: str) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    apply_run_font(run, 9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def ensure_table_shape(table, rows: int, cols: int) -> None:
    while len(table.columns) < cols:
        table.add_column(Mm(24))
    while len(table.rows) < rows:
        table.add_row()
    while len(table.rows) > rows:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)


def update_table_from_config(doc: Document, table_config: dict) -> None:
    table = doc.tables[table_config["table_index"]]
    rows = [table_config["headers"]] + table_config["rows"]
    ensure_table_shape(table, len(rows), len(rows[0]))
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            set_cell_text(table.rows[r_idx].cells[c_idx], value)
            if r_idx == 0:
                for run in table.rows[r_idx].cells[c_idx].paragraphs[0].runs:
                    run.bold = True


def update_database_tables(doc: Document, tables_config: list[dict]) -> None:
    for table_config in tables_config:
        update_table_from_config(doc, table_config)


def remove_template_artifacts(doc: Document) -> None:
    final_header_rid = None
    for paragraph in list(doc.paragraphs):
        text = paragraph_text(paragraph)
        if text.startswith("千万不要删除行尾的分节符") and "更新整个目录" in text:
            header_refs = paragraph._p.xpath(".//w:sectPr/w:headerReference[@w:type='default']")
            if header_refs:
                final_header_rid = header_refs[0].get(qn("r:id"))
        if text.startswith("千万不要删除行尾的分节符") or text == "/":
            paragraph._p.getparent().remove(paragraph._p)
    if final_header_rid:
        section = doc.element.body.sectPr
        if section is not None:
            header_refs = section.xpath("./w:headerReference[@w:type='default']")
            if header_refs:
                header_refs[0].set(qn("r:id"), final_header_rid)
            else:
                header_ref = OxmlElement("w:headerReference")
                header_ref.set(qn("w:type"), "default")
                header_ref.set(qn("r:id"), final_header_rid)
                section.insert(0, header_ref)
            for title_page in section.xpath("./w:titlePg"):
                section.remove(title_page)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in kwargs[edge].items():
                element.set(qn(f"w:{key}"), str(value))


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def format_tables(doc: Document) -> None:
    for table_index, table in enumerate(doc.tables):
        if table_index == 0:
            continue
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if table.rows:
            repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        apply_run_font(run, 9.5)
                        run.font.highlight_color = None
                        if row_index == 0:
                            run.bold = True

                nil = {"val": "nil", "sz": "0", "color": "FFFFFF"}
                set_cell_border(cell, top=nil, bottom=nil, left=nil, right=nil)

        if not table.rows:
            continue
        first = table.rows[0]
        last = table.rows[-1]
        thick = {"val": "single", "sz": "12", "color": "000000"}
        mid = {"val": "single", "sz": "8", "color": "000000"}
        for cell in first.cells:
            set_cell_border(cell, top=thick, bottom=mid, left={"val": "nil"}, right={"val": "nil"})
        for cell in last.cells:
            set_cell_border(cell, bottom=thick, left={"val": "nil"}, right={"val": "nil"})


def normalize_body_fonts(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.highlight_color = None


def copy_reference_toc_styles(doc: Document) -> None:
    if not FORMAT_REFERENCE.exists():
        return
    reference = Document(FORMAT_REFERENCE)
    for style_name in ("TOC Title", "toc 1", "toc 2", "toc 3"):
        try:
            source_style = reference.styles[style_name]._element
            target_style = doc.styles[style_name]._element
        except Exception:
            continue
        target_style.getparent().replace(target_style, deepcopy(source_style))


def set_single_run(paragraph: Paragraph, text: str, *, east_asia: str, size: float, bold: bool, ascii_font: str = "Times New Roman") -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    apply_run_typography(run, east_asia=east_asia, ascii_font=ascii_font, size=size, bold=bold)


def format_front_body_paragraph(paragraph: Paragraph, *, keyword: bool = False) -> None:
    try:
        paragraph.style = paragraph.part.document.styles["Normal"]
    except Exception:
        pass
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(12) if keyword else Pt(0)
    fmt.first_line_indent = None if keyword else Mm(11.3)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in paragraph.runs:
        if not run.text:
            continue
        apply_run_typography(run, east_asia="宋体", size=12, bold=False)


def format_keyword_line(paragraph: Paragraph) -> None:
    text = paragraph_text(paragraph)
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(12)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = None
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    clear_paragraph(paragraph)
    if text.startswith("关键词："):
        label = paragraph.add_run("关键词：")
        apply_run_typography(label, east_asia="黑体", size=12, bold=True)
        content = paragraph.add_run(text.removeprefix("关键词："))
        apply_run_typography(content, east_asia="宋体", size=12, bold=False)
    elif text.startswith("Keywords:"):
        label = paragraph.add_run("Keywords: ")
        apply_run_typography(label, east_asia="宋体", size=12, bold=True)
        content = paragraph.add_run(text.removeprefix("Keywords:").strip())
        apply_run_typography(content, east_asia="宋体", size=12, bold=False)


def set_paragraph_marker_typography(paragraph: Paragraph, *, east_asia: str, ascii_font: str, size: float, bold: bool) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    for tag, value in (("w:sz", str(int(size * 2))), ("w:szCs", str(int(size * 2)))):
        element = rpr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            rpr.append(element)
        element.set(qn("w:val"), value)
    bold_element = rpr.find(qn("w:b"))
    if bold_element is None:
        bold_element = OxmlElement("w:b")
        rpr.append(bold_element)
    bold_element.set(qn("w:val"), "1" if bold else "0")


def format_body_heading_runs(doc: Document) -> None:
    heading_sizes = {"2": 16, "5": 14, "6": 14}
    for paragraph in doc.paragraphs:
        ppr = paragraph._p.pPr
        pstyle = None
        if ppr is not None and ppr.pStyle is not None:
            pstyle = ppr.pStyle.val
        size = heading_sizes.get(str(pstyle))
        if size is None and paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            size = {"Heading 1": 16, "Heading 2": 14, "Heading 3": 14}[paragraph.style.name]
        if size is None:
            continue
        set_paragraph_marker_typography(paragraph, east_asia="黑体", ascii_font="Times New Roman", size=size, bold=True)
        for run in paragraph.runs:
            if run.text:
                apply_run_typography(run, east_asia="黑体", size=size, bold=True)


def apply_reference_front_matter_format(doc: Document) -> list[str]:
    copy_reference_toc_styles(doc)
    missing: list[str] = []
    paragraphs = list(doc.paragraphs)

    toc_idx = next((i for i, p in enumerate(paragraphs) if compact_text(paragraph_text(p)) in {"目录", "目　　录"}), None)
    abstract_idx = None
    if toc_idx is not None:
        abstract_idx = next(
            (
                i
                for i, p in enumerate(paragraphs[toc_idx + 1 :], start=toc_idx + 1)
                if compact_text(paragraph_text(p)) in {"摘要", "摘　要"} and not p.style.name.lower().startswith("toc")
            ),
            None,
        )
    if toc_idx is None:
        missing.append("目录标题")
    else:
        paragraphs[toc_idx].style = doc.styles["TOC Title"] if "TOC Title" in [s.name for s in doc.styles] else paragraphs[toc_idx].style
        set_single_run(paragraphs[toc_idx], "目　　录", east_asia="黑体", size=18, bold=True)
        paragraphs[toc_idx].alignment = WD_ALIGN_PARAGRAPH.CENTER

    if toc_idx is not None and abstract_idx is not None:
        toc_right_tabs = toc_right_dot_leader_tabs(doc)
        for paragraph in paragraphs[toc_idx + 1 : abstract_idx]:
            text = paragraph_text(paragraph)
            if "\t" not in text:
                continue
            level = 1
            toc_title = text.rsplit("\t", 1)[0].strip()
            parts = toc_title.split(" ", 1)[0].split(".")
            if len(parts) >= 3 and all(part.isdigit() for part in parts[:3]):
                level = 3
            elif len(parts) >= 2 and all(part.isdigit() for part in parts[:2]):
                level = 2
            try:
                paragraph.style = doc.styles[f"toc {level}"]
            except Exception:
                pass
            if "\t" in text:
                title, page = text.rsplit("\t", 1)
                if compact_text(title) in {"摘要", "摘　　要"}:
                    text = f"摘　　要\t{page}"
            set_single_run(paragraph, text, east_asia="宋体", size=12, bold=(level == 1))
            set_right_dot_leader_tab(paragraph, toc_right_tabs.get(level, text_width_twips(doc)))

    if abstract_idx is None:
        missing.append("中文摘要")
        return missing

    set_single_run(paragraphs[abstract_idx], "摘　要", east_asia="黑体", size=16, bold=True)
    paragraphs[abstract_idx].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphs[abstract_idx].paragraph_format.space_after = Pt(5)

    keyword_idx = next((i for i, p in enumerate(paragraphs[abstract_idx + 1 :], start=abstract_idx + 1) if paragraph_text(p).startswith("关键词")), None)
    english_idx = None
    if keyword_idx is not None:
        english_idx = next(
            (
                i
                for i, p in enumerate(paragraphs[keyword_idx + 1 :], start=keyword_idx + 1)
                if paragraph_text(p) == "Abstract"
            ),
            None,
        )
    if keyword_idx is None:
        missing.append("关键词")
    else:
        for paragraph in paragraphs[abstract_idx + 1 : keyword_idx]:
            if paragraph_text(paragraph):
                format_front_body_paragraph(paragraph)
        format_keyword_line(paragraphs[keyword_idx])

    if english_idx is None:
        missing.append("英文摘要")
        return missing

    set_single_run(paragraphs[english_idx], "Abstract", east_asia="黑体", size=16, bold=True)
    paragraphs[english_idx].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphs[english_idx].paragraph_format.space_after = Pt(5)

    english_keyword_idx = next((i for i, p in enumerate(paragraphs[english_idx + 1 :], start=english_idx + 1) if paragraph_text(p).startswith("Keywords")), None)
    if english_keyword_idx is None:
        missing.append("Keywords")
    else:
        for paragraph in paragraphs[english_idx + 1 : english_keyword_idx]:
            if paragraph_text(paragraph):
                format_front_body_paragraph(paragraph)
        format_keyword_line(paragraphs[english_keyword_idx])
    return missing


def set_header_text(section, text: str) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    apply_run_typography(run, east_asia="宋体", size=10.5, bold=False)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def ensure_reference_section_split(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    reference_idx = next(
        (
            i
            for i, paragraph in enumerate(paragraphs)
            if paragraph_text(paragraph) == "参考文献" and paragraph.style.name == "Title"
        ),
        None,
    )
    if reference_idx is None or reference_idx == 0:
        return

    previous = paragraphs[reference_idx - 1]
    ppr = previous._p.get_or_add_pPr()
    if ppr.find(qn("w:sectPr")) is not None:
        return

    source_sect_pr = None
    for paragraph in paragraphs[reference_idx:]:
        sects = paragraph._p.xpath("./w:pPr/w:sectPr")
        if sects:
            source_sect_pr = sects[0]
            break
    if source_sect_pr is None:
        source_sect_pr = doc.element.body.sectPr
    if source_sect_pr is None:
        return

    new_sect_pr = deepcopy(source_sect_pr)
    for header_ref in new_sect_pr.xpath("./w:headerReference"):
        new_sect_pr.remove(header_ref)
    ppr.append(new_sect_pr)


def normalize_front_headers(doc: Document) -> None:
    ensure_reference_section_split(doc)
    labels = {
        0: "目录",
        1: "摘要",
        2: "Abstract",
        3: "前言",
        4: "相关技术与理论",
        5: "系统需求分析",
        6: "系统设计",
        7: "系统实现",
        8: "系统测试",
        9: "总结与展望",
        10: "参考文献",
        11: "致谢",
    }
    for index, label in labels.items():
        if index < len(doc.sections):
            set_header_text(doc.sections[index], label)


def set_update_fields(docx_path: Path) -> None:
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                from lxml import etree

                root = etree.fromstring(data)
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                update = root.find("w:updateFields", ns)
                if update is None:
                    update = etree.Element(qn("w:updateFields"))
                    root.insert(0, update)
                update.set(qn("w:val"), "true")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    tmp.replace(docx_path)


def build(output: Path = DEFAULT_OUTPUT) -> Path:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEMPLATE, output)

    doc = Document(output)
    figures = load_json("figures.json")
    replacements = load_json("paragraph_replacements.json")
    tables = load_json("tables.json")
    trim_sections = load_json("trim_sections.json")
    toc_entries = load_json("toc_pages.json")
    code_blocks = load_json("code_blocks.json") if (CONTENT / "code_blocks.json").exists() else []

    missing_removals = []
    for item in figures.get("removals", []):
        ok = remove_figure_near_caption(doc, item["caption_anchor"])
        if not ok:
            missing_removals.append(item["caption_anchor"])

    missing_figures = []
    for item in figures["replacements"]:
        image = HERE / item["image"]
        ok = replace_image_near_caption(doc, item["caption_anchor"], image, item["width_mm"], item["caption_text"])
        if not ok:
            missing_figures.append(item["caption_anchor"])

    missing_replacements = apply_paragraph_replacements(doc, replacements)
    missing_trims = apply_trim_sections(doc, trim_sections)
    update_database_tables(doc, tables.get("database_tables", []))
    update_table_from_config(doc, tables["test_environment"])
    update_database_tables(doc, tables.get("function_test_tables", []))
    remove_template_artifacts(doc)
    format_tables(doc)
    normalize_body_fonts(doc)
    missing_code_blocks = insert_code_blocks(doc, code_blocks)
    retitle_system_test_section(doc)
    missing_object_model = insert_object_model_section(doc)
    missing_test_rewrites = rewrite_function_test_notes(doc)

    missing_insertions = []
    for item in figures["insertions"]:
        image = HERE / item["image"]
        ok = insert_figure_after(doc, item["after_anchor"], item["intro"], image, item["caption"], item["width_mm"])
        if not ok:
            missing_insertions.append(item["caption"] or item["after_anchor"][:40])

    missing_toc = update_toc_cache(doc, toc_entries)
    missing_reference_format = apply_reference_front_matter_format(doc)
    format_body_heading_runs(doc)
    normalize_front_headers(doc)

    doc.save(output)
    set_update_fields(output)

    diagnostics = {
        "missing_removals": missing_removals,
        "missing_figures": missing_figures,
        "missing_insertions": missing_insertions,
        "missing_code_blocks": missing_code_blocks,
        "missing_object_model": missing_object_model,
        "missing_test_rewrites": missing_test_rewrites,
        "missing_replacements": missing_replacements,
        "missing_trims": missing_trims,
        "missing_toc": missing_toc,
        "missing_reference_format": missing_reference_format,
        "output": str(output),
    }
    (HERE / "data" / "build_diagnostics.json").write_text(
        json.dumps(diagnostics, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return output


if __name__ == "__main__":
    out = build()
    print(out)
