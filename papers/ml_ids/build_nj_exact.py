#!/usr/bin/env python3
"""Assemble the paper on the real Nanjing vocational university DOCX shell.

This builder uses ``templates/ml_ids_source/source.docx`` as the base document,
keeps its cover/front matter/header/footer/styles, and replaces only the paper
content regions with the current JSON-driven thesis content.
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
IMG_DIR = HERE / "images"
SOURCE_DOCX = ROOT / "templates" / "ml_ids_source" / "source.docx"
PANDOC = Path("/Users/a136/.homebrew/bin/pandoc")
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MNS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def load_json(name: str):
    return json.loads((HERE / name).read_text(encoding="utf-8"))


def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip()


def default_output_path() -> Path:
    title = load_json("meta.json").get("title_zh", "论文")
    return HERE / f"{safe_filename(title)}_论文初稿.docx"


def clear_paragraph(para: Paragraph):
    for child in list(para._p):
        if child.tag != qn("w:pPr"):
            para._p.remove(child)


def set_run_font(run, size=12, bold=None, east_asia="宋体", latin="Times New Roman"):
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)


def set_code_run_font(run, size=9):
    set_run_font(run, size=size, east_asia="宋体", latin="Courier New")


def replace_paragraph_text(para: Paragraph, text: str, size=None, bold=None):
    clear_paragraph(para)
    run = para.add_run(text)
    if size is not None or bold is not None:
        set_run_font(run, size=size or 12, bold=bold)
    return para


def insert_paragraph_after(anchor: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if style:
        try:
            para.style = style
        except KeyError:
            pass
    if text:
        para.add_run(text)
    return para


def set_dot_leader_tab(para: Paragraph, pos=9000):
    p_pr = para._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        p_pr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), str(pos))
    tabs.append(tab)


def insert_index_entry_after(anchor: Paragraph, title: str, page: str, style: str) -> Paragraph:
    para = insert_paragraph_after(anchor, "", style)
    set_dot_leader_tab(para)
    run = para.add_run(title + ("\t" + page if page else ""))
    set_run_font(run, size=10.5)
    return para


def add_complex_field_begin(para: Paragraph, instruction: str):
    run = para.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    run = para.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)

    run = para.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)


def insert_complex_field_end_after(anchor: Paragraph) -> Paragraph:
    para = insert_paragraph_after(anchor, "")
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)
    return para


def remove_paragraph(para: Paragraph):
    parent = para._p.getparent()
    if parent is not None:
        parent.remove(para._p)


def paragraph_text(element) -> str:
    return "".join(element.xpath(".//w:t/text()"))


def find_paragraph(doc: Document, text: str, start=0, style: str | None = None) -> Paragraph:
    for para in doc.paragraphs[start:]:
        if (para.text or "").strip() == text and (style is None or para.style.name == style):
            return para
    raise ValueError(f"paragraph not found: {text}")


def remove_between(anchor: Paragraph, stop: Paragraph):
    node = anchor._p.getnext()
    while node is not None and node is not stop._p:
        nxt = node.getnext()
        parent = node.getparent()
        has_section_break = node.tag == qn("w:p") and bool(node.findall(".//w:sectPr", namespaces=node.nsmap))
        if parent is not None and not has_section_break:
            parent.remove(node)
        node = nxt


def remove_after_until_sectpr(anchor: Paragraph):
    node = anchor._p.getnext()
    while node is not None and node.tag != qn("w:sectPr"):
        nxt = node.getnext()
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)
        node = nxt


def find_first_body_heading(doc: Document) -> Paragraph:
    seen_fig_list = False
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text == "图表清单":
            seen_fig_list = True
            continue
        if seen_fig_list and para.style.name == "Heading 1" and text:
            return para
    raise ValueError("body Heading 1 not found")


def strip_chapter_title(title: str) -> str:
    match = re.match(r"^第[一二三四五六七八九十]+章\s*(.+)$", title.strip())
    return match.group(1) if match else title.strip()


def strip_numbered_title(title: str) -> str:
    return re.sub(r"^[0-9]+(?:\.[0-9]+)*\s*", "", title.strip())


def normalize_body_text(text: str) -> str:
    text = text.replace("图表", "表")
    text = re.sub(r"图\s*([0-9]+)[-－]\s*([0-9]+)", r"图\1. \2", text)
    text = re.sub(r"表\s*([0-9]+)[-－]\s*([0-9]+)", r"表\1. \2", text)
    text = re.sub(r"21\s*维", "19维", text)
    return text


def split_sentences(text: str):
    parts = re.split(r"(?<=[。！？；;])", text)
    return [part.strip() for part in parts if part.strip()]


def compact_body_text(text: str) -> str:
    if is_code_like(text) or len(text) < 170:
        return text
    sentences = split_sentences(text)
    if len(sentences) < 3:
        return text

    must_keep = [
        "指标",
        "准确率",
        "精确率",
        "召回率",
        "F1",
        "误报率",
        "漏报率",
        "AUC",
        "如图",
        "如表",
        "式（",
        "公式",
    ]
    kept = []
    if len(text) >= 360:
        kept = sentences[:2] + sentences[-1:]
    elif len(text) >= 240:
        kept = sentences[:1] + sentences[-2:]
    else:
        kept = sentences[:1] + sentences[-1:]

    for sentence in sentences:
        if sentence not in kept and any(token in sentence for token in must_keep):
            kept.insert(max(1, len(kept) - 1), sentence)

    # Keep original order after adding protected metric/figure/formula sentences.
    ordered = []
    for sentence in sentences:
        if sentence in kept and sentence not in ordered:
            ordered.append(sentence)
    compacted = "".join(ordered)
    return compacted if len(compacted) < len(text) else text


def format_caption(text: str) -> str:
    text = normalize_body_text(text.strip())
    return re.sub(r"^([图表])\s*([0-9]+)\.\s*([0-9]+)\s*", r"\1 \2. \3 ", text)


def fit_image_width_mm(img_path: Path, requested_width: float = 120) -> float:
    max_width = 150.0
    max_height = 118.0
    width = min(float(requested_width), max_width)
    with Image.open(img_path) as im:
        px_w, px_h = im.size
    if px_w and px_h:
        height = width * px_h / px_w
        if height > max_height:
            width = max_height * px_w / px_h
    return max(42.0, round(width, 1))


def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.first_child_found_in("w:tcBorders")
    if old is not None:
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        if spec is None:
            continue
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), spec.get("val", "single"))
        border.set(qn("w:sz"), spec.get("sz", "4"))
        border.set(qn("w:color"), "000000")
        border.set(qn("w:space"), "0")
        borders.append(border)
    tc_pr.append(borders)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    if not widths:
        return
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue
            cell = row.cells[idx]
            cell.width = Mm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 56.7)))
            tc_w.set(qn("w:type"), "dxa")


def format_table(table, font_size=10.5):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    line_spacing = max(13, font_size + 6)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.line_spacing = Pt(line_spacing)
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    set_run_font(run, size=font_size, bold=(r_idx == 0))
            set_cell_border(cell, left={"val": "nil"}, right={"val": "nil"})
            if r_idx == 0:
                set_cell_border(
                    cell,
                    top={"val": "single", "sz": "12"},
                    bottom={"val": "single", "sz": "6"},
                    left={"val": "nil"},
                    right={"val": "nil"},
                )
            elif r_idx == len(table.rows) - 1:
                set_cell_border(
                    cell,
                    bottom={"val": "single", "sz": "12"},
                    left={"val": "nil"},
                    right={"val": "nil"},
                )


def latex_to_omml(latex: str):
    if not PANDOC.exists():
        raise FileNotFoundError(f"Pandoc not found: {PANDOC}")
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        md_path = tmp / "formula.md"
        docx_path = tmp / "formula.docx"
        md_path.write_text(f"$$\n{latex}\n$$\n", encoding="utf-8")
        subprocess.run(
            [str(PANDOC), "-f", "markdown+tex_math_dollars", "-t", "docx", "-o", str(docx_path), str(md_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        with zipfile.ZipFile(docx_path, "r") as zf:
            xml = zf.read("word/document.xml")

    from lxml import etree

    root = etree.fromstring(xml)
    math_nodes = root.xpath(".//m:oMathPara | .//m:oMath", namespaces={"m": MNS})
    if not math_nodes:
        raise ValueError(f"No OMML generated for formula: {latex}")
    return deepcopy(math_nodes[0])


def clear_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "nil"},
                bottom={"val": "nil"},
                left={"val": "nil"},
                right={"val": "nil"},
            )


def as_inline_math(node):
    from lxml import etree

    if etree.QName(node).localname == "oMathPara":
        math_nodes = node.xpath(".//m:oMath", namespaces={"m": MNS})
        if math_nodes:
            return deepcopy(math_nodes[0])
    return deepcopy(node)


def insert_formula_after(anchor: Paragraph, formula_data: dict) -> Paragraph:
    para = insert_paragraph_after(anchor, "", "论文正文")
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.first_line_indent = None
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = Pt(20)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    tabs = para.paragraph_format.tab_stops
    tabs.add_tab_stop(Mm(75), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.SPACES)
    tabs.add_tab_stop(Mm(150), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)

    para.add_run("\t")
    para._p.append(as_inline_math(latex_to_omml(formula_data["latex"])))
    run = para.add_run("\t" + formula_data.get("number", ""))
    set_run_font(run, size=12)
    return para


def patch_update_fields(docx_path: Path):
    with zipfile.ZipFile(docx_path, "r") as zf:
        files = {name: zf.read(name) for name in zf.namelist()}
    xml = files.get(
        "word/settings.xml",
        (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:settings xmlns:w="{WNS}"></w:settings>'
        ).encode("utf-8"),
    ).decode("utf-8")
    if "w:updateFields" not in xml:
        xml = xml.replace("</w:settings>", '<w:updateFields w:val="true"/></w:settings>')
    files["word/settings.xml"] = xml.encode("utf-8")
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    tmp.replace(docx_path)


def set_paragraph_bottom_border(para: Paragraph):
    p_pr = para._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")


def set_page_number_start(section, start=1, fmt: str | None = None):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))
    if fmt:
        pg_num.set(qn("w:fmt"), fmt)


def add_page_field(para: Paragraph, instruction="PAGE  \\* MERGEFORMAT", cached="1"):
    clear_paragraph(para)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruction)
    result = para.add_run(cached)
    set_run_font(result, size=10.5)
    para._p.remove(result._r)
    fld.append(result._r)
    para._p.append(fld)


def ensure_footer_page_number(section, instruction="PAGE  \\* MERGEFORMAT", cached="1"):
    section.footer.is_linked_to_previous = False
    section.footer_distance = Mm(12)
    footer = section.footer
    footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    add_page_field(footer_para, instruction=instruction, cached=cached)


def apply_main_body_section_format(doc: Document):
    # The source Word carries the main-body header in a later section.  After
    # replacing the old body, the final remaining section is the body section.
    for idx, front_section in enumerate(doc.sections[:-1]):
        if idx == 2:
            set_page_number_start(front_section, 1, "lowerRoman")
            ensure_footer_page_number(front_section, instruction="PAGE  \\* roman", cached="i")
        elif idx == 3:
            ensure_footer_page_number(front_section, instruction="PAGE  \\* roman", cached="iii")

    section = doc.sections[-1]
    section.header.is_linked_to_previous = False
    set_page_number_start(section, 1, "decimal")

    header = section.header
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    clear_paragraph(para)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("南京工业职业技术大学毕业设计（论文）")
    set_run_font(run, size=10.5)
    set_paragraph_bottom_border(para)

    ensure_footer_page_number(section, instruction="PAGE  \\* MERGEFORMAT", cached="1")


def collect_chapters():
    chapters = []
    chapter_indices = sorted(
        int(match.group(1))
        for path in HERE.glob("ch*.json")
        if (match := re.match(r"ch([0-9]+)\.json$", path.name))
    )
    for idx in chapter_indices:
        chapters.append(load_json(f"ch{idx}.json"))
    return chapters


def collect_toc_entries(chapters):
    cached_pages = {}
    current_toc_path = HERE / "toc_pages_current.json"
    legacy_toc_path = HERE / "toc_pages.json"
    if current_toc_path.exists():
        for title, page, _level in load_json(current_toc_path.name):
            cached_pages[title.replace("致  谢", "致谢").replace("附  录", "附录")] = str(page)
    elif legacy_toc_path.exists():
        for title, page, _level in load_json(legacy_toc_path.name):
            try:
                old_page = int(page)
            except (TypeError, ValueError):
                continue
            # The old draft counted abstract/TOC pages as Arabic pages.  The
            # Nanjing template restarts the main body at page 1.
            cached_pages[title.replace("致  谢", "致谢").replace("附  录", "附录")] = str(max(1, old_page - 5))

    entries = []
    for chapter in chapters:
        entries.append((chapter["title"], cached_pages.get(chapter["title"], ""), 1))
        for section in chapter.get("sections", []):
            entries.append((section["title"], cached_pages.get(section["title"], ""), 2))
            for subsection in section.get("subsections", []):
                entries.append((subsection["title"], cached_pages.get(subsection["title"], ""), 3))
    entries.append(("参考文献", cached_pages.get("参考文献", ""), 1))
    entries.append(("致谢", cached_pages.get("致谢", ""), 1))
    return entries


def collect_figure_entries(chapters):
    figures, tables = [], []

    def visit_items(items):
        for item in items:
            if not isinstance(item, dict):
                continue
            caption = item.get("caption")
            if item.get("type") == "image" and caption:
                figures.append(format_caption(caption))
            elif item.get("type") == "table" and caption:
                tables.append(format_caption(caption))

    for chapter in chapters:
        visit_items(chapter.get("content", []))
        for section in chapter.get("sections", []):
            visit_items(section.get("content", []))
            for subsection in section.get("subsections", []):
                visit_items(subsection.get("content", []))
    return figures, tables


def replace_cover_and_abstracts(doc: Document, meta: dict):
    if doc.tables:
        title_cell = doc.tables[0].rows[0].cells[-1]
        for para in title_cell.paragraphs:
            clear_paragraph(para)
        title_cell.paragraphs[0].add_run(meta.get("title_zh", ""))

    zh_heading = find_paragraph(doc, "摘要")
    zh_keyword = next(p for p in doc.paragraphs if (p.text or "").strip().startswith("关键词"))
    remove_between(zh_heading, zh_keyword)
    anchor = zh_heading
    for text in meta.get("abstract_zh_list", []):
        anchor = insert_paragraph_after(anchor, normalize_body_text(text), "论文正文")
    replace_paragraph_text(zh_keyword, "关键词：" + meta.get("keywords_zh", ""))

    en_heading = find_paragraph(doc, "ABSTRACT")
    en_keyword = next(p for p in doc.paragraphs if (p.text or "").strip().startswith("KEY WORDS"))
    remove_between(en_heading, en_keyword)
    anchor = en_heading
    for text in meta.get("abstract_en_list", []):
        anchor = insert_paragraph_after(anchor, text, "论文正文")
    replace_paragraph_text(en_keyword, "KEY WORDS: " + meta.get("keywords_en", ""))


def replace_toc_and_figure_list(doc: Document, chapters):
    toc_heading = find_paragraph(doc, "目录")
    fig_heading = find_paragraph(doc, "图表清单")
    remove_between(toc_heading, fig_heading)
    anchor = toc_heading
    toc_entries = collect_toc_entries(chapters)
    for idx, (title, page, level) in enumerate(toc_entries):
        para = insert_index_entry_after(anchor, title, page, f"toc {level}")
        if idx == 0:
            add_complex_field_begin(para, 'TOC \\o "1-3" \\h \\z \\u')
        anchor = para
    if toc_entries:
        anchor = insert_complex_field_end_after(anchor)

    body_start = find_first_body_heading(doc)
    remove_between(fig_heading, body_start)
    fig_heading.paragraph_format.page_break_before = True
    anchor = fig_heading
    figures, tables = collect_figure_entries(chapters)
    for caption in figures + tables:
        anchor = insert_index_entry_after(anchor, caption, "", "table of figures")
        anchor.paragraph_format.space_before = Pt(0)
        anchor.paragraph_format.space_after = Pt(0)
        anchor.paragraph_format.line_spacing = Pt(13)
        anchor.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        for run in anchor.runs:
            set_run_font(run, size=10)


def is_code_like(text: str) -> bool:
    stripped = text.strip()
    return bool(
        re.match(r"^(def |return |if |for |x_train|x_val|y_train|scaler|estimator|preds|proba|attack_type|[{}]|\")", stripped)
        or text.startswith("    ")
    )


def insert_text(anchor: Paragraph, text: str) -> Paragraph:
    text = compact_body_text(normalize_body_text(text))
    para = insert_paragraph_after(anchor, text, "论文正文")
    if is_code_like(text):
        para.paragraph_format.first_line_indent = None
        for run in para.runs:
            set_code_run_font(run, size=9)
    return para


def insert_image(anchor: Paragraph, item: dict) -> Paragraph:
    img = IMG_DIR / item["path"]
    width = fit_image_width_mm(img, item.get("width", 120))
    para = insert_paragraph_after(anchor, "", "论文正文")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(2)
    para.add_run().add_picture(str(img), width=Mm(width))
    anchor = para
    if item.get("caption"):
        cap = insert_paragraph_after(anchor, format_caption(item["caption"]), "图题样式")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_with_next = False
        anchor = cap
    return anchor


def insert_table(anchor: Paragraph, item: dict) -> Paragraph:
    if item.get("caption"):
        cap = insert_paragraph_after(anchor, format_caption(item["caption"]), "表题样式")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_with_next = True
        anchor = cap
    headers = item["headers"]
    rows = item.get("rows", [])
    widths = item.get("widths")
    table_width = sum(widths) if widths else item.get("width", 150)
    table = anchor._parent.add_table(rows=1, cols=len(headers), width=Mm(table_width))
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = str(value)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = str(value)
    set_table_widths(table, widths)
    format_table(table, item.get("font_size", 10.5))
    anchor._p.addnext(table._tbl)
    tail = OxmlElement("w:p")
    table._tbl.addnext(tail)
    return Paragraph(tail, anchor._parent)


def insert_content_item(anchor: Paragraph, item) -> Paragraph:
    if isinstance(item, str):
        return insert_text(anchor, item)
    if not isinstance(item, dict):
        return anchor
    item_type = item.get("type")
    if item_type == "image":
        return insert_image(anchor, item)
    if item_type == "table":
        return insert_table(anchor, item)
    if item_type == "formula":
        return insert_formula_after(anchor, item)
    if item.get("text"):
        return insert_text(anchor, item["text"])
    return anchor


def insert_section(anchor: Paragraph, section: dict, level: int) -> Paragraph:
    style = "Heading 2" if level == 2 else "Heading 3"
    title = strip_numbered_title(section["title"])
    anchor = insert_paragraph_after(anchor, title, style)
    for item in section.get("content", []):
        anchor = insert_content_item(anchor, item)
    for subsection in section.get("subsections", []):
        anchor = insert_section(anchor, subsection, level + 1)
    return anchor


def replace_body(doc: Document, chapters):
    first_heading = find_first_body_heading(doc)
    remove_after_until_sectpr(first_heading)
    replace_paragraph_text(first_heading, strip_chapter_title(chapters[0]["title"]))
    anchor = first_heading
    for item in chapters[0].get("content", []):
        anchor = insert_content_item(anchor, item)
    for section in chapters[0].get("sections", []):
        anchor = insert_section(anchor, section, 2)

    for chapter in chapters[1:]:
        anchor = insert_paragraph_after(anchor, strip_chapter_title(chapter["title"]), "Heading 1")
        for item in chapter.get("content", []):
            anchor = insert_content_item(anchor, item)
        for section in chapter.get("sections", []):
            anchor = insert_section(anchor, section, 2)

    anchor = insert_paragraph_after(anchor, "参考文献", "目录索引标题")
    anchor.paragraph_format.page_break_before = True
    for idx, ref in enumerate(load_json("references.json"), start=1):
        anchor = insert_paragraph_after(anchor, f"[{idx}] {ref}", "引用文献著录")

    ack = load_json("acknowledgement.json").get("text", "")
    anchor = insert_paragraph_after(anchor, "致谢", "目录索引加宽标题")
    anchor.paragraph_format.page_break_before = True
    for para_text in [p for p in ack.split("\n\n") if p.strip()]:
        anchor = insert_text(anchor, para_text.strip())


def normalize_final_styles(doc: Document):
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text.startswith("关键词：") or text.startswith("KEY WORDS:"):
            for run in para.runs:
                set_run_font(run, size=12)
        if para.style.name in {"图题样式", "表题样式"}:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_run_font(run, size=10.5, bold=False)
        if para.style.name == "论文正文":
            for run in para.runs:
                if not is_code_like(run.text):
                    set_run_font(run, size=12)
        if para.style.name == "引用文献著录":
            for run in para.runs:
                set_run_font(run, size=10.5)


def main() -> int:
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else default_output_path()
    output.parent.mkdir(parents=True, exist_ok=True)
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(SOURCE_DOCX)

    doc = Document(SOURCE_DOCX)
    meta = load_json("meta.json")
    chapters = collect_chapters()
    replace_cover_and_abstracts(doc, meta)
    replace_toc_and_figure_list(doc, chapters)
    replace_body(doc, chapters)
    normalize_final_styles(doc)
    apply_main_body_section_format(doc)
    doc.save(output)
    patch_update_fields(output)
    print(f"完成: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
